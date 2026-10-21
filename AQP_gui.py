#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
AQP_gui.py  --  Python port of AQP_gui_merged.m  (AutoQuant Data-Analytic Pipeline)
=====================================================================================

What it does (same pipeline as the MATLAB GUI):

    dataset folder with  CS_*.xlsx  (master calibration set, may hold the master
                                     transfer samples with Ref = NaN)
                         XRS_*.xlsx (transfer / scouting set on the target unit)
                         VAL_*.xlsx (validation set on the target unit)
      -> load + pick analyte (+ optional Spectra_Avg_Mean)
      -> match wavelength grids (master interpolated onto the target grid)
      -> calibration transfer  (woCabXfer | MDC | PDS99[wN] | ...)
      -> pretreatment PP1 then PP2  (SNV, SG 1st/2nd derivative, ...)
      -> PLS with training-set cross validation (sqrt(N)-folds by sample)
      -> optimal PLS factor (KneePt+1 / KneePt / Min RMSECV)
      -> prediction of the validation set, RMSEP, final model, plots
      -> Result4OUT_cln_AQP / Result4FinalModels / Results4PRO / StepByStep

Run it
    python AQP_gui.py                       # the GUI (tkinter + matplotlib)
    python AQP_gui.py -d 1 -w 1             # GUI, verbose console (like the -d/-w switches)
    python AQP_gui.py --cli --ds <folder> [options]     # no GUI, e.g. for batch work
    python AQP_gui.py --cli --help          # all command-line options

Requirements:  Python 3.9+, numpy, scipy, openpyxl, matplotlib  (tkinter for the GUI)

Output files are written under the folder the program is started from (the
MATLAB build used pwd the same way), or under --workdir.

Notes on the port
  * The numerical engine follows the MATLAB code line by line: MATLAB's
    plsregress/SIMPLS, sgolay, knee_pt, stdgen99/stdize99, RMSE definitions,
    fold assignment, median-based transfer standards, etc.
  * The MATLAB build handed stages to one another through Atrainpketc_*.mat
    files (later the in-memory AQP_mem store).  Here the stages pass Python
    objects; the deliverables (Results_AQP_*.mat, Beta_etc_FinalModel_*.mat,
    OpmModel_*.mat, StepByStep Atrainpketc_*.mat, PRO csv/xlsx) are still
    written, in MATLAB .mat format, so the MATLAB tools can read them.
  * .sam export (VIAVI SamLibrary.dll, .NET) and text-to-speech are Windows /
    MATLAB only and are not part of the port.  The csv/xlsx files for PRO are
    written as before.
  * Transfer schemes that need the commercial PLS_Toolbox (GLSw, OSC, stdgen
    based PDS) are not in the file and are reported as unavailable; the ones
    the GUI offers (woCabXfer, MDC, PDS99[w1/5/9]) and LS / GLSstd2 /
    MDC_GLSstd3 / STDgenize99 are implemented.
  * Figures open in their own windows (GUI) or are saved as PNG (--cli);
    MATLAB's .fig copies in Figs4OpmModel_AQP are PNG files here.
  * Where the MATLAB build stops with an error the port behaves as follows:
      - a dataset folder with only CS_*.xlsx (no XRS, no VAL) runs here with
        the training-set cross validation; the MATLAB build failed after the
        Self-Pred step;
      - Spectra_Avg_T-Mean_P-All together with XRS/VAL raises the same
        "mismatch in length between Xhat and Xmeasure" error as MATLAB;
      - Result4FinalModels/4json is written even when a VAL set exists (the
        MATLAB json step errored there and its catch skipped it).
  * Spectra_Avg_T-Mean_P-All is only listed on the panel for the administrator
    account, as in AQP_gui_OutputFcn (it is always available with --cli).

Verified against the original MATLAB code (run in GNU Octave with
MATLAB-compatible plsregress/sgolay) on the sugarcane CS / XRS / VAL data:
11 settings combinations (woCabXfer, MDC, PDS99[w1/5/9]; SNV, 1st/2nd SG
derivatives, none, PP1_PP2 xlsx; Spectra_Avg_All/Mean; KneePt+1 / KneePt /
Min; pro / lite; with and without XRS / VAL).  RMSEP, RMSECV curves and every
.mat / .csv / .xlsx deliverable agree to <= 1e-11 relative, and the same set
of files is written.
"""
from __future__ import annotations

import argparse
import base64
import copy
import datetime as _dt
import glob
import itertools
import math
import os
import re
import shutil
import sys
import time
import traceback
import warnings
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Sequence, Tuple

import numpy as np

__version__ = "2026.10-py1"

# =============================================================================
#   INTERNAL SETTINGS  --  EDIT THE VALUES IN THIS ONE PLACE
#   (the former INP_wo_handles.mat / inp4PLS.mat contents that the program uses)
# =============================================================================
INTERNAL_SETTINGS = {
    # inp4PLS.PlsfactorScan  -> PlsfactorScan_default (trimmed per dataset)
    "PlsfactorScan": list(range(2, 13)),
    # sample selection for the transfer standards (INP.cXM_Slct_scheme)
    "cXM_Slct_scheme": ["KSall"],
    "TP_includeTrn_Yes": 1,
    # cross-validation scheme used inside PLS_inside_PLS_predict_ONLY_MLtool
    "Tcv_scheme": "sqrtNSfolds-Conc",
    # Option_Output_SAM  (the .sam export itself is Windows-only; kept as a flag)
    "Option_Output_SAM": 1,
    # Results comparison folder name: include analyte / CabXfer scheme ?
    "NEWDIR_include_AnaName_yes": 0,
    "NEWDIR_include_CabXfer_yes": 0,
    # XTick guard policy when storing a Results_AQP file: 'divert'|'save-anyway'|'skip'
    "XTick_policy": "divert",
    # keep this many Results4PRO\tmpN folders
    "Ntmp_keep": 2,
}

# The menus of AQP_gui_merged.fig, entry by entry (row 1 is the title row
# where the .fig has one).
GUI_MENUS = {
    "Load_DS": ["Load_DS", "Browse"],
    "AQP_class": ["pro", "lite"],
    "Spectra_Avg_Method": ["Spectra_Avg_All", "Spectra_Avg_Mean", "Spectra_Avg_T-Mean_P-All"],
    "cPP1": ["PP1", "PP1_PP2_xlsx", "SNV", "1stDerSGFL5[PO2]", "1stDerSGFL7[PO2]",
             "1stDerSGFL9[PO2]", "1stDerSGFL11[PO2]", "1stDerSGFL13[PO2]",
             "2ndDerSGFL5[PO2]", "2ndDerSGFL7[PO2]", "2ndDerSGFL9[PO2]",
             "2ndDerSGFL11[PO2]", "2ndDerSGFL13[PO2]", "none"],
    "cPP2": ["PP2", "PP1_PP2_xlsx", "SNV", "1stDerSGFL5[PO2]", "1stDerSGFL7[PO2]",
             "1stDerSGFL9[PO2]", "1stDerSGFL11[PO2]", "1stDerSGFL13[PO2]",
             "2ndDerSGFL5[PO2]", "2ndDerSGFL7[PO2]", "2ndDerSGFL9[PO2]",
             "2ndDerSGFL11[PO2]", "2ndDerSGFL13[PO2]", "none"],
    "cDataFlow": ["Xfer-PP"],
    "CabXfer_scheme": ["CabXfer_scheme", "woCabXfer", "MDC", "PDS99[w1]", "PDS99[w5]", "PDS99[w9]"],
    "UDMas_scheme": ["woUDM", "Scv", "UDMasCS", "ALL_Above_UDM_schemes",
                     "-----------------------------", "woUDM", "OpmPLS-ONLY", "UDM-ONLY",
                     "Split_OddUDM_EvenXRS", "Scv", "Both_UDM_Opm", "UDMasCS", "ALL_UDM_schemes"],
    "TcvModelParaOpmScheme": ["TcvModelParaOpmScheme", "KneePt+1_RMSECV (default)",
                              "KneePt-RMSECV", "Min-RMSECV"],
    "Lines_Xticks": ["Lines_PP1--Xticks_CabXfer", "Lines_PP1--Xticks_UDM",
                     "Lines_UDM--Xticks_DSn", "Lines_DSn--Xticks_UDM", "Lines--Xticks"],
    "StepbyStep_plots": ["StepbyStep", "CS_XRSmst_Bef_MGs", "CS_Val_CabXfer_PP"],
}

# MicroNIR wavelength grid (get_MN_wvl)
MN_WVL = np.array([
    908.1, 914.29435, 920.48871, 926.68306, 932.87742, 939.07177, 945.26613, 951.46048,
    957.65484, 963.84919, 970.04355, 976.2379, 982.43226, 988.62661, 994.82097, 1001.01532,
    1007.20968, 1013.40403, 1019.59839, 1025.79274, 1031.9871, 1038.18145, 1044.37581,
    1050.57016, 1056.76452, 1062.95887, 1069.15323, 1075.34758, 1081.54194, 1087.73629,
    1093.93065, 1100.125, 1106.31935, 1112.51371, 1118.70806, 1124.90242, 1131.09677,
    1137.29113, 1143.48548, 1149.67984, 1155.87419, 1162.06855, 1168.2629, 1174.45726,
    1180.65161, 1186.84597, 1193.04032, 1199.23468, 1205.42903, 1211.62339, 1217.81774,
    1224.0121, 1230.20645, 1236.40081, 1242.59516, 1248.78952, 1254.98387, 1261.17823,
    1267.37258, 1273.56694, 1279.76129, 1285.95565, 1292.15, 1298.34435, 1304.53871,
    1310.73306, 1316.92742, 1323.12177, 1329.31613, 1335.51048, 1341.70484, 1347.89919,
    1354.09355, 1360.2879, 1366.48226, 1372.67661, 1378.87097, 1385.06532, 1391.25968,
    1397.45403, 1403.64839, 1409.84274, 1416.0371, 1422.23145, 1428.42581, 1434.62016,
    1440.81452, 1447.00887, 1453.20323, 1459.39758, 1465.59194, 1471.78629, 1477.98065,
    1484.175, 1490.36935, 1496.56371, 1502.75806, 1508.95242, 1515.14677, 1521.34113,
    1527.53548, 1533.72984, 1539.92419, 1546.11855, 1552.3129, 1558.50726, 1564.70161,
    1570.89597, 1577.09032, 1583.28468, 1589.47903, 1595.67339, 1601.86774, 1608.0621,
    1614.25645, 1620.45081, 1626.64516, 1632.83952, 1639.03387, 1645.22823, 1651.42258,
    1657.61694, 1663.81129, 1670.00565, 1676.2])


class AQPError(Exception):
    """An error the pipeline raises on purpose (shown to the user as-is)."""


# =============================================================================
#   CONSOLE VERBOSITY  (-d / -w switches of the MATLAB build)
# =============================================================================
class _Verbosity:
    disp = 0       # -d : 0 = quiet, 1 = print progress
    warn = 0       # -w : 0 = no warnings, 1 = warnings


VERB = _Verbosity()


def disp(*msg):
    if VERB.disp:
        print(*msg, flush=True)


def disp_with_border(msg):
    if VERB.disp:
        b = "+" * (len(str(msg)) + 16)
        print(b + "\n        " + str(msg) + "        \n" + b, flush=True)


def warn(msg):
    if VERB.warn:
        print("Warning: " + str(msg), file=sys.stderr, flush=True)


def note(msg):
    """Always printed: things the user must know (the MATLAB build spoke them)."""
    print(str(msg), flush=True)


MSGBOX_HOOK = None      # the GUI sets this to show MATLAB's (non-modal) msgbox


def msgbox(msg):
    note(msg)
    if MSGBOX_HOOK is not None:
        try:
            MSGBOX_HOOK(str(msg))
        except Exception:
            pass


# =============================================================================
#   MATLAB-COMPATIBILITY HELPERS
# =============================================================================
def num2str(x) -> str:
    """MATLAB num2str for a real scalar (or a 1-D vector, space separated)."""
    if isinstance(x, str):
        return x
    a = np.atleast_1d(np.asarray(x, dtype=float))
    if a.size == 0:
        return ""
    if a.size > 1:
        return "  ".join(num2str(v) for v in a)
    v = float(a[0])
    if math.isnan(v):
        return "NaN"
    if math.isinf(v):
        return "Inf" if v > 0 else "-Inf"
    if v == int(v) and abs(v) < 1e15:
        return str(int(v))
    xmax = abs(v)
    dgt = math.floor(math.log10(xmax)) if xmax > 0 else 0
    dgt = max(dgt + 5, 5)
    dgt = min(dgt, 16)
    return ("%." + str(dgt) + "g") % v


def roundns(x, d=None) -> str:
    """roundns(x,d): round to d decimals and return num2str of it (a char)."""
    x = float(x)
    if d is None:
        y = round_half_away(x)
    else:
        y = round_half_away(10 ** d * x) / 10 ** d
    return num2str(y)


def round_half_away(x):
    """MATLAB round: halves away from zero."""
    return float(np.sign(x) * np.floor(np.abs(x) + 0.5))


def unique_sorted(cstr: Sequence[str]) -> List[str]:
    """MATLAB unique() on a cellstr: sorted by char code."""
    return sorted(set(cstr))


def unique_appear_order(cstr: Sequence[str]) -> List[str]:
    out, seen = [], set()
    for s in cstr:
        if s not in seen:
            seen.add(s)
            out.append(s)
    return out


def sortnat(c: Sequence[str]) -> List[str]:
    """Natural-order sort, same order as MATLAB sortnat (Cobeldick 2012):
    case-insensitive; runs of digits compare by value, every other character
    one at a time by ASCII; at each position a shorter string < a number < a
    character.  Stable, like MATLAB's sortrows.
    e.g. 1stDerSGFL11, 1stDerSGFL5, 1stDerSGFL9 -> ...SGFL5, ...SGFL9, ...SGFL11"""
    return sorted(c, key=_natkey)


def _natkey(s) -> list:
    """sort key behind sortnat (use it to get a natural-order index)"""
    return [(0, int(num), "") if num else (1, 0, ch)
            for num, ch in re.findall(r"(\d+)|(.)", str(s).lower(), flags=re.S)]


def _findall_overlap(s: str, sub: str) -> List[int]:
    """findstr: every (overlapping) occurrence, 0-based."""
    out, k = [], s.find(sub)
    while k >= 0 and sub:
        out.append(k)
        k = s.find(sub, k + 1)
    return out


def find_keyword_between_markers(s: str, m1: str, m2: str):
    """find_keyword_between_markers: text between the LAST marker1 and the
    first marker2 that follows it; None when not found ([] in MATLAB)."""
    if s is None:
        return None
    if m1 == "":
        if m2 == "":
            return s
        k = s.find(m2)
        return None if k < 0 else s[:k]
    starts = [k + len(m1) for k in _findall_overlap(s, m1)]
    if not starts:
        return None
    st = starts[-1]
    if m2 == "":
        return s[st:]
    later = [k for k in _findall_overlap(s, m2) if k > st]
    if not later:
        return None
    return s[st:later[0]]


def find_keynumber_after(s: str, marker: str) -> float:
    """find_keynumber_numeric_AFTER_marker: the integer right after marker."""
    m = re.search(re.escape(marker) + r"\D*?(\d+)", s)
    if not m:
        return float("nan")
    # MATLAB: pat1 = [marker,'\D*\d{1}'] end ... pat2 = [marker,'\D*\d+'] end
    m1 = re.search(re.escape(marker) + r"\D*\d", s)
    m2 = re.search(re.escape(marker) + r"\D*\d+", s)
    try:
        return float(s[m1.end() - 1:m2.end()])
    except Exception:
        return float("nan")


def strip_after(s: str, kw: str) -> str:
    """textual_eraseAfter_wo_kw1 for a char: drop kw and everything after it."""
    k = s.find(kw)
    if k < 0:
        return s
    out = s[:k]
    return out if out else s


def is_odd(n) -> bool:
    return int(n) % 2 == 1


def mat_median_rows(A: np.ndarray) -> np.ndarray:
    """median(A) for a matrix (column medians); a single row is returned as is."""
    A = np.atleast_2d(A)
    if A.shape[0] == 1:
        return A[0].copy()
    return np.median(A, axis=0)


# =============================================================================
#   NUMERICAL CORE
# =============================================================================
def plsregress(X: np.ndarray, Y: np.ndarray, ncomp: int):
    """MATLAB Statistics Toolbox plsregress (SIMPLS), exactly as in plsregress.m.

    Returns Xloadings, Yloadings, Xscores, Yscores, beta ((p+1) x m, intercept
    first), Weights.
    """
    X = np.asarray(X, dtype=float)
    Y = np.asarray(Y, dtype=float)
    if Y.ndim == 1:
        Y = Y[:, None]
    n, dx = X.shape
    maxncomp = min(n - 1, dx)
    if ncomp > maxncomp:
        raise AQPError(f"plsregress: NCOMP={ncomp} must be <= {maxncomp} "
                       f"(the number of samples minus one)")
    meanX = X.mean(axis=0)
    meanY = Y.mean(axis=0)
    X0 = X - meanX
    Y0 = Y - meanY
    dy = Y0.shape[1]
    Xl = np.zeros((dx, ncomp))
    Yl = np.zeros((dy, ncomp))
    Xs = np.zeros((n, ncomp))
    Ys = np.zeros((n, ncomp))
    W = np.zeros((dx, ncomp))
    V = np.zeros((dx, ncomp))
    Cov = X0.T @ Y0
    for i in range(ncomp):
        if dy == 1:
            c = Cov[:, 0]
            si = np.sqrt(c @ c)
            ri = c / si
            ci = np.array([1.0])
        else:
            u, s, vt = np.linalg.svd(Cov, full_matrices=False)
            ri, si, ci = u[:, 0], s[0], vt[0, :]
        ti = X0 @ ri
        normti = np.sqrt(ti @ ti)
        ti = ti / normti
        Xl[:, i] = X0.T @ ti
        qi = si * ci / normti
        Yl[:, i] = qi
        Xs[:, i] = ti
        Ys[:, i] = Y0 @ qi
        W[:, i] = ri / normti
        vi = Xl[:, i].copy()
        for _ in range(2):
            for j in range(i):
                vj = V[:, j]
                vi = vi - (vj @ vi) * vj
        vi = vi / np.sqrt(vi @ vi)
        V[:, i] = vi
        Cov = Cov - np.outer(vi, vi @ Cov)
        Vi = V[:, :i + 1]
        Cov = Cov - Vi @ (Vi.T @ Cov)
    for i in range(ncomp):
        ui = Ys[:, i].copy()
        for _ in range(2):
            for j in range(i):
                tj = Xs[:, j]
                ui = ui - (tj @ ui) * tj
        Ys[:, i] = ui
    beta = W @ Yl.T
    beta = np.vstack([meanY - meanX @ beta, beta])
    return Xl, Yl, Xs, Ys, beta, W


def pls_predict(X: np.ndarray, beta: np.ndarray) -> np.ndarray:
    """[ones(n,1) X]*beta  (single response -> 1-D)."""
    X = np.atleast_2d(X)
    y = beta[0] + X @ beta[1:]
    return y[:, 0] if y.ndim == 2 and y.shape[1] == 1 else y


def rms_error_woNaN(xhat, xmeas) -> float:
    """RMS_error_woNaN: sqrt(nansum(d^2)/(N-1)); NaN when N<=1."""
    xhat = np.ravel(np.asarray(xhat, dtype=float))
    xmeas = np.ravel(np.asarray(xmeas, dtype=float))
    if xhat.size != xmeas.size:
        raise AQPError("mismatch in length between Xhat and Xmeasure")
    if xhat.size == 0:
        return float("nan")
    d = xhat - xmeas
    N = int(np.sum(~np.isnan(d)))
    if N <= 1:
        return float("nan")
    return float(np.sqrt(np.nansum(d ** 2) / (N - 1)))


def rms_error_woNaN_N(xhat, xmeas) -> float:
    """RMS_error_woNaN_N: RMSE based on N (sqrt(RMSE_Nm1^2*(L-1)/L))."""
    xhat = np.ravel(np.asarray(xhat, dtype=float))
    xmeas = np.ravel(np.asarray(xmeas, dtype=float))
    if xhat.size == 1:
        return float(abs(xhat[0] - xmeas[0]))
    r = rms_error_woNaN(xhat, xmeas)
    L = xhat.size
    return float(np.sqrt(r ** 2 * (L - 1) / L))


def knee_pt(y, x=None):
    """knee_pt.m (Dmitry Kaplan) with use_absolute_dev_p = true."""
    y = np.ravel(np.asarray(y, dtype=float))
    if y.size < 3:
        raise AQPError("knee_pt: y must be at least 3 elements long")
    if x is None:
        x = np.arange(1, y.size + 1, dtype=float)
    else:
        x = np.ravel(np.asarray(x, dtype=float))
    if np.any(np.diff(x) < 0):
        idx = np.argsort(x, kind="stable")
        y = y[idx]
        x = x[idx]
    else:
        idx = np.arange(x.size)
    with np.errstate(divide="ignore", invalid="ignore"):
        n = np.arange(1, y.size + 1, dtype=float)
        sxy = np.cumsum(x * y); sx = np.cumsum(x); sy = np.cumsum(y); sxx = np.cumsum(x * x)
        det = n * sxx - sx * sx
        mfwd = (n * sxy - sx * sy) / det
        bfwd = -(sx * sxy - sxx * sy) / det
        xr, yr = x[::-1], y[::-1]
        sxy = np.cumsum(xr * yr); sx = np.cumsum(xr); sy = np.cumsum(yr); sxx = np.cumsum(xr * xr)
        det = n * sxx - sx * sx
        mbck = ((n * sxy - sx * sy) / det)[::-1]
        bbck = (-(sx * sxy - sxx * sy) / det)[::-1]
        err = np.full(y.size, np.nan)
        for bp in range(1, y.size):          # MATLAB breakpt = 2:length(y)
            dfwd = (mfwd[bp] * x[:bp + 1] + bfwd[bp]) - y[:bp + 1]
            dbck = (mbck[bp] * x[bp:] + bbck[bp]) - y[bp:]
            err[bp] = np.sum(np.abs(dfwd)) + np.sum(np.abs(dbck))
    if np.all(np.isnan(err)):
        loc = 0
    else:
        loc = int(np.nanargmin(err))
    return float(x[loc]), int(idx[loc])


def sgolay_G(k: int, F: int) -> np.ndarray:
    """G output of MATLAB sgolay(k,F): matrix of differentiation filters."""
    S = np.arange(-(F - 1) / 2, (F - 1) / 2 + 1)[:, None] ** np.arange(k + 1)[None, :]
    Q, R = np.linalg.qr(S)
    Rinv = np.linalg.inv(R)
    return S @ Rinv @ Rinv.T


def sgolay_B(k: int, F: int) -> np.ndarray:
    S = np.arange(-(F - 1) / 2, (F - 1) / 2 + 1)[:, None] ** np.arange(k + 1)[None, :]
    G = sgolay_G(k, F)
    return G @ S.T


def apply_NstDer(X: np.ndarray, poly_order: int, width: int, DN: int) -> np.ndarray:
    """apply_NstDer: SG derivative by conv with p!/(-1)^p * G(:,p+1), 'same', trimmed."""
    if poly_order < DN:
        raise AQPError("poly_order should be at least equal or larger than DN i.e. DNthDer")
    g = sgolay_G(poly_order, width)
    p = DN
    kern = math.factorial(p) / ((-1.0) ** p) * g[:, p]
    X = np.atleast_2d(np.asarray(X, dtype=float))
    out = np.empty((X.shape[0], X.shape[1]))
    for i in range(X.shape[0]):
        out[i] = np.convolve(X[i], kern, mode="same")
    h = width // 2
    return out[:, h:X.shape[1] - h]


def savitzky_golay_coeffs(F: int, N: int, DN: int) -> np.ndarray:
    """savitzkyGolay(pp,N,DN,pp,ones) of the file (Diederick Niehorster)."""
    x = np.arange(-(F // 2), F // 2 + 1, dtype=float)
    V = np.cumprod(np.hstack([np.ones((F, 1)), x[:, None] * np.ones((1, N))]), axis=1)
    df = np.linalg.lstsq(V, np.eye(F), rcond=None)[0].T          # (V \ eye).'
    x0 = x
    hx = np.hstack([np.zeros((F, DN)), np.ones((F, 1)) * math.prod(range(1, DN + 1))])
    for k in range(1, N - DN + 1):
        hx = np.hstack([hx, (x0 ** k)[:, None] * math.prod(range(k + 1, DN + k + 1))])
    return df @ hx.T


def savitzky_golay_filt_rows(X: np.ndarray, N: int, DN: int, F: int) -> np.ndarray:
    """savitzkyGolayFilt applied along each row (spectra in rows)."""
    B = savitzky_golay_coeffs(F, N, DN)
    X = np.atleast_2d(X)
    out = np.empty_like(X, dtype=float)
    h = (F - 1) // 2
    for r in range(X.shape[0]):
        x = X[r]
        y = np.zeros(x.size)
        y[:h] = np.fliplr(B[:, h + 1:]).T @ x[:F][::-1]
        yc = np.convolve(x, B[:, h], mode="full")
        y[h:x.size - h] = yc[F - 1:x.size]
        y[x.size - h:] = np.fliplr(B[:, :h]).T @ x[x.size - F:][::-1]
        out[r] = y
    return out


def sgolayfilt_rows(X: np.ndarray, K: int, F: int) -> np.ndarray:
    """MATLAB sgolayfilt(x',K,F)' : smoothing along each spectrum."""
    B = sgolay_B(K, F)
    X = np.atleast_2d(X)
    out = np.empty_like(X, dtype=float)
    h = (F - 1) // 2
    for r in range(X.shape[0]):
        x = X[r]
        n = x.size
        ybegin = np.fliplr(B[:, h + 1:]).T @ x[:F][::-1]
        yc = np.convolve(x, B[:, h], mode="full")[:n]
        yend = np.fliplr(B[:, :h]).T @ x[n - F:][::-1]
        out[r] = np.concatenate([ybegin, yc[F - 1:], yend])
    return out


def apply_SNV(X: np.ndarray) -> np.ndarray:
    """auto(x')': each spectrum to mean 0 / std 1 (std with N-1)."""
    X = np.atleast_2d(np.asarray(X, dtype=float))
    m = X.mean(axis=1, keepdims=True)
    s = X.std(axis=1, ddof=1, keepdims=True)
    return (X - m) / s


def mncn(X: np.ndarray):
    X = np.atleast_2d(np.asarray(X, dtype=float))
    mx = X.mean(axis=0)
    return X - mx, mx


def mscorr(x: np.ndarray, xref: np.ndarray):
    """mscorr(x,xref) with mean centring (mc=1)."""
    x = np.atleast_2d(x)
    xref = np.ravel(xref)
    sx_, alpha = mncn(x.T)              # columns = samples
    xr, mx = mncn(xref[:, None])
    beta = np.linalg.lstsq(xr, sx_, rcond=None)[0].ravel()
    alpha = alpha - mx[0] * beta
    return (x - alpha[:, None]) / beta[:, None]


def preprocess_NIR_spectra(X: np.ndarray, pp: str, xref_msc=None):
    """Port of preprocess_NIR_spectra.  Returns (X_pp, xref_msc_or_None)."""
    X = np.atleast_2d(np.asarray(X, dtype=float))
    out_xref = None
    if "DerSG" in pp:
        if "DerSGw" in pp:
            pp_alt = pp[:pp.find("DerSGw")] + "DerSGw"
        elif "DerSGFL" in pp:
            pp_alt = pp[:pp.find("DerSGFL")] + "DerSGFL"
        else:
            raise AQPError("unexpected case ???")
    else:
        pp_alt = pp
    isPRO = find_keyword_between_markers(pp, "{", "}") == "PRO"
    if pp_alt[:3] == "SGw":
        pp_alt = "SGw"
    if pp_alt in ("SGw", "SG"):
        sF = find_keyword_between_markers(pp, "SGw", "")
        try:
            F = int(sF) if sF else 5
        except ValueError:
            F = 5
        return sgolayfilt_rows(X, 3, F), None
    if pp_alt == "1stDerSGDiederick":
        Y = -1 * savitzky_golay_filt_rows(X, 3, 1, 5)
        return Y[:, 2:X.shape[1] - 2], None
    if pp_alt in ("0thDerSGFL", "1stDerSGw", "1stDerSGFL", "2ndDerSGFL", "3rdDerSGFL", "4thDerSGFL"):
        po = None
        k = find_keyword_between_markers(pp, "_PO", "")
        if k:
            try:
                po = int(k)
            except ValueError:
                po = None
        if po is None:
            k = find_keyword_between_markers(pp, "[PO", "]")
            try:
                po = int(k) if k else 3
            except ValueError:
                po = 3
        width = find_keynumber_after(pp, pp_alt)
        if math.isnan(width) or not is_odd(int(width)):
            raise AQPError("width for SG filter must be odd number")
        width = int(width)
        if pp_alt == "1stDerSGw":
            Y = -1 * savitzky_golay_filt_rows(X, po, 1, width)
            h = (width - 1) // 2
            return Y[:, h:X.shape[1] - h], None
        DN = {"0thDerSGFL": 0, "1stDerSGFL": 1, "2ndDerSGFL": 2, "3rdDerSGFL": 3, "4thDerSGFL": 4}[pp_alt]
        Y = apply_NstDer(X, po, width, DN)
        if isPRO:
            nz = (X.shape[1] - Y.shape[1]) // 2
            Y = np.hstack([np.zeros((Y.shape[0], nz)), Y, np.zeros((Y.shape[0], nz))])
        return Y, None
    if pp_alt == "2ndDer":
        raise AQPError("'2ndDer' (Andy Hulse loc_preprocess) is not part of the port; use 2ndDerSGFLn[POk]")
    if pp_alt == "1stDer":
        raise AQPError("'1stDer' (Andy Hulse loc_preprocess) is not part of the port; use 1stDerSGFLn[POk]")
    if pp_alt in ("SNV", "SNV{PRO}"):
        if isPRO:
            Y = X.copy()
            keep = ~np.all(X == 0, axis=0)
            Y[:, keep] = apply_SNV(X[:, keep])
            return Y, None
        return apply_SNV(X), None
    if pp_alt in ("SampMncn", "SMC"):
        return X - X.mean(axis=1, keepdims=True), None
    if pp_alt == "MSC":
        if xref_msc is None:
            xref = X.mean(axis=0)
            out_xref = xref
        else:
            xref = xref_msc
        return mscorr(X, xref), out_xref
    if pp_alt == "none":
        return X.copy(), None
    raise AQPError(f"{pp} not supported !!!")


def apply_pp1_pp2(Xt: np.ndarray, pp1: str, pp2: str, Xp: Optional[np.ndarray] = None):
    """pretreat_preprocess_RawSpectra_pp1_pp2_AT_TP for a T (and P) pair.

    The P set uses the T set's MSC reference (if any PP is MSC)."""
    t1, xr1 = preprocess_NIR_spectra(Xt, pp1)
    t2, xr2 = preprocess_NIR_spectra(t1, pp2)
    if Xp is None:
        return t2, None
    if np.all(np.isnan(Xp)):
        return t2, None
    p1, _ = preprocess_NIR_spectra(Xp, pp1, xr1)
    p2, _ = preprocess_NIR_spectra(p1, pp2, xr2)
    return t2, p2


# ----- calibration transfer standards ---------------------------------------
def stdgen99(spec1: np.ndarray, spec2: np.ndarray, window: int, tol: float = 1e-2, maxpc=None):
    """stdgen99 (PLS_Toolbox 1999, in the MATLAB file): PDS transfer matrix.

    spec1: target-unit transfer spectra, spec2: master-unit transfer spectra.
    Returns (stdmat, stdvect) such that  spec2_new*stdmat + stdvect ~ spec1."""
    spec1 = np.asarray(spec1, dtype=float)
    spec2 = np.asarray(spec2, dtype=float)
    ms, ns = spec1.shape
    if spec2.shape != spec1.shape:
        raise AQPError("Both spectra must have the same number of samples and channels")
    if window != 0 and window % 2 == 0:
        warn("The number of channels in the window should really be an odd number")
    mspec1, mns1 = mncn(spec1)
    mspec2, mns2 = mncn(spec2)
    if maxpc is None:
        maxpc = ms
    if window == 0:
        if ms <= ns:
            u, s, vt = np.linalg.svd(mspec2.T, full_matrices=False)
            invs = np.zeros((ms, ms))
            invs[:ms - 1, :ms - 1] = np.diag(1.0 / s[:ms - 1])
            spec2inv = u @ invs @ vt
        else:
            spec2inv = np.linalg.pinv(mspec2)
        stdmat = spec2inv @ mspec1
    else:
        stdmat = np.zeros((ns, ns))
        ind1 = window // 2
        ind2 = window - ind1 - 1
        for i in range(ns):                    # 0-based column i  (MATLAB i+1)
            i1 = i + 1
            if i1 <= ind1:
                lo, hi = 0, i1 + ind2
            elif i1 >= ns - ind2 + 1:
                lo, hi = i1 - ind1 - 1, ns
            else:
                lo, hi = i1 - ind1 - 1, i1 + ind2
            xs2 = mspec2[:, lo:hi]
            wind = hi - lo
            A = xs2.T @ xs2
            u, s, vt = np.linalg.svd(A)
            ratio = s / (s[0] * np.ones(wind)) if s[0] != 0 else np.zeros(wind)
            sinds = int(np.sum(ratio > tol))
            sinds = min(sinds, maxpc)
            sinv = np.zeros((wind, wind))
            if sinds > 0:
                sinv[:sinds, :sinds] = np.linalg.inv(np.diag(s[:sinds]))
            mod = u @ sinv @ vt @ xs2.T @ spec1[:, i]
            stdmat[lo:hi, i] = mod
    stdvect = (mns1 - stdmat.T @ mns2)
    return stdmat, stdvect


def stdize99(nspec: np.ndarray, stdmat: np.ndarray, stdvect: np.ndarray) -> np.ndarray:
    return np.asarray(nspec) @ stdmat + stdvect[None, :]


def GLSstd(mx: np.ndarray, sx: np.ndarray, alpha: float, version: int = 2):
    """GLSstd2 (version=2, transfer samples adjusted to the same mean) and
    GLSstd3 (version=3) of the MATLAB file.  Returns BadWgtMatrX, meandif."""
    mx = np.asarray(mx, dtype=float)
    sx = np.asarray(sx, dtype=float)
    m, n = mx.shape
    if m == 1:
        raise AQPError("you can not run GLS with only one transfer sample")
    if sx.shape[0] != m:
        raise AQPError("mx and sx should have same number of samples")
    meandif = np.mean(mx - sx, axis=0)
    dif = mx - (sx + meandif) if version == 2 else mx - sx
    BadCovX = np.cov(dif, rowvar=False, ddof=1) + np.eye(n) * alpha
    v, s, _ = np.linalg.svd(BadCovX)
    ev = s / np.mean(s)
    sv = np.sqrt(ev)
    inv = np.zeros_like(sv)
    keep = sv > n * np.max(sv) * np.finfo(float).eps
    inv[keep] = 1.0 / sv[keep]
    return 2 * v @ np.diag(inv) @ v.T, meandif


# =============================================================================
#   DATA STRUCTURES
#   SpecSet ~ one "Atrainpketc" block with saConc (one spectrum per row)
#   TPSet   ~ an Atrainpketc file with a T (training) and a P (prediction) set
# =============================================================================
@dataclass
class SpecSet:
    Atrainpk: np.ndarray                 # spectra after the current pretreatment
    RawSpectra: np.ndarray               # spectra before pretreatment
    wvl: np.ndarray                      # wvl_standardize (grid of RawSpectra)
    AclabelT: List[str]                  # sample ID per row (saConc.SampleName)
    conc: np.ndarray                     # reference value per row (NaN = none)
    clsname: str                         # analyte name
    clistclslabel: List[str] = field(default_factory=list)
    AclassinfoT: Optional[np.ndarray] = None
    AInfo_1: Any = ""
    SAT_orig_Avg_All: Optional["SpecSet"] = None    # only for Spectra_Avg_T-Mean_P-All
    seqnum: Optional[np.ndarray] = None             # saConc.seqnum (row number in the sheet)
    extra: Dict[str, Any] = field(default_factory=dict)   # other LAT fields to keep (wvl, wvl_AT)

    def __post_init__(self):
        if self.seqnum is None:
            self.seqnum = np.arange(1, np.atleast_2d(self.Atrainpk).shape[0] + 1, dtype=float)
        self.Atrainpk = np.atleast_2d(np.asarray(self.Atrainpk, dtype=float))
        self.RawSpectra = np.atleast_2d(np.asarray(self.RawSpectra, dtype=float))
        self.conc = np.ravel(np.asarray(self.conc, dtype=float))
        self.wvl = np.ravel(np.asarray(self.wvl, dtype=float))
        if not self.clistclslabel:
            self.clistclslabel = unique_sorted(self.AclabelT)
        if self.AclassinfoT is None:
            pos = {c: i + 1 for i, c in enumerate(self.clistclslabel)}
            self.AclassinfoT = np.array([pos.get(a, np.nan) for a in self.AclabelT], dtype=float)

    @property
    def n(self) -> int:
        return self.Atrainpk.shape[0]

    def subset(self, rows) -> "SpecSet":
        rows = list(rows)
        return SpecSet(self.Atrainpk[rows].copy(), self.RawSpectra[rows].copy(), self.wvl.copy(),
                       [self.AclabelT[i] for i in rows], self.conc[rows].copy(), self.clsname,
                       list(self.clistclslabel), np.asarray(self.AclassinfoT)[rows].copy(),
                       self.AInfo_1, None, np.asarray(self.seqnum)[rows].copy(), dict(self.extra))

    def copy(self) -> "SpecSet":
        return copy.deepcopy(self)


@dataclass
class TPSet:
    T: SpecSet
    P: Optional[SpecSet]
    tag: str = ""                          # what the MATLAB file name carried, e.g. '{T-CS_P-Val}'
    saCTCP: Any = None
    RawSpectra_Tset_orig_woCabXfer: Optional[np.ndarray] = None

    def copy(self) -> "TPSet":
        return copy.deepcopy(self)


# =============================================================================
#   XLSX READING   (xlsread [NUM,TXT,RAW] semantics + clistfilename2AT_AQP)
# =============================================================================
_XLS_CACHE: Dict[Tuple[str, float], List[List[Any]]] = {}


def read_xlsx_raw(path: str) -> List[List[Any]]:
    """RAW of xlsread: numbers as float, text as str, empty as NaN (first sheet)."""
    key = (os.path.abspath(path), os.path.getmtime(path))
    if key in _XLS_CACHE:
        return _XLS_CACHE[key]
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = []
    for r in ws.iter_rows(values_only=True):
        row = []
        for v in r:
            if v is None:
                row.append(float("nan"))
            elif isinstance(v, bool):
                row.append(float(v))
            elif isinstance(v, (int, float)):
                row.append(float(v))
            elif isinstance(v, (_dt.datetime, _dt.date)):
                row.append(str(v))
            else:
                sv = str(v)
                row.append(sv)
        rows.append(row)
    wb.close()
    # trailing all-empty rows/cols are not part of xlsread's range
    def empty(v):
        return isinstance(v, float) and math.isnan(v)
    while rows and all(empty(v) for v in rows[-1]):
        rows.pop()
    ncol = max((len(r) for r in rows), default=0)
    rows = [r + [float("nan")] * (ncol - len(r)) for r in rows]
    while ncol > 0 and all(empty(r[ncol - 1]) for r in rows):
        ncol -= 1
        rows = [r[:ncol] for r in rows]
    _XLS_CACHE[key] = rows
    return rows


def _is_num(v) -> bool:
    return isinstance(v, float)


def _is_num_not_nan(v) -> bool:
    return isinstance(v, float) and not math.isnan(v)


def _cell_num2str(v) -> str:
    """cellfun(@(x) num2str(x), ...) on a RAW cell."""
    if isinstance(v, str):
        return v
    return num2str(v)


def detect_info_etc(TXT: List[str]) -> List[bool]:
    out = []
    for t in TXT:
        t2 = t.replace(" ", "").lower()
        out.append(t2 == "info" or re.match(r"^info[_-]?\d", t2) is not None)
    return out


def AQP_rm_replicate_seq_sam(s: str) -> str:
    """Strip '_12032019T152102.sam' time stamps and '-#.sam' / '_#.sam' replicate
    counters so replicate scans of one sample share one ID."""
    m_TS = re.findall(r"[_]\d{8}[T]\d{6}[.]sam", s)
    if len(m_TS) == 1 and s.find(m_TS[0]) == len(s) - len(m_TS[0]):
        s = s.replace(m_TS[0], ".sam")
    m = re.findall(r"[-_]\d+[.]sam", s)
    if not m:
        return s
    q = m[-1]
    last = s.rfind(q)
    end = last + len(q)
    if end != len(s):
        return s
    if len(m) == 1:
        return s.replace(q, ".sam")
    return s[:last] + ".sam"


def read_ID_column(path: str) -> List[str]:
    """LoadXlsx4AQP with XRS_ID_only_yes: the ID column of the XRS sheet."""
    RAW, TXT, NUMh = _xls_trim(read_xlsx_raw(path))
    if "ID" not in TXT:
        raise AQPError(f"no 'ID' column in {os.path.basename(path)}")
    k = TXT.index("ID")
    return [_cell_num2str(r[k]) for r in RAW[1:]]


def _xls_trim(RAW0):
    """The common front part of LoadXlsx4AQP / clistfilename2AT_AQP: drop all-NaN
    rows/cols, find the wavelength block from the header, cut everything to the
    right of it.  Returns RAW, TXT (header texts), loc_wvl (columns)."""
    RAW = [list(r) for r in RAW0]
    if not RAW:
        raise AQPError("empty sheet")
    # trim all-NaN rows / cols (numeric NaN everywhere)
    keep_r = [i for i, r in enumerate(RAW) if not all(_is_num(v) and math.isnan(v) for v in r)]
    RAW = [RAW[i] for i in keep_r]
    ncol = len(RAW[0])
    keep_c = [j for j in range(ncol) if not all(_is_num(r[j]) and math.isnan(r[j]) for r in RAW)]
    RAW = [[r[j] for j in keep_c] for r in RAW]
    cheading = RAW[0]
    loc_wvl = [j for j, v in enumerate(cheading) if _is_num_not_nan(v)]
    if not loc_wvl:
        raise AQPError("no wavelength columns found in the header row")
    RAW = [r[:loc_wvl[-1] + 1] for r in RAW]
    TXT = [v for v in RAW[0] if isinstance(v, str) and v.strip() != ""]
    return RAW, TXT, loc_wvl


@dataclass
class LoadedSheet:
    SAT: SpecSet                       # what was saved as Atrainpketc_saConc_(..)
    XSmst: Optional[SpecSet] = None    # CS rows with Ref = NaN that are XRS IDs
    kind: str = ""                     # 'CS' | 'XRS-woRef' | 'Val'
    analytes: List[str] = field(default_factory=list)
    path: str = ""


def load_sheet(path: str, keyword4TP: str, AnaName: Optional[str], ID_XRS: Sequence[str],
               Spectra_Avg_Method: str, pick_analyte=None) -> LoadedSheet:
    """Port of LoadXlsx4AQP + prep_IDRC_shootout_MLtool + clistfilename2AT_AQP."""
    RAW, TXT, loc_wvl = _xls_trim(read_xlsx_raw(path))
    nvar = len(RAW[0]) - len(TXT)
    if nvar <= 0:
        raise AQPError(f"{os.path.basename(path)}: cannot find the spectra block")
    wvl = np.array([RAW[0][j] for j in range(len(RAW[0]) - nvar, len(RAW[0]))], dtype=float)
    body = RAW[1:]
    RS = np.array([[v if _is_num(v) else np.nan for v in r[len(r) - nvar:]] for r in body], dtype=float)
    if np.isnan(RS).any():
        raise AQPError(f"{int(np.isnan(RS).sum())} spectra readings were NaN inside RawSpectra "
                       f"({os.path.basename(path)})")
    nsamp = RS.shape[0]
    analytes: List[str] = []
    if keyword4TP == "CS":
        loc_info = detect_info_etc(TXT)
        ncol4info = sum(loc_info)
        analytes = TXT[1 + ncol4info:]
        if len(analytes) > 1:
            if AnaName in analytes:
                pass
            elif pick_analyte is not None:
                AnaName = pick_analyte(analytes)
            else:
                AnaName = analytes[0]
            if AnaName not in analytes:
                raise AQPError("no analyte picked")
        elif len(analytes) == 1:
            AnaName = analytes[0]
        else:
            raise AQPError(f"{os.path.basename(path)}: no analyte column found")
    if AnaName not in TXT:
        if keyword4TP.lower() == "val":
            raise AQPError(f"analyte '{AnaName}' not found in {os.path.basename(path)}")
        loc_Ana = None
    else:
        loc_Ana = TXT.index(AnaName)
    # ---- reference values (same control flow as the MATLAB try/catch) ----
    loc_NaN: Optional[List[int]] = []
    Conc_T_all = None
    if loc_Ana is None:
        Conc_T = np.full(nsamp, np.nan)
        loc_NaN = list(range(nsamp))
        loc_numericRef = []
    else:
        cells = [r[loc_Ana] for r in body[:nsamp]]
        if all(_is_num(v) for v in cells):                        # cell2mat succeeded
            Conc_num = np.array(cells, dtype=float)
            branch_ok = True
            loc_numericRef = list(range(nsamp))
            if ID_XRS:
                loc_ID = TXT.index("ID") if "ID" in TXT else None
                lab_all = [_cell_num2str(r[loc_ID]) for r in body] if loc_ID is not None else []
                idx = [i for i, a in enumerate(lab_all) if a in set(ID_XRS)]
                if [lab_all[i] for i in idx] == list(ID_XRS):
                    loc_NaN = idx
                    loc_numericRef = [i for i in range(nsamp) if i not in set(idx)]
                    Conc_T_all = Conc_num.copy()
                    Conc_T = np.zeros(nsamp)
                    Conc_T[loc_numericRef] = Conc_num[loc_numericRef]
                    Conc_T[idx] = np.nan
                    if np.any(Conc_T == 0):
                        raise AQPError("some values in Conc_T not updated ?")
                else:
                    branch_ok = False          # MATLAB: error -> catch branch below
            else:
                loc_NaN = []
                Conc_T = Conc_num
            if not branch_ok:
                loc_NaN = []                   # catch branch: no char cells -> loc_NaN = []
                Conc_T = Conc_num
                loc_numericRef = list(range(nsamp))
                Conc_T_all = None
        else:
            idx_char = [isinstance(v, str) for v in cells]
            vals_char = []
            for v in cells:
                if isinstance(v, str):
                    try:
                        vals_char.append(float(v.strip()) if v.strip().lower() != "nan" else np.nan)
                    except ValueError:
                        vals_char.append(None)
            if any(v is None for v in vals_char) or not all(np.isnan(v) for v in vals_char):
                raise AQPError(f"something other than NaN found in the '{AnaName}' column "
                               f"of {os.path.basename(path)}")
            loc_NaN = [i for i, c in enumerate(idx_char) if c]
            loc_numericRef = [i for i, c in enumerate(idx_char) if not c]
            Conc_T = np.array([np.nan if isinstance(v, str) else v for v in cells], dtype=float)
    # ---- sample IDs ----
    if "ID" in TXT:
        loc_lab = TXT.index("ID")
    else:
        norm = [t.lower().replace("_", "").replace(" ", "").replace("-", "") for t in TXT]
        if "samplename" not in norm:
            raise AQPError('1st Column should be either "ID" or "SampleName" (case insensitive '
                           'and can have "_" or "-" or space)')
        loc_lab = norm.index("samplename")
    AclabelT = [_cell_num2str(r[loc_lab]) for r in body]
    if len(AclabelT) != len(Conc_T):
        raise AQPError("length of AclabelT not matched with Conc")
    # ---- pretreatment at load time is 'none' (all-zero columns are dropped) ----
    allzero = np.all(RS == 0, axis=0)
    if allzero.any():
        Atrainpk = RS[:, ~allzero]
        RS = np.full(RS.shape, np.nan)
    else:
        Atrainpk = RS.copy()
    AclabelT = [AQP_rm_replicate_seq_sam(a) for a in AclabelT]
    conc_used = Conc_T_all if (Conc_T_all is not None and not np.isnan(Conc_T_all).any()) else Conc_T
    SAT = SpecSet(Atrainpk, RS, wvl, AclabelT, conc_used, AnaName or "")
    # ---- split / average ----
    xsmst = None
    if loc_NaN is not None and len(loc_NaN) == 0:
        SATnew = SAT
        if Spectra_Avg_Method != "Spectra_Avg_All" and keyword4TP.lower() == "cs":
            SATnew = apply_spectra_avg(SAT, Spectra_Avg_Method, AnaName)
        kind = keyword4TP
    elif loc_NaN is not None and len(loc_NaN) == len(Conc_T):
        SATnew = SAT
        kind = "XStrg_wRefNaN" if keyword4TP == "XRS" else "UnknownDS_wRefNaN"
    else:
        if Conc_T_all is not None and not np.isnan(Conc_T_all).any():
            SATnew = SAT
        else:
            SATnew = extract_rows_quirk(SAT, loc_numericRef)
        if Spectra_Avg_Method != "Spectra_Avg_All" and keyword4TP.lower() == "cs":
            SATnew = apply_spectra_avg(SATnew, Spectra_Avg_Method, AnaName)
        kind = keyword4TP
        ids = set(ID_XRS or [])
        loc_XS1 = [i for i, a in enumerate(SAT.AclabelT) if a in ids]
        loc_XS = sorted(set(loc_NaN) & set(loc_XS1))
        if loc_XS:
            if loc_XS1 != loc_XS:
                raise AQPError("locations of XSmst not a subset of all rows with NaN ? this code is "
                               "not ready to handle XSmst with ref values yet")
            xsmst = extract_rows_quirk(SAT, loc_XS)
    return LoadedSheet(SATnew, xsmst, kind, analytes, path)


def extract_rows_quirk(L: SpecSet, keep: Sequence[int]) -> SpecSet:
    """ATsaConc_extract_selective_samples: rows not in keep are deleted.  The
    MATLAB code also tries clistclslabel(loc_OLs)=[] inside a try; that
    deletion only happens when every row index is <= numel(clistclslabel)."""
    keep = list(keep)
    drop = sorted(set(range(L.n)) - set(keep))
    new = L.subset([i for i in range(L.n) if i not in set(drop)])
    if drop and max(drop) < len(L.clistclslabel):
        new.clistclslabel = [c for i, c in enumerate(L.clistclslabel) if i not in set(drop)]
    else:
        new.clistclslabel = list(L.clistclslabel)
    return new


def _mode_smallest(v: np.ndarray) -> float:
    vals, cnt = np.unique(v[~np.isnan(v)], return_counts=True)
    if vals.size == 0:
        return float("nan")
    return float(vals[np.argmax(cnt)])


def apply_spectra_avg(SAT: SpecSet, method: str, AnaName: str) -> SpecSet:
    """AQP_Apply_Spectra_Avg: one mean spectrum per sample ID ('<ID>_Mean')."""
    if method in ("Spectra_Avg_All",):
        return SAT
    if method not in ("Spectra_Avg_Mean", "Spectra_Avg_T-Mean_P-All", "one_scan_each_cls", "Mean_PDS"):
        raise AQPError("Spectra_Avg_Method Not supported")
    labels = SAT.AclabelT
    if method == "one_scan_each_cls":
        labels = ["cls-" + num2str(c) for c in SAT.AclassinfoT]
    Q = unique_appear_order(labels)
    AT, RS, lab, cinfo, conc = [], [], [], [], []
    for q in Q:
        loc = [i for i, a in enumerate(labels) if a == q]
        AT.append(SAT.Atrainpk[loc].mean(axis=0))
        RS.append(SAT.RawSpectra[loc].mean(axis=0))
        lab.append(q + "_Mean")
        cinfo.append(_mode_smallest(np.asarray(SAT.AclassinfoT, dtype=float)[loc]))
        conc.append(float(np.mean(SAT.conc[loc])))
    new = SpecSet(np.array(AT), np.array(RS), SAT.wvl.copy(), lab, np.array(conc), AnaName,
                  list(SAT.clistclslabel), np.array(cinfo), SAT.AInfo_1, None, None, dict(SAT.extra))
    if method == "Spectra_Avg_T-Mean_P-All":
        new.SAT_orig_Avg_All = SAT.copy()
    return new


# =============================================================================
#   DATASET  (XLSX2MAT_AQP)  and  MATCH GRIDS
# =============================================================================
def find_prefixed_xlsx(folder: str, prefix: str) -> Optional[str]:
    """get_OnlyOne_XLSX_wPrefix: exactly one *.xlsx whose name starts with prefix
    (case-insensitive, as on Windows).  None when there is none."""
    hits = []
    for f in sorted(os.listdir(folder), key=str.upper):
        if f.lower().endswith(".xlsx") and f.lower().startswith(prefix.lower()) and not f.startswith("~$"):
            hits.append(os.path.join(folder, f))
    if not hits:
        return None
    if len(hits) > 1:
        raise AQPError(f"non-unique {prefix}*.xlsx in {folder}")
    return hits[0]


def interp_rows(x: np.ndarray, Y: np.ndarray, xi: np.ndarray) -> np.ndarray:
    """interp1_CH (linear; NaN outside the range, like interp1)."""
    Y = np.atleast_2d(Y)
    out = np.empty((Y.shape[0], xi.size))
    for r in range(Y.shape[0]):
        out[r] = np.interp(xi, x, Y[r], left=np.nan, right=np.nan)
    return out


def ismembertol(a, b, tol=1e-3):
    a = np.ravel(a)
    b = np.ravel(b)
    sc = max(np.max(np.abs(a)) if a.size else 0, np.max(np.abs(b)) if b.size else 0)
    lia = np.zeros(a.size, dtype=bool)
    locb = np.zeros(a.size, dtype=int) - 1
    for i, v in enumerate(a):
        k = np.nonzero(np.abs(b - v) <= tol * sc)[0]
        if k.size:
            lia[i] = True
            locb[i] = k[0]
    return lia, locb


def is_MN_wvl(w) -> bool:
    lia, _ = ismembertol(w, MN_WVL, 1e-3)
    return bool(lia.size == MN_WVL.size and lia.all())


def to_common_grid(L: SpecSet, wvl_common: np.ndarray) -> SpecSet:
    """ATsaConc2wvlcommon with 'pp-After-MatchGrids' and pp none."""
    keep = [i for i, w in enumerate(L.wvl) if np.any(wvl_common == w)]
    new = L.copy()
    new.wvl = L.wvl[keep].copy()
    new.RawSpectra = L.RawSpectra[:, keep].copy()
    if not np.all(np.isnan(new.RawSpectra)):
        new.Atrainpk = new.RawSpectra.copy()
    return new


@dataclass
class Dataset:
    """Everything XLSX2MAT_AQP leaves behind (one dataset folder)."""
    path_XLSX: str
    AnaName: str
    CS: SpecSet                        # CS after match grids (T of every TP)
    CS_orig: SpecSet                   # CS before match grids (StepByStep)
    XSmst: Optional[SpecSet]           # master transfer spectra after match grids
    XSmst_orig: Optional[SpecSet]
    XRS: Optional[SpecSet]             # target transfer spectra (after match grids)
    Val: Optional[SpecSet]             # validation set (after match grids)
    cCabXfer_forced: Optional[str] = None
    files: Dict[str, Optional[str]] = field(default_factory=dict)
    analytes: List[str] = field(default_factory=list)


def load_dataset(path_XLSX: str, AnaName: Optional[str], Spectra_Avg_Method: str,
                 AQP_class: str = "pro", pick_analyte=None) -> Dataset:
    """XLSX2MAT_AQP + Match_wvl_HiRes2LoRes_and_Diff_Ranges_MLtool."""
    if not path_XLSX or not os.path.isdir(path_XLSX):
        raise AQPError("pls load a dataset folder first (Load_DS)")
    f_XRS = find_prefixed_xlsx(path_XLSX, "XRS_")
    f_CS = find_prefixed_xlsx(path_XLSX, "CS_")
    f_Val = find_prefixed_xlsx(path_XLSX, "Val_")
    if f_CS is None:
        raise AQPError(f"no CS_*.xlsx found in {path_XLSX}")
    forced = None
    if f_XRS:
        ID_XRS = read_ID_column(f_XRS)
    else:
        if AQP_class == "pro":
            forced = "woCabXfer"
            note("Running without calibration transfer (no XRS_*.xlsx in the dataset folder)")
        else:
            raise AQPError("can not find XRS set and stop running AQP")
        ID_XRS = []
    cs = load_sheet(f_CS, "CS", AnaName, ID_XRS, Spectra_Avg_Method, pick_analyte)
    AnaName = cs.SAT.clsname
    xsmst = cs.XSmst
    if xsmst is None and f_XRS:
        warn("XSmst not able to be extracted ??? Assume you are NOT going to run CabXfer")
    xrs = None
    if xsmst is not None and f_XRS:
        xrs = load_sheet(f_XRS, "XRS", AnaName, ID_XRS, Spectra_Avg_Method).SAT
    val = None
    if f_Val:
        kw = os.path.basename(f_Val).split("_")[0]
        val = load_sheet(f_Val, kw, AnaName, ID_XRS, Spectra_Avg_Method).SAT
    else:
        msgbox("NO Validation set exist  !!!")
    # ---------------- match grids ----------------
    Lhr = cs.SAT
    CS_m, XRS_m, Val_m, XS_m = Lhr, xrs, val, xsmst
    do_match = True
    if xrs is None and val is not None:
        if is_MN_wvl(Lhr.wvl) and is_MN_wvl(val.wvl):
            do_match = False
        Llr = val
    elif xrs is not None:
        Llr = xrs
    else:
        do_match = False
        Llr = None
    if do_match:
        wvl_HR, wvl_LR = Lhr.wvl, Llr.wvl
        if wvl_HR.size == wvl_LR.size and np.all(wvl_HR == wvl_LR):
            wvl_common = wvl_LR.copy()
            same = True
        else:
            same = False
            lo = max(wvl_LR.min(), wvl_HR.min())
            hi = min(wvl_LR.max(), wvl_HR.max())
            i0 = int(np.nonzero(wvl_LR >= lo)[0][0])
            i1 = int(np.nonzero(wvl_LR <= hi)[0][-1])
            wvl_common = wvl_LR[i0:i1 + 1].copy()
        Llr_m = to_common_grid(Llr, wvl_common)
        Lh = Lhr.copy()
        Lh.RawSpectra = interp_rows(Lhr.wvl, Lhr.RawSpectra, wvl_common)
        Lh.Atrainpk = Lhr.Atrainpk.copy() if same else Lh.RawSpectra.copy()
        Lh.wvl = wvl_common.copy()
        CS_m = to_common_grid(Lh, wvl_common)
        CS_m.extra = {"wvl": "", "wvl_AT": wvl_common.copy()}   # (try Lhr.wvl='' / Lhr.wvl_AT=..)
        if xsmst is not None:
            Lx = xsmst.copy()
            Lx.RawSpectra = interp_rows(xsmst.wvl, xsmst.RawSpectra, wvl_common)
            Lx.wvl = wvl_common.copy()
            XS_m = to_common_grid(Lx, wvl_common)
        if xrs is not None:
            XRS_m = Llr_m
        if val is not None:
            Val_m = to_common_grid(val, wvl_common) if Llr is not val else Llr_m
    ds = Dataset(path_XLSX, AnaName, CS_m, cs.SAT, XS_m, xsmst, XRS_m, Val_m, forced,
                 {"CS": f_CS, "XRS": f_XRS, "Val": f_Val}, cs.analytes)
    if ds.Val is not None and ds.Val.Atrainpk.shape[1] != ds.CS.Atrainpk.shape[1]:
        raise AQPError("the VAL grid does not match the common grid of CS/XRS")
    return ds


# =============================================================================
#   CALIBRATION TRANSFER   (CabXfer_Siesler48_MLtool + run_each_CabXfer_AQP)
# =============================================================================
CABXFER_NEEDS_PLS_TOOLBOX = ("GLSw", "MDCextr_GLSw", "GLSw_default", "LS-GLSw", "MDC-GLSw",
                             "LS-OSC", "OSC", "PDS", "LS-PDS", "STDgenize")
CABXFER_SUPPORTED = ("woCabXfer", "MDC", "LS", "PDS99", "STDgenize99", "GLSstd2", "MDC_GLSstd3",
                     "LS_GLSstd3")


def split_scheme(s: str) -> Tuple[str, str]:
    """'PDS99[w5]' -> ('PDS99', 'w5')"""
    add = find_keyword_between_markers(s, "[", "]") or ""
    base = s.replace("[" + add + "]", "") if add else s
    return base, add


@dataclass
class TransferStandards:
    CT: Optional[np.ndarray]           # master transfer spectra (median per sample)
    CP: Optional[np.ndarray]           # target transfer spectra (median per sample)
    labels: List[str]
    snXS: str
    forced_woCabXfer: bool = False


def transfer_standards(T: SpecSet, P: SpecSet, XSmst: Optional[SpecSet]) -> TransferStandards:
    """The 'KSall' selection of CabXfer_Siesler48_MLtool (the way it actually runs in
    AQP): one median spectrum per physical transfer sample on each unit."""
    Plabs = list(P.AclabelT)
    Pset = set(Plabs)
    locXT = [i for i, a in enumerate(T.AclabelT) if a in Pset]
    if not locXT:
        # no reference values for the transfer set: XSmst came out of the CS sheet
        if XSmst is None:
            return TransferStandards(None, None, [], "", True)
        qT = unique_sorted(XSmst.AclabelT)
        qP = unique_sorted(Plabs)
        if len(qT) != len(qP):
            raise AQPError("mismatch in lengths between qLabelXT qLabelXP ")
        CT, CP = [], []
        for a, b in zip(qT, qP):
            lt = [i for i, x in enumerate(XSmst.AclabelT) if x == a]
            lp = [i for i, x in enumerate(Plabs) if x == b]
            CT.append(mat_median_rows(XSmst.Atrainpk[lt]))
            CP.append(mat_median_rows(P.Atrainpk[lp]))
        return TransferStandards(np.array(CT), np.array(CP), qT, f"_nsampXS{len(qT)}")
    # transfer samples carry reference values and are part of CS ("KSn" branch)
    qT = unique_sorted([T.AclabelT[i] for i in locXT])
    qP = unique_sorted(Plabs)
    if len(qT) != len(qP):
        raise AQPError("mismatch in lengths between qLabelXT qLabelXP ")
    nKS = len(set(unique_sorted(Plabs)) & set(unique_sorted(T.AclabelT)))
    CT, CP = [], []
    for a, b in zip(qT, qP):
        lt = [i for i, x in enumerate(T.AclabelT) if x == a]
        lp = [i for i, x in enumerate(Plabs) if x == b]
        if nKS == 1:
            if len(lt) == len(lp):
                CT.extend(T.Atrainpk[lt]); CP.extend(P.Atrainpk[lp])
            else:
                raise AQPError("KS1 with unequal replicate counts is not handled")
        else:
            CT.append(mat_median_rows(T.Atrainpk[lt]))
            CP.append(mat_median_rows(P.Atrainpk[lp]))
    return TransferStandards(np.array(CT), np.array(CP), qT, f"_nsampXS{len(qT)}")


@dataclass
class XferModel:
    scheme: str
    apply_T: Any          # function(T_prep) -> T_cab
    apply_P: Any          # function(P_prep) -> P_cab
    info: Dict[str, Any] = field(default_factory=dict)


def build_transfer(scheme_orig: str, std: TransferStandards, DataFlow: str = "Xfer-PP") -> XferModel:
    base, add = split_scheme(scheme_orig)
    CT, CP = std.CT, std.CP
    ident = lambda X: np.array(X, dtype=float, copy=True)
    if base == "woCabXfer":
        return XferModel(scheme_orig, ident, ident)
    if CT is None or CP is None or len(CT) == 0:
        raise AQPError("no transfer standards available for " + scheme_orig)
    if base == "MDC":
        mCT, mCP = CT.mean(axis=0), CP.mean(axis=0)
        return XferModel(scheme_orig, lambda X: (np.asarray(X) - mCT) + mCP, ident,
                         {"mean_CT": mCT, "mean_CP": mCP})
    if base == "LS":
        mCT, mCP = CT.mean(axis=0), CP.mean(axis=0)
        return XferModel(scheme_orig, lambda X: np.asarray(X) - mCT, lambda X: np.asarray(X) - mCP,
                         {"mean_CT": mCT, "mean_CP": mCP})
    if base in ("STDgenize99", "PDS99"):
        if add:
            win = find_keynumber_after(add, "win")
            if math.isnan(win):
                win = find_keynumber_after(add, "w")
            if math.isnan(win):
                raise AQPError('something wrong with "win" width')
            win = int(win)
        else:
            win = 5
        stdmat, stdvect = stdgen99(CP, CT, win)
        return XferModel(scheme_orig, lambda X: stdize99(X, stdmat, stdvect), ident,
                         {"stdmat": stdmat, "stdvect": stdvect, "win": win})
    if base in ("GLSstd2", "MDC_GLSstd3", "LS_GLSstd3"):
        alpha = 1e-6
        if add:
            sa = find_keyword_between_markers(add, "alpha", "")
            try:
                alpha = float((sa or "").replace("p", "."))
            except ValueError:
                raise AQPError('something wrong with "alpha" parameter')
        if base == "GLSstd2":
            W, meandif = GLSstd(CP, CT, alpha, 2)
            return XferModel(scheme_orig, lambda X: (np.asarray(X) + meandif) @ W,
                             lambda X: np.asarray(X) @ W, {"BadWgtMatrX": W, "meandif": meandif})
        if base == "MDC_GLSstd3":
            CT_MDC = CT - CT.mean(axis=0) + CP.mean(axis=0)
            W, meandif = GLSstd(CP, CT_MDC, alpha, 3)
            mCT, mCP = CT.mean(axis=0), CP.mean(axis=0)
            return XferModel(scheme_orig, lambda X: (np.asarray(X) - mCT + mCP) @ W,
                             lambda X: np.asarray(X) @ W, {"BadWgtMatrX": W, "meandif": meandif})
        if DataFlow != "PP-Xfer":
            raise AQPError('DataFlow should be "PP-Xfer" to this Xfer_scheme to make sense')
        W, meandif = GLSstd(CP - CP.mean(axis=0), CT - CT.mean(axis=0), alpha, 3)
        mCT, mCP = CT.mean(axis=0), CP.mean(axis=0)
        return XferModel(scheme_orig, lambda X: (np.asarray(X) - mCT) @ W,
                         lambda X: (np.asarray(X) - mCP) @ W, {"BadWgtMatrX": W})
    if base in CABXFER_NEEDS_PLS_TOOLBOX:
        raise AQPError(f"CabXfer scheme '{scheme_orig}' needs Eigenvector's PLS_Toolbox "
                       f"(glsw/stdgen/osccalc), which is not part of AQP_gui_merged.m")
    raise AQPError("CabXfer scheme not supported: " + scheme_orig)


# =============================================================================
#   PLS ENGINE   (PLS_Tcv, PLS_indv, PLS_inside_PLS_predict_ONLY_MLtool)
# =============================================================================
def tcv_folds(cSampleName: Sequence[str], scheme: str = "sqrtNSfolds-Conc"):
    """Fold assignment of PLS_Tcv.  Returns (folds as list of P-row lists, nFolds,
    QSample_appear_order, idx_table rows)."""
    QSample = unique_sorted(cSampleName)
    order = unique_appear_order(cSampleName)
    nSample = len(QSample)
    if scheme in ("10folds-Conc", "5folds-Conc", "sqrtNSfolds-Conc"):
        if scheme == "sqrtNSfolds-Conc":
            nFolds = int(math.floor(math.sqrt(nSample)))
        else:
            nFolds = min(int(find_keyword_between_markers(scheme, "", "fold")), nSample)
        ncol = int(math.ceil(nSample / nFolds))
        rows = []
        for i in range(nFolds):
            rows.append([i + k * nFolds for k in range(ncol) if i + k * nFolds < nSample])
        return rows, nFolds, order, QSample
    if scheme == "Leave-OneConc-Out":
        srt = sortnat(QSample)
        return [[order.index(q)] for q in srt], nSample, order, QSample
    raise AQPError("Tcv scheme NOT supported")


def PLS_Tcv(X, Y, names, PLSfactor: int, scheme: str = "sqrtNSfolds-Conc"):
    X = np.asarray(X, dtype=float)
    Y = np.ravel(np.asarray(Y, dtype=float))
    if scheme in ("RMSEC", "Self-Pred"):
        beta = plsregress(X, Y, PLSfactor)[4]
        Yest = pls_predict(X, beta)
        return {"Yest_all": Yest, "RMSE": rms_error_woNaN_N(Yest, Y),
                "CV": 100 * rms_error_woNaN(Yest, Y) / np.mean(Y), "PLSfactor": PLSfactor,
                "Tcv_scheme": scheme, "nFolds": float("nan")}
    rows, nFolds, order, QSample = tcv_folds(names, scheme)
    Yest = np.full(Y.size, np.nan)
    fold_of = np.full(Y.size, np.nan)
    names = list(names)
    for iQS, r in enumerate(rows):
        held = set(order[k] for k in r) if scheme != "Leave-OneConc-Out" else {sortnat(QSample)[iQS]}
        locP = [i for i, a in enumerate(names) if a in held]
        locT = [i for i in range(Y.size) if i not in set(locP)]
        if len(set(names[i] for i in locP)) + len(set(names[i] for i in locT)) != len(QSample):
            raise AQPError("not all spectra parsed into training vs prediction set in cross validation")
        beta = plsregress(X[locT], Y[locT], PLSfactor)[4]
        Yest[locP] = pls_predict(X[locP], beta)
        fold_of[locP] = iQS + 1
    return {"Yest_all": Yest, "RMSE": rms_error_woNaN_N(Yest, Y),
            "CV": 100 * rms_error_woNaN(Yest, Y) / np.mean(Y), "PLSfactor": PLSfactor,
            "nFolds": nFolds, "Tcv_scheme": scheme, "fold": fold_of}


def PLS_Tcv_P_Avg_All(X, Y, names, PLSfactor, XP, YP, namesP, scheme="sqrtNSfolds-Conc"):
    """PLS_Tcv_P_Avg_All: folds built on the averaged T set, every scan of the
    held-out samples predicted (Spectra_Avg_T-Mean_P-All)."""
    X = np.asarray(X, dtype=float)
    Y = np.ravel(np.asarray(Y, dtype=float))
    XP = np.asarray(XP, dtype=float)
    YP = np.ravel(np.asarray(YP, dtype=float))
    namesP = [a + "_Mean" for a in namesP]
    rows, nFolds, order, QSample = tcv_folds(names, scheme)
    Yest = np.full(Y.size, np.nan)          # grows like MATLAB when P has more rows
    for r in rows:
        held = set(order[k] for k in r)
        locP = [i for i, a in enumerate(namesP) if a in held]
        locT = [i for i in range(Y.size) if i not in set(r)]
        beta = plsregress(X[locT], Y[locT], PLSfactor)[4]
        if locP:
            if max(locP) >= Yest.size:
                Yest = np.concatenate([Yest, np.full(max(locP) + 1 - Yest.size, np.nan)])
            Yest[locP] = pls_predict(XP[locP], beta)
    if Yest.size != YP.size:
        raise AQPError("mismatch in length between Xhat and Xmeasure (Spectra_Avg_T-Mean_P-All "
                       "cross-validates the mean spectra against every scan of the same CS, so it "
                       "only runs when the dataset folder holds just the CS)")
    return {"Yest_all": Yest, "RMSE": rms_error_woNaN_N(Yest, YP),
            "CV": 100 * rms_error_woNaN(Yest, YP) / np.mean(YP), "PLSfactor": PLSfactor,
            "nFolds": nFolds, "Tcv_scheme": scheme}


def PLS_indv(XT, XP, YT, YP, PlsfactorScan):
    out = []
    for k in PlsfactorScan:
        beta = plsregress(XT, YT, int(k))[4]
        Yest = pls_predict(XP, beta)
        out.append({"PLSfactor": int(k), "RMSE": rms_error_woNaN_N(Yest, YP), "Yest_PLS": Yest,
                    "Y_iAna_P": np.ravel(YP).copy(), "beta": beta})
    return out


def pick_opm_factor(scheme: str, scan: Sequence[int], rmsecv: np.ndarray, user_pick=None):
    """CurTcvModelParaOpmScheme -> (PLSfactor_Opm, RMSE at it, knee point res_x)."""
    scan = list(scan)
    rmsecv = np.asarray(rmsecv, dtype=float)
    res_x = None
    if scheme == "Min-RMSECV":
        k = int(np.nanargmin(rmsecv)) if not np.all(np.isnan(rmsecv)) else 0
        opm = scan[k]
    elif scheme in ("KneePt-RMSECV", "KneePt+1_RMSECV", "User-Pick", "LSUX_RMSECV"):
        try:
            if max(scan) > 2:
                res_x, _ = knee_pt(rmsecv, scan)
            else:
                k = int(np.nanargmin(rmsecv))
                res_x = k + 1          # (sic) the MATLAB code stores the location here
        except Exception as e:      # [BUGFIX Aug 26, 2026] fall back to Min-RMSECV
            warn(f"knee_pt failed ({e}) -- falling back to Min-RMSECV.")
            res_x = scan[int(np.nanargmin(rmsecv))]
        if scheme == "KneePt+1_RMSECV":
            opm = res_x + 1 if max(scan) > res_x else res_x
        elif scheme == "User-Pick":
            res_x = float("nan")
            opm = user_pick
        elif scheme == "LSUX_RMSECV":
            raise AQPError("LSUX_RMSECV (PLS_wTcv_StandAlone_AQPpu) is not part of the GUI port")
        else:
            opm = res_x
        opm = int(opm)
    else:
        raise AQPError("unknown TcvModelParaOpmScheme: " + str(scheme))
    m = [i for i, k in enumerate(scan) if k == opm]
    rm = float(rmsecv[m[0]]) if m else float("nan")
    return int(opm), rm, res_x


# =============================================================================
#   FIGURES
#   GUI mode: every figure opens in its own window (like MATLAB figure windows).
#   File mode (CLI): figures are saved as PNG files and closed.
# =============================================================================
class FigMgr:
    def __init__(self, mode: str = "file", outdir: Optional[str] = None, on_new=None):
        self.mode = mode                  # 'gui' | 'file' | 'none'
        self.outdir = outdir
        self.on_new = on_new              # GUI hook: on_new(fig, name)
        self.figs: List[Any] = []         # figures of the current iteration
        self.kept: List[Any] = []         # figures kept from earlier iterations
        self.counter = 0
        self.saved: List[str] = []

    def new(self, name: str, size=(9.4, 5.3)):
        if self.mode == "none":
            return None
        from matplotlib.figure import Figure
        fig = Figure(figsize=size, dpi=100)
        self.counter += 1
        fig._aqp_name = name
        fig._aqp_num = self.counter
        self.figs.append(fig)
        return fig

    def show(self, fig):
        """Called once a figure is complete."""
        if fig is None:
            return
        try:
            fig.tight_layout()
        except Exception:
            pass
        if self.mode == "gui" and self.on_new is not None:
            self.on_new(fig, fig._aqp_name)
        elif self.mode == "file" and self.outdir:
            os.makedirs(self.outdir, exist_ok=True)
            from matplotlib.backends.backend_agg import FigureCanvasAgg
            FigureCanvasAgg(fig)
            fn = os.path.join(self.outdir, f"Fig{fig._aqp_num:03d}_{_safe_name(fig._aqp_name)}.png")
            fig.savefig(fn)
            self.saved.append(fn)

    def close_fig(self, fig):
        """close one figure (GUI: destroy its window)"""
        try:
            w = getattr(fig, "_aqp_window", None)
            if w is not None:
                w.destroy()
        except Exception:
            pass
        if fig in self.figs:
            self.figs.remove(fig)

    def save_set(self, figs, folder, prefix):
        """copy_all_open_fig_AQP: store figures as PNG (MATLAB: .fig)."""
        os.makedirs(folder, exist_ok=True)
        from matplotlib.backends.backend_agg import FigureCanvasAgg
        out = []
        for f in figs:
            if f is None:
                continue
            if f.canvas is None or not hasattr(f.canvas, "print_png"):
                FigureCanvasAgg(f)
            fn = os.path.join(folder, f"{prefix}_AQP_{f._aqp_num}.png")
            try:
                f.savefig(fn)
                out.append(fn)
            except Exception as e:
                warn(f"could not save figure {fn}: {e}")
        return out


def _safe_name(s: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.+\-\[\]{}()]+", "_", s)[:80]


# =============================================================================
#   MAT-FILE WRITERS  (deliverables readable by the MATLAB tools)
# =============================================================================
def _cellcol(lst):
    a = np.empty((len(lst), 1), dtype=object)
    for i, v in enumerate(lst):
        a[i, 0] = v
    return a


def _cellrow(lst):
    a = np.empty((1, len(lst)), dtype=object)
    for i, v in enumerate(lst):
        a[0, i] = v
    return a


def _struct_array(records: List[Dict[str, Any]], fields: Sequence[str]):
    dt = [(f, object) for f in fields]
    a = np.zeros((len(records), 1), dtype=dt)
    for i, r in enumerate(records):
        for f in fields:
            a[i, 0][f] = r.get(f, np.zeros((0, 0)))
    return a


def saConc_records(S: SpecSet, X: np.ndarray):
    recs = []
    for i in range(S.n):
        recs.append({"seqnum": float(S.seqnum[i]), "clsname": S.clsname, "Conc": float(S.conc[i]),
                     "SampleName": _cellrow([S.AclabelT[i]]), "Atrainpk": X[i:i + 1, :]})
    return _struct_array(recs, ["seqnum", "clsname", "Conc", "SampleName", "Atrainpk"])


def _ainfo(S: SpecSet):
    return S.AInfo_1 if isinstance(S.AInfo_1, str) else ""


def LAT_T_only(S: SpecSet) -> Dict[str, Any]:
    """An Atrainpketc file of one set (saved straight from the xlsx loader)."""
    d = {"saConc": saConc_records(S, S.Atrainpk), "Atrainpk": S.Atrainpk,
         "AclabelT": _cellcol(S.AclabelT), "clistclslabel": _cellrow(S.clistclslabel),
         "AclassinfoT": np.asarray(S.AclassinfoT, dtype=float).reshape(-1, 1),
         "RawSpectra": S.RawSpectra, "wvl_standardize": S.wvl.reshape(1, -1), "AInfo_1": _ainfo(S)}
    d.update(S.extra)
    return d


def LAT_merge_TP(T: SpecSet, P: SpecSet) -> Dict[str, Any]:
    """Atrainpk_merge_Apred / ssds TP_pair: T's own fields + the P set."""
    d = LAT_T_only(T)
    d.pop("saConc")
    d.update({"Apred": P.Atrainpk, "AclassinfoP": np.asarray(P.AclassinfoT, float).reshape(-1, 1),
              "AclabelP": _cellcol(P.AclabelT),
              "RawSpectra": {"Tset": T.RawSpectra, "Pset": P.RawSpectra},
              "PLS": {"Tset": {"saConc": saConc_records(T, T.Atrainpk)},
                      "Pset": {"saConc": saConc_records(P, P.Atrainpk)}}})
    if list(P.clistclslabel) != list(T.clistclslabel):
        d["clistclslabel_P"] = _cellrow(P.clistclslabel)
    return d


def _classinfo_in(labels, clist):
    """cellfun(@(x) strmatch(x,clist,'exact'),labels) with the MATLAB fallbacks:
    any label missing (or repeated in clist) -> the whole vector is NaN."""
    out = []
    for a in labels:
        k = [i for i, c in enumerate(clist) if c == a]
        if len(k) != 1:
            return np.full(len(labels), np.nan)
        out.append(k[0] + 1)
    return np.array(out, dtype=float)


def LAT_cab_TP(tp: TPSet, with_orig_raw: bool = False) -> Dict[str, Any]:
    """SAT_cab of run_each_CabXfer_AQP (and its pretreated descendants)."""
    T, P = tp.T, tp.P
    d = {"wvl_standardize": T.wvl.reshape(1, -1), "clistclslabel": _cellrow(T.clistclslabel),
         "Atrainpk": T.Atrainpk, "Apred": P.Atrainpk,
         "RawSpectra": {"Tset": T.RawSpectra, "Pset": P.RawSpectra},
         "AclabelT": _cellcol(T.AclabelT), "AclabelP": _cellcol(P.AclabelT),
         "AclassinfoT": _classinfo_in(T.AclabelT, T.clistclslabel).reshape(-1, 1),
         "AclassinfoP": _classinfo_in(P.AclabelT, T.clistclslabel).reshape(-1, 1),
         "PLS": {"Tset": {"saConc": saConc_records(T, T.Atrainpk)},
                 "Pset": {"saConc": saConc_records(P, P.Atrainpk)}},
         "saCTCP": np.zeros((0, 0))}
    if with_orig_raw and tp.RawSpectra_Tset_orig_woCabXfer is not None:
        d["RawSpectra_Tset_orig_woCabXfer"] = tp.RawSpectra_Tset_orig_woCabXfer
    return d


def LAT_cab_T_only(T: SpecSet) -> Dict[str, Any]:
    """rm_Pset of a SAT_cab: the T part only."""
    return {"wvl_standardize": T.wvl.reshape(1, -1), "clistclslabel": _cellrow(T.clistclslabel),
            "Atrainpk": T.Atrainpk, "RawSpectra": T.RawSpectra, "AclabelT": _cellcol(T.AclabelT),
            "AclassinfoT": _classinfo_in(T.AclabelT, T.clistclslabel).reshape(-1, 1),
            "saCTCP": np.zeros((0, 0)), "saConc": saConc_records(T, T.Atrainpk)}


def LAT_TP(tp: TPSet) -> Dict[str, Any]:
    if tp.P is None:
        return LAT_T_only(tp.T)
    return LAT_merge_TP(tp.T, tp.P)


def LAT_TP_AT2RS(tp: TPSet) -> Dict[str, Any]:
    """A merged TP that went through AT2RS + PP (no transfer step ran: woXRS)."""
    d = LAT_merge_TP(tp.T, tp.P)
    if tp.RawSpectra_Tset_orig_woCabXfer is not None:
        d["RawSpectra_Tset_orig_woCabXfer"] = tp.RawSpectra_Tset_orig_woCabXfer
    return d


def savemat(path: str, d: Dict[str, Any]):
    import scipy.io as sio
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    sio.savemat(path, d, do_compression=True, long_field_names=True, oned_as="column")


def atfile_name(corename: str, tp_or_set, kind="TP") -> str:
    """saveAT_fname: Atrainpketc_<corename>_nvar.._ncls.._nsampT.._nsampP...mat"""
    if isinstance(tp_or_set, TPSet):
        T, P = tp_or_set.T, tp_or_set.P
    else:
        T, P = tp_or_set, None
    svar = f"_nvar{T.Atrainpk.shape[1]}"
    sncls = f"_ncls{len(T.clistclslabel)}"
    if P is not None:
        sn = f"_nsampT{T.n}_nsampP{P.n}"
    else:
        sn = f"_nsamp{T.n}"
    core = "_" + corename if corename else ""
    return ("Atrainpketc_" + core + svar + sncls + sn + ".mat").replace("__", "_")


# =============================================================================
#   ONE RUN  (AutoQuant_DA_pipeline for one PP1/PP2 and one CabXfer scheme)
# =============================================================================
@dataclass
class RunSettings:
    path_XLSX: str
    AnaName: Optional[str] = None
    cPP1: List[str] = field(default_factory=lambda: ["1stDerSGFL7[PO2]"])
    cPP2: List[str] = field(default_factory=lambda: ["SNV"])
    cCabXfer_scheme: List[str] = field(default_factory=lambda: ["MDC"])
    cDataFlow: List[str] = field(default_factory=lambda: ["Xfer-PP"])
    cUDMas_scheme: List[str] = field(default_factory=lambda: ["woUDM"])
    Spectra_Avg_Method: str = "Spectra_Avg_All"
    Lines_Xticks_picked: str = "Lines_PP1--Xticks_CabXfer"
    CurTcvModelParaOpmScheme: str = "KneePt+1_RMSECV"
    PLSfactor_Opm_User_Pick: float = float("nan")
    AQP_class: str = "pro"
    pathfname_PPn_xlsx: Optional[str] = None
    PlsfactorScan_default: List[int] = field(default_factory=lambda: list(INTERNAL_SETTINGS["PlsfactorScan"]))
    workdir: str = "."
    Option_Output_SAM: int = 1


class RunContext:
    """Paths and shared state of one Run AQP (what the TMP folders held)."""

    def __init__(self, st: RunSettings, figs: FigMgr, progress=None):
        self.st = st
        self.figs = figs
        self.progress = progress or (lambda s: None)
        self.root = find_last_nonTMP_path(os.path.abspath(st.workdir))
        self.tmp = os.path.join(self.root, "TMP_AQP")
        self.sbs_pre = os.path.join(self.root, "TMP_AQP_StepByStep")
        self.final_models = os.path.join(self.root, "Result4FinalModels")
        self.results4pro = os.path.join(self.root, "Results4PRO")
        self.sam_folder = None

    def save_sbs(self, name: str, d: Dict[str, Any]):
        savemat(os.path.join(self.sbs_pre, name), d)


def find_last_nonTMP_path(p: str) -> str:
    parts = os.path.normpath(p).split(os.sep)
    for i, s in enumerate(parts):
        if s.startswith("TMP"):
            return os.sep.join(parts[:i]) or os.sep
    return p


def tmp_folder_rm_mk(name: str, parent: str) -> str:
    """Create parent/name; when it exists, delete the files in it (not subfolders)."""
    p = os.path.join(parent, name)
    if os.path.isfile(p):
        os.remove(p)
    os.makedirs(p, exist_ok=True)
    for f in os.listdir(p):
        fp = os.path.join(p, f)
        if os.path.isfile(fp):
            try:
                os.remove(fp)
            except OSError as e:
                warn(f"could not delete {fp}: {e}")
    return p


def create_next_subfolder(parent: str, prefix: str, keep: int = None) -> str:
    keep = INTERNAL_SETTINGS["Ntmp_keep"] if keep is None else keep
    os.makedirs(parent, exist_ok=True)
    sub = [d for d in os.listdir(parent) if os.path.isdir(os.path.join(parent, d)) and prefix in d]
    if sub:
        if len(sub) >= keep:
            srt = sortnat(sub)
            for d in srt[:len(sub) - keep + 1]:
                shutil.rmtree(os.path.join(parent, d), ignore_errors=True)
        nums = []
        for d in sub:
            m = re.search(re.escape(prefix) + r"(\d+)", d)
            if m:
                nums.append(int(m.group(1)))
        nxt = (max(nums) + 1) if nums else 1
        return tmp_folder_rm_mk(f"{prefix}{nxt}", parent)
    return tmp_folder_rm_mk(f"{prefix}1", parent)


def write_cell_xlsx(path: str, rows: List[List[Any]]):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append([(None if (isinstance(v, float) and math.isnan(v)) else v) for v in r])
    wb.save(path)


def write_cell_csv(path: str, rows: List[List[Any]]):
    """AQP_gui_merged__cell2csv: comma separated, numbers as num2str."""
    with open(path, "w", newline="") as f:
        for r in rows:
            out = []
            for v in r:
                if isinstance(v, str):
                    out.append(v)
                elif v is None:
                    out.append("")
                else:
                    out.append(num2str(v))
            f.write(",".join(out) + "\n")


# ---- string helpers used to give files the names the MATLAB build gave them ----
def regexp_extract_mk1_mk2(s: str, mk1: str, mk2) -> str:
    p1 = "^" if not mk1 else re.escape(mk1)
    if mk2 is None or mk2 == "":
        m = re.search(p1 + "(.*)", s, re.S)
    else:
        m = re.search(p1 + "(.*?)" + re.escape(mk2), s, re.S)
    return m.group(1) if m else ""


def find_keyword_between_markers_wlistRHS(s: str, m1: str, m2list) -> str:
    if isinstance(m2list, str):
        m2list = [m2list]
    found = s
    for m2 in m2list:
        kw = regexp_extract_mk1_mk2(s, m1, None if m2.lower() == "end" else m2)
        if kw and len(kw) <= len(found):
            found = kw
    return "" if found == s else found


def remove_keyword_between_markers(s: str, m1: str, m2: str, opt=None) -> str:
    kw = find_keyword_between_markers(s, m1, m2)
    if not kw:
        return s
    if opt is None:
        return s.replace(m1 + kw + m2, "")
    if opt == "keepLHS":
        return s.replace(kw + m2, "")
    if opt == "keepRHS":
        return s.replace(m1 + kw, "")
    raise AQPError("not supported opt")


def strrep_keyword_between_markers_wlistRHS(s: str, m1: str, m2list, rep: str, opt=None) -> str:
    kw = find_keyword_between_markers_wlistRHS(s, m1, m2list)
    if not kw:
        return s
    used = [m2 for m2 in m2list if (kw + m2) in s]
    used = used[0] if used else ""
    if opt is None:
        return s.replace(m1 + kw + used, rep)
    if opt == "keepLHS":
        return s.replace(kw + used, rep)
    if opt == "keepRHS":
        return s.replace(m1 + kw, rep)
    if opt == "keepBoth":
        return s.replace(kw, rep)
    raise AQPError("not supported opt")


def remove_keyword_between_markers_wlistRHS(s: str, m1: str, m2list, opt=None) -> str:
    kw = find_keyword_between_markers_wlistRHS(s, m1, m2list)
    if not kw:
        return s
    used = [m2 for m2 in m2list if (kw + m2) in s]
    used = used[0] if used else ""
    if opt is None:
        return s.replace(m1 + kw + used, "")
    if opt == "keepLHS":
        return s.replace(kw + used, "")
    if opt == "keepRHS":
        return s.replace(m1 + kw, "")
    raise AQPError("not supported opt")


def pp_rename(fname: str, pp1: str, pp2: str, nvar_new: int) -> str:
    """The file name pretreat_preprocess_RawSpectra_pp1_pp2_AT_TP gives its output."""
    f4p = f"_pp1-{pp1}_pp2-{pp2}"
    stem = fname[:-4] if fname.endswith(".mat") else fname
    kw = find_keyword_between_markers(fname, "_pp1-", "}")
    if "{" in fname and "}" in fname and kw:
        new1 = find_keyword_between_markers(f4p, "_pp1-", "_") or ""
        new = fname.replace("_pp1-" + kw, "_pp1-" + new1)
    elif "_pp1-" not in stem and "_pp2-" not in stem:
        new = fname.replace(".mat", f4p + ".mat")
    else:
        new = remove_keyword_between_markers(fname, "pp1-", "_")
        new = remove_keyword_between_markers_wlistRHS(new, "pp2-", ["_", ".mat"], "keepRHS")
        new = new.replace("__", f4p + "_")
    new = strrep_keyword_between_markers_wlistRHS(new, "_nvar", ["_", ".mat"], str(nvar_new), "keepBoth")
    if "nvar" not in new:
        new = new.replace(".mat", f"_nvar{nvar_new}.mat")
    return new


def cabxfer_fname(scheme_orig: str, pfT_name: str, nT: int, nP: int, snXS: str,
                  Test_or_Val: str = "Test") -> str:
    """fname_cab of run_each_CabXfer_AQP."""
    stem = pfT_name[:-4] if pfT_name.endswith(".mat") else pfT_name
    spp1 = find_keyword_between_markers_wlistRHS(stem, "_pp1-", ["_", "}"])
    if not spp1:
        spp1 = find_keyword_between_markers_wlistRHS(stem, "_pp-", ["_", "}"])
    sTSM = find_keyword_between_markers(stem, "{T-", "_") or ""
    sPSM = find_keyword_between_markers(stem, "_P-", "_") or ""
    f = (f"Atrainpketc_saConc_TestCabXfer_{{{scheme_orig}_KSall__pp1-{spp1}}}_[T-{sTSM}_P-{sPSM}]"
         f"_nsampT{nT}_nsampP{nP}.mat")
    f = f.replace("__", "_").replace("_pp1-1stDerSGDiederick", "_pp-1stDer")
    if split_scheme(scheme_orig)[0] != "woCabXfer":
        f = f.replace(".mat", snXS + ".mat")
    f = f.replace("_saConc", "_saConc_TPwTrn")
    if Test_or_Val == "Val":
        f = f.replace("]", "_Val]").replace(".mat", "_wo_AT2RS.mat")
    return f


def pro_csv_name(atname: str, ext: str) -> str:
    s = atname[:-4] + ext
    s = s.replace("P-Val", "")
    k1 = s.find("_nsampP")
    if k1 >= 0:
        rest = s[k1 + len("_nsampP"):]
        k2 = rest.find(".xls") if ext.startswith(".xls") else rest.find(".csv")
        if k2 < 0:
            k2 = rest.find(".")
        s = s[:k1] + rest[k2:]
    return s.replace("_}", "}")


# =============================================================================
#   STEP 1+2 : transfer on raw spectra, then pretreatment  (Xfer-PP data flow)
# =============================================================================
@dataclass
class XferPPResult:
    L_main: TPSet                    # PP'd TP the Tcv / Pwhole step runs on
    L_val: Optional[TPSet]           # PP'd CS & Val (None when there is no Val)
    name_main: str                   # its MATLAB file name (titles / model names)
    name_val: str
    scheme_used: str
    std: Optional[TransferStandards] = None
    xfer: Optional[XferModel] = None
    tp_val_cab: Optional[TPSet] = None


def _with_X(S: SpecSet, X: np.ndarray, raw: Optional[np.ndarray] = None) -> SpecSet:
    n = S.copy()
    n.Atrainpk = np.atleast_2d(np.asarray(X, dtype=float)).copy()
    if raw is not None:
        n.RawSpectra = np.atleast_2d(np.asarray(raw, dtype=float)).copy()
    return n


def _pp_TP(tp: TPSet, pp1: str, pp2: str) -> TPSet:
    """AT2RS (Atrainpk -> RawSpectra) and then pp1+pp2 on T and P."""
    T_raw = tp.T.Atrainpk
    P_raw = tp.P.Atrainpk if tp.P is not None else None
    Xt, Xp = apply_pp1_pp2(T_raw, pp1, pp2, P_raw)
    T = _with_X(tp.T, Xt, T_raw)
    P = None
    if tp.P is not None:
        P = _with_X(tp.P, Xp if Xp is not None else tp.P.Atrainpk, P_raw)
    out = TPSet(T, P, tp.tag, tp.saCTCP, tp.T.RawSpectra.copy())
    return out


def xfer_then_pp(ds: Dataset, ctx: RunContext, pp1: str, pp2: str, scheme: str) -> XferPPResult:
    st = ctx.st
    CS = ds.CS
    base = split_scheme(scheme)[0]
    tp_xrs_cab = tp_val_cab = None
    std = xm = None
    scheme_used = scheme
    name_xrs_cab = name_val_cab = ""
    if ds.XRS is not None:
        std = transfer_standards(CS, ds.XRS, ds.XSmst)
        if std.forced_woCabXfer and scheme != "woCabXfer":
            note("XferScheme has been forced to switch to \"woCabXfer\" (no matching spectra between "
                 "scouting set and calibration set)")
            scheme_used = "woCabXfer"
            base = "woCabXfer"
        xm = build_transfer(scheme_used, std)
        T_cab = xm.apply_T(CS.Atrainpk)
        P_cab = xm.apply_P(ds.XRS.Atrainpk)
        tp_xrs_cab = TPSet(_with_X(CS, T_cab), _with_X(ds.XRS, P_cab), "{T-CS_P-XRS}", None)
        pfT = atfile_name("{T-CS_P-XRS}", TPSet(CS, ds.XRS))
        name_xrs_cab = cabxfer_fname(scheme_used, pfT, CS.n, ds.XRS.n, std.snXS, "Test")
        # Tset-only AT file for AAQP (CabXferd master RawSpectra, before PP) -> StepByStep
        T_only = _with_X(CS, T_cab)
        ctx.save_sbs(atfile_name("{" + base + "-Mst_RawSpectra_Tset-Only_for-AAQP}", T_only),
                     LAT_cab_T_only(T_only))
        if ds.Val is not None:
            V_cab = xm.apply_P(ds.Val.Atrainpk)
            tp_val_cab = TPSet(_with_X(CS, T_cab), _with_X(ds.Val, V_cab), "{T-CS_P-Val}", None)
            name_val_cab = cabxfer_fname(scheme_used, pfT, CS.n, ds.Val.n, std.snXS, "Val")
            write_pro_files(ctx, tp_val_cab, base, with_val=True)
        # without a Val set MATLAB's Val transfer fails on an empty path (caught by its
        # try/end), so no PRO csv/xlsx is written and the PP'd T + XRS become "CS&Val"
    elif ds.Val is not None:
        tp_val_cab = TPSet(CS, ds.Val, "{T-CS_P-Val}")
        name_val_cab = atfile_name("{T-CS_P-Val}", tp_val_cab)
        scheme_used = scheme
    # ---- PP on the transferred spectra -----------------------------------
    L_val = L_xrs = None
    name_val = name_xrs = ""
    sPP12 = f"_pp1-{pp1}_pp2-{pp2}"
    if tp_val_cab is not None:
        L_val = _pp_TP(tp_val_cab, pp1, pp2)
        name_val = atfile_name("(" + scheme_used + sPP12 + "_CS&Val_for-AAQP)", L_val)
        ctx.save_sbs(name_val, LAT_cab_TP(L_val, True) if tp_xrs_cab is not None else LAT_TP_AT2RS(L_val))
    if tp_xrs_cab is not None:
        L_xrs = _pp_TP(tp_xrs_cab, pp1, pp2)
        nm = name_xrs_cab.replace(".mat", "_AT2RS_XRS.mat")
        name_xrs = pp_rename(nm, pp1, pp2, L_xrs.T.Atrainpk.shape[1])
        if tp_val_cab is None:
            name_cv = atfile_name("(" + scheme_used + sPP12 + "_CS&Val_for-AAQP)", L_xrs)
            ctx.save_sbs(name_cv, LAT_cab_TP(L_xrs, True))
            return XferPPResult(L_xrs, None, name_cv, "", scheme_used, std, xm, None)
    if L_xrs is None and L_val is None:
        # CS only: Tset alone (T = P = CS, Tcv decides everything)
        src = CS
        P_src = CS.SAT_orig_Avg_All if CS.SAT_orig_Avg_All is not None else CS
        tp_cs = TPSet(src, P_src, "{CS-ONLY}")
        L_cs = _pp_TP(tp_cs, pp1, pp2)
        ctx.save_sbs(atfile_name("(" + scheme_used + sPP12 + "_Tset-Only_for-AAQP)", L_cs.T),
                     LAT_T_only(L_cs.T))
        nm = atfile_name("{CS-ONLY}", tp_cs).replace(".mat", "_wAT2RS_For_PRO_testing.mat")
        return XferPPResult(L_cs, None, pp_rename(nm, pp1, pp2, L_cs.T.Atrainpk.shape[1]), "",
                            scheme_used, std, xm, None)
    if L_xrs is not None:
        return XferPPResult(L_xrs, L_val, name_xrs, name_val, scheme_used, std, xm, tp_val_cab)
    return XferPPResult(L_val, L_val, name_val, name_val, scheme_used, std, xm, tp_val_cab)


def write_pro_files(ctx: RunContext, tp: TPSet, base: str, with_val: bool):
    """Transferred CS (not PPd) for MicroNIR PRO: xlsx+csv in TMP_AQP and csv in
    Results4PRO\\tmpN  (the .sam export needs SamLibrary.dll and is skipped)."""
    T = tp.T
    wvl = T.wvl
    ids = list(T.AclabelT)
    conc = list(T.conc)
    X = T.Atrainpk
    sPZ = ""
    if wvl.size < MN_WVL.size:
        Npz = MN_WVL.size - wvl.size
        sPZ = f"_Pad{Npz}ZerosCol"
        lia, locb = ismembertol(wvl, MN_WVL, 1e-3)
        Xp = np.zeros((X.shape[0], MN_WVL.size))
        Xp[:, locb[lia]] = X[:, lia]
        X, wvl = Xp, MN_WVL
    rows = [["ID", T.clsname] + list(wvl)]
    for i in range(len(ids)):
        rows.append([ids[i], conc[i]] + list(X[i]))
    atname = atfile_name("{T-" + base + "-CS-woPPd_P-Val}", tp)
    os.makedirs(ctx.tmp, exist_ok=True)
    csv_name = pro_csv_name(atname, ".csv").replace(".csv", sPZ + ".csv")
    if with_val:
        xlsx_name = pro_csv_name(atname, ".xlsx").replace(".xlsx", sPZ + ".xlsx")
        if len(xlsx_name) > 135:
            warn("xls_fname exceeds 135 -> did NOT write to xlsx file")
        else:
            write_cell_xlsx(os.path.join(ctx.tmp, xlsx_name), rows)
            disp_with_border(xlsx_name + " has been created")
    write_cell_csv(os.path.join(ctx.tmp, csv_name), rows)
    if (not with_val) or ctx.st.Option_Output_SAM:
        ctx.sam_folder = create_next_subfolder(ctx.results4pro, "tmp")
        shutil.copy2(os.path.join(ctx.tmp, csv_name), os.path.join(ctx.sam_folder, csv_name))


def get_snow() -> str:
    return _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S").replace(":", "_").replace(" ", "_")


def seconds2human(secs: float) -> str:
    if secs < 1:
        return roundns(secs, 3) + " sec."
    secs = int(round(secs))
    parts = []
    for unit, n in (("hour", 3600), ("minute", 60), ("second", 1)):
        if secs >= n:
            q, secs = divmod(secs, n)
            parts.append(f"{q} {unit}{'s' if q != 1 else ''}")
    return " ".join(parts)


# =============================================================================
#   STEP 3 : PLS  (SVR_PLS_Opm_by_Testset_then_AppliedToPset_MLtool +
#                  PLS_inside_PLS_predict_ONLY_MLtool)
# =============================================================================
def pls_stage(xp: XferPPResult, ctx: RunContext, cnt: int, pp1: str, pp2: str,
              scheme_label: str) -> Dict[str, Any]:
    st = ctx.st
    figs = ctx.figs
    pro = st.AQP_class == "pro"
    L = xp.L_main
    T = L.T
    Ana = T.clsname
    CurTcv = st.CurTcvModelParaOpmScheme
    ModelPara_Opm_Scheme = "User-Pick" if CurTcv == "User-Pick" else "ModelParaOpmBy-Tcv"
    tcv_scheme = INTERNAL_SETTINGS["Tcv_scheme"]
    # PlsfactorScan trimmed so every fold can still be fitted (Apr 19, 2023)
    nQ = len(set(T.AclabelT))
    nFolds = int(math.floor(math.sqrt(nQ)))
    maxP = nQ - int(math.ceil(nQ / nFolds)) - 1
    scan = [k for k in st.PlsfactorScan_default if k <= maxP]
    if not scan:
        raise AQPError(f"too few samples ({nQ}) for the PLS factor scan")
    XT, YT, names = T.Atrainpk, T.conc, list(T.AclabelT)
    avgTP = st.Spectra_Avg_Method == "Spectra_Avg_T-Mean_P-All"
    if avgTP:
        XP_tcv, YP_tcv, names_P = L.P.Atrainpk, L.P.conc, list(L.P.AclabelT)
    nameTP = xp.name_main
    out: Dict[str, Any] = {"figs": []}
    ctx.progress("running cross validation ...")
    # ---- Self-Pred (RMSEC)  (pro only) ----
    if pro:
        rmsec = [PLS_Tcv(XT, YT, names, k, "Self-Pred")["RMSE"] for k in scan]
        f = figs.new("RMSEC")
        if f is not None:
            ax = f.add_subplot(111)
            ax.plot(scan, rmsec, "b-*")
            ax.set_xlabel("PLS factor"); ax.set_ylabel("RMSEC")
            ax.set_title(f"{Ana}    Tset Self Prediction\n{nameTP}", fontsize=8)
            figs.show(f)
        out["RMSEC"] = rmsec
    # ---- Tcv ----
    t0 = time.time()
    OUT_Tcv = []
    for k in scan:
        if avgTP:
            OUT_Tcv.append(PLS_Tcv_P_Avg_All(XT, YT, names, k, XP_tcv, YP_tcv, names_P, tcv_scheme))
        else:
            OUT_Tcv.append(PLS_Tcv(XT, YT, names, k, tcv_scheme))
    sTime = seconds2human(time.time() - t0)
    rmsecv = np.array([o["RMSE"] for o in OUT_Tcv])
    opm, rm_opm, res_x = pick_opm_factor(CurTcv, scan, rmsecv, st.PLSfactor_Opm_User_Pick)
    out.update({"PlsfactorScan": scan, "RMSECV": rmsecv, "PLSfactor_Opm_Tcv": opm,
                "RMSE_PLS_Opm_Tcv": rm_opm, "knee_res_x": res_x})
    f = figs.new("RMSECV")
    if f is not None:
        ax = f.add_subplot(111)
        ax.plot(scan, rmsecv, "k-*", lw=2)
        ax.plot([opm], [rm_opm], "o", color=ML_G, ms=10, mfc=ML_G)
        if res_x is not None and not (isinstance(res_x, float) and math.isnan(res_x)):
            ax.axvline(res_x, color=(1, 0.5, 0))
        ax.set_xlabel("PLS factor"); ax.set_ylabel("RMSECV")
        if pro:
            ax.set_title(f"{nameTP}\n{Ana}  {ModelPara_Opm_Scheme}   PLSfactor_opm ={opm}  RMSE={roundns(rm_opm, 3)}\n"
                         f"Tcv_scheme = {tcv_scheme}   nFolds={OUT_Tcv[0]['nFolds']}   Time Spent: {sTime}"
                         f"     Opm-LV by: {CurTcv}", fontsize=8)
        else:
            ax.set_title(f"PLSfactor_opm ={opm}  RMSECV={roundns(rm_opm, 3)}")
        figs.show(f)
    # ---- whole Pset of L (XRS when there is one) ----
    XPw = L.P.Atrainpk if L.P is not None else XT
    YPw = L.P.conc if L.P is not None else YT
    sa_Pw = PLS_indv(XT, XPw, YT, YPw, scan)
    allR = np.array([r["RMSE"] for r in sa_Pw])
    if np.all(np.isnan(allR)):
        imin = 0
    else:
        imin = int(np.nanargmin(allR))
    out["PLSfactor_PLS_min"] = sa_Pw[imin]["PLSfactor"]
    out["RMSE_PLS_min"] = float(allR[imin])
    # ---- validation set ----
    sAltOpmPLS = ""
    has_val = xp.L_val is not None
    if has_val:
        Lv = xp.L_val
        X_V, Y_V = Lv.P.Atrainpk, Lv.P.conc
        XT_v, YT_v = (Lv.T.Atrainpk, Lv.T.conc) if Lv.T.n > XT.shape[0] else (XT, YT)
        rv = PLS_indv(XT_v, X_V, YT_v, Y_V, [opm])[0]
    else:
        tc = (PLS_Tcv_P_Avg_All(XT, YT, names, opm, XP_tcv, YP_tcv, names_P, tcv_scheme) if avgTP
              else PLS_Tcv(XT, YT, names, opm, tcv_scheme))
        if not (tc["RMSE"] == rm_opm or (math.isnan(tc["RMSE"]) and math.isnan(rm_opm))):
            raise AQPError("mismatch between outPLS_Tcv_Opm.RMSE vs RMSE_PLS_Opm_Tcv")
        rv = {"PLSfactor": opm, "RMSE": rm_opm, "Yest_PLS": tc["Yest_all"],
              "Y_iAna_P": (YP_tcv if avgTP else YT).copy(), "beta": float("nan")}
    out["saPLS_Results_Val"] = rv
    # ---- final model (pro) ----
    if pro:
        if has_val:
            beta = rv["beta"]
            Pest = pls_predict(X_V, beta)
            FM = {"beta": beta, "Apred": X_V, "Pest": Pest.reshape(-1, 1), "Pref": rv["Y_iAna_P"].reshape(-1, 1),
                  "RMSEP": rms_error_woNaN_N(Pest, rv["Y_iAna_P"]), "pp1": pp1, "pp2": pp2,
                  "RawSpectra_Val": Lv.P.RawSpectra,
                  "RawSpectra_CS_aft_CabXfer_bef_PPs": L.T.RawSpectra,
                  "Atrainpk_CS_aft_CabXfer_aft_PPs": L.T.Atrainpk,
                  "meanXT": L.T.Atrainpk.mean(axis=0).reshape(1, -1), "meanYT": float(np.mean(L.T.conc))}
            core = find_keyword_between_markers(xp.name_val, "Atrainpketc_", "_nsampXS")
            if not core:
                core = find_keyword_between_markers(xp.name_val, "{", "}") or ""
            fn = f"Beta_etc_FinalModel_cnt-{cnt}_{Ana}_PLSfactor{rv['PLSfactor']}_{{{core}}}.mat"
        else:
            FM = {"RMSEP": rv["RMSE"], "pp1": pp1, "pp2": pp2}
            core = find_keyword_between_markers(nameTP, "Atrainpketc_", "_nsampP") or ""
            fn = f"Beta_etc_FinalModel_cnt-{cnt}_{Ana}_PLSfactor{rv['PLSfactor']}_{core}.mat"
            fn = fn.replace("}]", "}")
        if "}_nsampT" in fn:
            fn = fn.replace("}_nsampT", f"_pp2-{pp2}}}_nsampT")
        else:
            fn = fn.replace("}.mat", f"_pp2-{pp2}}}.mat")
        if "_pp2" in fn and "pp1" not in fn:
            fn = fn.replace("_pp2", f"_pp1-{pp1}_pp2")
        savemat(os.path.join(ctx.final_models, fn), FM)
        out["FinalModel_file"] = fn
    # ---- 45 degree plot (outlier analysis is disabled: RMSE_thres4OLs = Inf) ----
    Yref = rv["Y_iAna_P"]
    Yest = rv["Yest_PLS"]
    res = Yref - Yest
    ok = ~np.isnan(res)                       # loc_Clean = find(dd <= RMSE_thres4OLs=Inf)
    RMSE_clean = rms_error_woNaN_N(Yref[ok], Yest[ok]) if ok.any() else float("nan")
    STD_clean = float(np.std(res[ok], ddof=1)) if ok.sum() > 1 else float("nan")
    Bias_clean = float(np.mean(res[ok])) if ok.any() else float("nan")
    f = figs.new("Pest_45deg")
    if f is not None:
        ax = f.add_subplot(111)
        ax.plot(Yref, Yest, "b*", ls="none")
        hi = max(ax.get_xlim()[1], ax.get_ylim()[1])
        ax.plot([0, hi], [0, hi], "-", color=ML_Y, lw=2)
        ax.set_xlabel("True Conc"); ax.set_ylabel("Est Conc")
        sm = (f"{CurTcv}-->OpmPLSfactor={rv['PLSfactor']}{sAltOpmPLS}   RMSE ={roundns(RMSE_clean, 3)}"
              f"      STD ={roundns(STD_clean, 3)}      Bias ={roundns(Bias_clean, 3)}")
        if has_val:
            if pro:
                tt = f"{Ana}    Pest by PLS \nTP: {nameTP.replace('{', '').replace('}', '')}\nVal: {xp.name_val}\n{sm}"
            else:
                tt = (f"OpmPLSfactor={rv['PLSfactor']}{sAltOpmPLS}   RMSEP ={roundns(RMSE_clean, 3)}"
                      f"      STD ={roundns(STD_clean, 3)}      Bias ={roundns(Bias_clean, 3)}")
        else:
            stp = "Tset Alone: " + (find_keyword_between_markers(nameTP, "", "_nsampP") or nameTP)
            tt = f"{Ana}    {pp1} + {pp2}\n{stp.replace('{', '').replace('}', '')}\n{sm}"
        tt += f"\n{nameTP}\n{st.Spectra_Avg_Method}"
        ax.set_title(tt, fontsize=7)
        figs.show(f)
    out["RMSE_clean"], out["STD_clean"], out["Bias_clean"] = RMSE_clean, STD_clean, Bias_clean
    # ---- RMSEP vs PLS factor on the validation set ----
    if has_val:
        sa_scan = PLS_indv(XT_v, X_V, YT_v, Y_V, scan)
        rmsep_scan = [r["RMSE"] for r in sa_scan]
        out["RMSEP_scan"] = rmsep_scan
        f = figs.new("RMSEP_vs_PLSfactor")
        if f is not None:
            ax = f.add_subplot(111)
            ax.plot([rv["PLSfactor"]], [rv["RMSE"]], "o", color=ML_Y, ms=12, mfc=ML_Y)
            ax.plot(scan, rmsep_scan, "b-*", lw=2)
            ax.set_xlabel("PLS factor"); ax.set_ylabel("RMSEP")
            if pro:
                ax.set_title(f"{Ana}    \n{xp.name_val}\n{st.cDataFlow[0]}  {pp1}  {scheme_label}  woUDM\n"
                             f"Val set {CurTcv}-->OpmPLSfactor={rv['PLSfactor']}{sAltOpmPLS}  RMSE={roundns(rv['RMSE'], 3)}",
                             fontsize=8)
            else:
                ax.set_title(f"OpmPLSfactor={rv['PLSfactor']}{sAltOpmPLS}  RMSEP={roundns(rv['RMSE'], 3)}")
            figs.show(f)
    # ---- what Run AQP reports (SVR_PLS_Opm_by_Testset_then_AppliedToPset_MLtool) ----
    out["OUT"] = {"Results4sRMSE_Val": roundns(rv["RMSE"], 4),
                  "Results4sOpmPLSfactor": roundns(rv["PLSfactor"], 0),
                  "Results4RMSEP": float(rv["RMSE"])}
    return out


# =============================================================================
#   BATCH RUN  (BatchRun_AutoQuant_DA_pipeline + Run_AQP)
# =============================================================================
@dataclass
class BatchResult:
    OUT_cln: List[Dict[str, Any]]
    Path4OUT_cln_AQP_FinalSubfolder: Optional[str]
    fname_Results_AQP_OUT_cln: Optional[str]
    runs: List[Dict[str, Any]]
    outOpmFM: Optional[Dict[str, Any]]
    root: str
    out_SbS: Optional[str] = None
    nfile_CR: int = 0
    dataset: Optional[Dataset] = None
    elapsed: float = 0.0


def read_PP1_PP2_xlsx(path: str) -> Tuple[List[str], List[str]]:
    """load_AQP_PP1_PP2_xlsx: the PP1 / PP2 columns of a PP1_PP2*.xlsx."""
    rows = read_xlsx_raw(path)
    head = [str(v).strip() if isinstance(v, str) else "" for v in rows[0]]
    if "PP1" not in head or "PP2" not in head:
        raise AQPError(f"{os.path.basename(path)} needs a 'PP1' and a 'PP2' column")
    i1, i2 = head.index("PP1"), head.index("PP2")
    cPP1, cPP2 = [], []
    for r in rows[1:]:
        a, b = r[i1], r[i2]
        if isinstance(a, float) and math.isnan(a) and isinstance(b, float) and math.isnan(b):
            continue
        cPP1.append(a if isinstance(a, str) else num2str(a))
        cPP2.append(b if isinstance(b, str) else num2str(b))
    return cPP1, cPP2


def _dsname_from_corename(core: str, ana: str) -> str:
    """find_keyword_merge_dual_curly_bracket_w_targetstring_remain(...).cb1"""
    opens = _findall_overlap(core, "{")
    closes = _findall_overlap(core, "}")
    if not opens or not closes:
        return find_keyword_between_markers(core, "", "}") or ""
    st = opens[-1] + 1
    if len(opens) > 1 and len(closes) != 1:
        cb2 = core[st:closes[-1]]
    else:
        later = [k for k in closes if k > st - 1]
        cb2 = core[st:later[-1]] if later else ""
    tmp1 = core.replace("{" + cb2 + "}", "")
    cb1 = find_keyword_between_markers(tmp1, "{", "}")
    ds = cb1 if cb1 else (find_keyword_between_markers(core, "{", "}") or "")
    return ds.replace("_" + ana, "")


def batch_run(st: RunSettings, figs: Optional[FigMgr] = None, progress=None,
              pick_analyte=None) -> BatchResult:
    t_start = time.time()
    progress = progress or (lambda s: None)
    figs = figs or FigMgr("none")
    ctx = RunContext(st, figs, progress)
    root = ctx.root
    tmp_folder_rm_mk("TMP_AQP", root)
    if os.path.isdir(ctx.sbs_pre):
        shutil.rmtree(ctx.sbs_pre, ignore_errors=True)
    os.makedirs(ctx.sbs_pre, exist_ok=True)
    if not (st.cUDMas_scheme and st.cUDMas_scheme[0] == "woUDM"):
        raise AQPError('current AQP only support "cUDMas_scheme" as "{woUDM}" !!!')
    if st.cDataFlow != ["Xfer-PP"]:
        raise AQPError("this port runs the 'Xfer-PP' data flow (the only one the GUI offers)")
    cPP1, cPP2 = list(st.cPP1), list(st.cPP2)
    if not cPP2:
        cPP2 = ["SNV"]
    progress("loading the dataset ...")
    ds = load_dataset(st.path_XLSX, st.AnaName, st.Spectra_Avg_Method, st.AQP_class, pick_analyte)
    cCabXfer = list(st.cCabXfer_scheme)
    if ds.cCabXfer_forced:
        cCabXfer = [ds.cCabXfer_forced]
    for sc in cCabXfer:
        b = split_scheme(sc)[0]
        if b not in CABXFER_SUPPORTED and b not in CABXFER_NEEDS_PLS_TOOLBOX:
            raise AQPError("CabXfer scheme not supported: " + sc)
    Ana = ds.AnaName
    # ---- StepByStep: CS / XRSmst before and after the grid matching ----
    sfold = os.path.basename(os.path.normpath(st.path_XLSX))
    if ds.XRS is not None:
        tp = TPSet(ds.CS, ds.XRS, "{T-CS_P-XRS}")
        ctx.save_sbs(atfile_name("{T-CS_P-XRS}", tp), LAT_TP(tp))
        ctx.save_sbs(atfile_name("After-MatchGrids_T-CS_P-MstXRS", tp), LAT_TP(tp))
    elif ds.Val is not None:
        tp = TPSet(ds.CS, ds.Val, "{T-CS_P-Val}")
        ctx.save_sbs(atfile_name("{T-CS_P-Val}", tp), LAT_TP(tp))
        ctx.save_sbs(atfile_name("After-MatchGrids_T-CS_P-MstVal", tp), LAT_TP(tp))
    if ds.XRS is not None or ds.Val is not None:
        c = ds.CS_orig
        ctx.save_sbs(f"Atrainpketc_CS-Before-MatchGrids_({sfold})_pp1-none_pp2-none_nvar{c.Atrainpk.shape[1]}"
                     f"_nsamp{c.n}.mat", LAT_T_only(c))
    if ds.XRS is not None and ds.XSmst_orig is not None:
        c = ds.XSmst_orig
        ctx.save_sbs(f"Atrainpketc_XSmst-Before-MatchGrids_({sfold})_pp1-none_pp2-none_nvar{c.Atrainpk.shape[1]}"
                     f"_nsamp{c.n}.mat", LAT_T_only(c))
    # ---- lines (legend) x x-ticks ----
    Lines = st.Lines_Xticks_picked
    if Lines == "Lines_PP1--Xticks_CabXfer":
        if len(cPP1) > 1 and len(cPP2) == 1:
            cLegendLoop = list(cPP1)
        elif len(cPP1) == 1 and len(cPP2) > 1:
            cLegendLoop = list(cPP2)
        elif len(cPP1) == 1 and len(cPP2) == 1:
            cLegendLoop = list(cPP1)
        else:
            cLegendLoop = [a + "+" + b for a, b in zip(cPP1, cPP2)]
        call_XLabel_Grids = list(cCabXfer)
    elif Lines == "Lines_DSn--Xticks_UDM":
        cLegendLoop = ["single_XLSX_folder"]
        call_XLabel_Grids = list(st.cUDMas_scheme)
    elif Lines == "Lines_PP1--Xticks_UDM":
        raise AQPError("can not handle this case of call_XLabel_Grids yet (Lines_PP1--Xticks_UDM)")
    else:
        raise AQPError(f"Lines_Xticks option '{Lines}' is not supported for a single XLSX folder")
    if st.AQP_class == "pro":
        N_jLine = len(cLegendLoop)
    else:
        N_jLine = min(len(cLegendLoop), len(GUI_MENUS["cPP1"]))
        cLegendLoop = cLegendLoop[:N_jLine]
    Total_cnt = len(cLegendLoop) * len(call_XLabel_Grids)
    corename = (f"{cPP1[0]}_{st.cDataFlow[0]}_{cCabXfer[0]}_nLine{len(cLegendLoop)}"
                f"_nXtick{len(call_XLabel_Grids)}").replace("{PRO}", "(PRO)")
    tmp_folder_rm_mk("Result4FinalModels", root)
    figs_opm_dir = os.path.join(root, "Figs4OpmModel_AQP")
    # ---- the loops ----
    runs = []
    cnt = 0
    best_RMSEP = float("inf")
    best_figs: List[Any] = []
    last_pp = (cPP1[0], cPP2[0])
    for j in range(N_jLine):
        if len(cPP1) == 1 and len(cPP2) > 1:
            pp1, pp2 = cPP1[0], cPP2[j]
        elif len(cPP2) == 1 and len(cPP1) > 1:
            pp1, pp2 = cPP1[j], cPP2[0]
        elif len(cPP2) == 1 and len(cPP1) == 1:
            pp1, pp2 = cPP1[0], cPP2[0]
        else:
            pp1, pp2 = cPP1[j], cPP2[j]
        last_pp = (pp1, pp2)
        for iX in range(len(call_XLabel_Grids)):
            cnt += 1
            scheme = cCabXfer[iX] if Lines == "Lines_PP1--Xticks_CabXfer" else cCabXfer[0]
            progress(f"iter-{cnt}/{Total_cnt}")
            figs.figs = []
            xp = xfer_then_pp(ds, ctx, pp1, pp2, scheme)
            res = pls_stage(xp, ctx, cnt, pp1, pp2, scheme)
            OUT = dict(res["OUT"])
            OUT["cLegendLoop_jLine"] = cLegendLoop[j]
            OUT["CabXfer_scheme"] = cCabXfer[iX] if iX < len(cCabXfer) else cCabXfer[0]
            res["pp1"], res["pp2"], res["scheme"], res["cnt"] = pp1, pp2, scheme, cnt
            res["OUT"] = OUT
            res["figs"] = list(figs.figs)
            runs.append(res)
            if Total_cnt > 1:
                if cnt == 1 or OUT["Results4RMSEP"] < best_RMSEP:
                    tmp_folder_rm_mk("Figs4OpmModel_AQP", root)
                    figs.save_set(figs.figs, figs_opm_dir, f"Cnt{cnt}")
                    if figs.mode == "gui":
                        for f in best_figs:
                            if f not in figs.figs:
                                figs.close_fig(f)
                    best_figs = list(figs.figs)
                elif figs.mode == "gui":
                    for f in figs.figs:
                        figs.close_fig(f)
                if not math.isnan(OUT["Results4RMSEP"]):
                    best_RMSEP = min(best_RMSEP, OUT["Results4RMSEP"])
    # ---- results file ----
    OUT_cln = []
    for r in runs:
        o = dict(r["OUT"])
        o["AnaName"] = Ana
        o["Spectra_Avg_Method"] = st.Spectra_Avg_Method
        OUT_cln.append(o)
    core = corename
    if len(cPP1) > 1:
        core = find_keyword_between_markers(core, "Xfer-PP_", "") or core
    core = core.replace("woCabXfer", "()")
    core = "{" + sfold + "}_" + core + "_" + st.Spectra_Avg_Method
    core = core.replace("}", "_" + Ana + "}")
    if "DerSGFL" in last_pp[0]:
        sBC = "_{" + last_pp[1] + "-PP2}"
    else:
        sBC = "_{" + last_pp[0] + "-PP1}"
    core = core + sBC
    fname_R = ("Results_AQP_" + core + ".mat").replace("(PRO)", "").replace("{PRO}", "")
    core = core.replace("SNV{PRO}", "SNV(PRO)")
    dsname = _dsname_from_corename(core, Ana)
    if st.pathfname_PPn_xlsx:
        PPn = os.path.splitext(os.path.basename(st.pathfname_PPn_xlsx))[0]
        PPn = PPn.replace("pp1_pp2", "").replace("PP1_PP2", "")
    else:
        PPn = "_" + cPP1[0] + "+" + cPP2[0]
    PPn = PPn.replace("{PRO}+SNV{PRO}", "+SNV_PRO")
    sAna = "[" + Ana + "]" if INTERNAL_SETTINGS["NEWDIR_include_AnaName_yes"] else ""
    sX = ("(" + cCabXfer[0] + ")").replace("woCabXfer", "") if (
        INTERNAL_SETTINGS["NEWDIR_include_CabXfer_yes"] and len(cCabXfer) == 1) else ""
    NEWDIR = dsname + sAna + sX + PPn
    P4 = tmp_folder_rm_mk("Result4OUT_cln_AQP", root)
    final_sub = os.path.join(P4, NEWDIR)
    os.makedirs(final_sub, exist_ok=True)
    xt = check_XTick_before_store(final_sub, [o["cLegendLoop_jLine"] for o in OUT_cln],
                                  INTERNAL_SETTINGS["XTick_policy"])
    final_sub = xt["dest_folder"]
    if xt["save_yes"]:
        pfn = os.path.join(final_sub, fname_R)
        if len(pfn) > 259:
            fname_R = "Results_" + _dt.datetime.now().strftime("%H%M%S-%m%d%y") + ".mat"
            pfn = os.path.join(final_sub, fname_R)
        write_results_mat(pfn, OUT_cln)
    else:
        note("[XTick guard] Results_AQP_OUT_cln was NOT copied into the comparison folder")
    # ---- AQP_App_Emulator : the optimal final model ----
    outOpmFM = None
    if st.AQP_class == "pro":
        try:
            outOpmFM = app_emulator(ctx.final_models, cCabXfer)
        except Exception as e:
            warn(f"AQP_App_Emulator: {e}")
            outOpmFM = None
        if outOpmFM is not None and outOpmFM.get("opm_file"):
            try:
                json_prepare(os.path.join(ctx.final_models, outOpmFM["opm_file"]), outOpmFM)
            except Exception as e:
                warn(f"json_AQP_OpmModel_Sequence: {e}")
    # ---- show the optimal model's figures (several runs) ----
    if Total_cnt > 1 and figs.mode == "gui":
        for r in runs:
            for f in r["figs"]:
                if f not in best_figs:
                    figs.close_fig(f)
    # ---- RMSEP summary across the scan ----
    if st.AQP_class == "pro" and outOpmFM is not None and Total_cnt > 1:
        plot_rmsep_summary(figs, runs, cLegendLoop, call_XLabel_Grids, outOpmFM, st, cPP1, cPP2,
                           cCabXfer, last_pp)
    # ---- StepByStep folder into TMP_AQP ----
    out_SbS = tmp_folder_rm_mk("StepByStep", ctx.tmp)
    for fn in glob.glob(os.path.join(ctx.sbs_pre, "Atrainpk*.mat")):
        shutil.move(fn, os.path.join(out_SbS, os.path.basename(fn)))
    shutil.rmtree(ctx.sbs_pre, ignore_errors=True)
    nCR = count_results(final_sub)
    return BatchResult(OUT_cln, final_sub, fname_R, runs, outOpmFM, root, out_SbS, nCR, ds,
                       time.time() - t_start)


def write_results_mat(pfn: str, OUT_cln: List[Dict[str, Any]]):
    fields = ["Results4sRMSE_Val", "Results4sOpmPLSfactor", "Results4RMSEP", "cLegendLoop_jLine",
              "CabXfer_scheme", "AnaName", "Spectra_Avg_Method"]
    savemat(pfn, {"OUT_cln": _struct_array(OUT_cln, fields)})


def read_results_mat(pfn: str) -> List[Dict[str, Any]]:
    import scipy.io as sio
    L = sio.loadmat(pfn, squeeze_me=True, struct_as_record=False)
    if "OUT_cln" not in L:
        raise AQPError(f"no OUT_cln inside --> {pfn}")
    O = np.atleast_1d(L["OUT_cln"])
    out = []
    for o in O:
        d = {}
        for f in o._fieldnames:
            v = getattr(o, f)
            if isinstance(v, np.ndarray) and v.size == 0:
                v = ""
            elif isinstance(v, np.ndarray) and v.dtype.kind in "U":
                v = str(v)
            d[f] = v
        out.append(d)
    return out


def count_results(folder: Optional[str]) -> int:
    if not folder or not os.path.isdir(folder):
        return 0
    return len([f for f in os.listdir(folder) if "Results_AQP" in f and f.endswith(".mat")
                and os.path.isfile(os.path.join(folder, f))])


def app_emulator(path_FM: str, cCabXfer: List[str]) -> Optional[Dict[str, Any]]:
    """AQP_App_Emulator: pick the final model with the smallest RMSEP."""
    import scipy.io as sio
    opm = [f for f in os.listdir(path_FM) if f.startswith("OpmModel_Beta_etc_FinalModel_") and f.endswith(".mat")]
    if opm:
        return None
    files = [f for f in os.listdir(path_FM) if f.startswith("Beta_etc_FinalModel_") and f.endswith(".mat")]
    if not files:
        return None
    files = sortnat(files)
    saFM = []
    for f in files:
        saFM.append(sio.loadmat(os.path.join(path_FM, f), squeeze_me=True, struct_as_record=False))
    allR = np.array([float(np.ravel(d["RMSEP"])[0]) for d in saFM])
    k = int(np.nanargmin(allR)) if not np.all(np.isnan(allR)) else 0
    seq = {"PP1": str(saFM[k]["pp1"]), "PP2": str(saFM[k]["pp2"]),
           "CabXfer": cCabXfer[0] if len(cCabXfer) == 1 else "woCabXfer"}
    if seq["CabXfer"] == "T-CS_P-Val":
        seq["CabXfer"] = "woCabXfer"
    write_cell_xlsx(os.path.join(path_FM, "Seq_Opm_PLS_Models.xlsx"),
                    [list(seq.keys()), list(seq.values())])
    FM = {kk: v for kk, v in saFM[k].items() if not kk.startswith("__")}
    FM["seq"] = seq
    fn = ("OpmModel_" + files[k]).replace("{{", "{").replace("}.", ".").replace("_Val", "")
    savemat(os.path.join(path_FM, fn), FM)
    disp_with_border(os.path.join(path_FM, fn) + " has been saved")
    return {"RMSEP_Opm": float(allR[k]), "seq": seq, "opm_file": fn, "index": k, "FM": FM}


def json_prepare(pfn_opm: str, outOpmFM: Dict[str, Any]):
    """json_AQP_OpmModel_Sequence: the SG / SNV / PLS steps of the optimal model
    as a text file to hand-build the json (beta base64, big-endian doubles).
    NOTE: in the MATLAB build this step stopped with an undefined variable
    whenever the model carried Apred (i.e. a Val set existed), so the file was
    only written for the Tcv-only case; here it is written in both cases."""
    FM = outOpmFM["FM"]
    seq = outOpmFM["seq"]
    if "Apred" in FM and "beta" in FM:
        b = np.atleast_2d(np.asarray(FM["beta"], dtype=float)).reshape(-1, 1)
        re_calc = pls_predict(np.atleast_2d(FM["Apred"]), b)
        if not np.allclose(re_calc, np.ravel(FM["Pest"]), rtol=0, atol=1e-12 * max(1, np.max(np.abs(re_calc)))):
            raise AQPError("mismatch among beta, Apred, and Pest")
    stem = os.path.splitext(os.path.basename(pfn_opm))[0]
    snow = get_snow()
    p4 = tmp_folder_rm_mk("4json", os.path.dirname(pfn_opm))
    head = find_keyword_between_markers(stem, "", "}") or stem
    fname = os.path.join(p4, f"Prepare4json_{head}}}_{snow}.txt")
    ana = find_keyword_between_markers(stem, "FinalModel_", "_PLSfactor") or ""
    L = [f"% {ana}", f"% {snow}", "% " + "-" * 58, None, "SavitskyGolayStep"]
    PP1 = seq["PP1"]
    L.append(f'"Order": "{PP1[:1]}",')
    L.append(f'"Polynomial": "{find_keyword_between_markers(PP1, "[PO", "]") or ""}",')
    L.append(f'"Points": "{find_keyword_between_markers(PP1, "FL", "[") or ""}",')
    L += [None, "SnvStep", None, "PlsStep"]
    if "beta" in FM:
        beta = np.ravel(np.asarray(FM["beta"], dtype=">f8"))
        L.append('"Coefficients": "' + base64.b64encode(beta.tobytes()).decode() + '",')
    with open(fname, "w", newline="") as f:
        for s in L:
            f.write("\r\n" if s is None else s + " \r\n")
    return fname


# MATLAB's one-letter colours.  b, r, k match matplotlib's; g, c, m, y do not
# (matplotlib 'y' is a dark olive (0.75, 0.75, 0); MATLAB 'y' is pure yellow).
ML_G, ML_C, ML_M, ML_Y = (0, 1, 0), (0, 1, 1), (1, 0, 1), (1, 1, 0)


def plot_rmsep_summary(figs: FigMgr, runs, cLegendLoop, call_X, outOpmFM, st, cPP1, cPP2, cCabXfer, last_pp):
    f = figs.new("RMSEP_summary", size=(11, 6))
    if f is None:
        return
    scol = "krgbymcpolasvhkrgbymcpolasvhkrgb" * 2
    smk = ">+<pxodh>+<pxodh>+<pxodh>+<pxodh"
    smk = smk + smk[1:] + smk[0]
    colmap = {"k": "k", "r": "r", "g": ML_G, "b": "b", "y": ML_Y, "m": ML_M, "c": ML_C,
              "p": (1, 0.6, 0.8), "o": (1, 0.5, 0), "l": (0.6, 1, 0.6), "a": (0.6, 0.6, 0.6),
              "s": (0.75, 0.75, 0.75), "v": (0.5, 0, 0.5), "h": (0.5, 0.5, 0)}
    mkmap = {">": ">", "+": "+", "<": "<", "p": "p", "x": "x", "o": "o", "d": "d", "h": "h"}
    ax = f.add_subplot(111)
    lab = []
    for j, cl in enumerate(cLegendLoop):
        y = [r["OUT"]["Results4RMSEP"] for r in runs if r["OUT"]["cLegendLoop_jLine"] == cl]
        c = colmap.get(scol[j], "k")
        m = mkmap.get(smk[j], "o")
        if len(call_X) == 1 and len(cLegendLoop) > 1:
            ax.plot([j + 1], y, color=c, marker=m, lw=2, ms=12, ls="none")
        else:
            ax.plot(range(1, len(y) + 1), y, color=c, marker=m, lw=2, ms=12)
        lab.append(cl + "  " + st.cDataFlow[0])
    seq = outOpmFM["seq"]
    if len(call_X) == 1:
        xs = [i + 1 for i, c in enumerate(cLegendLoop) if c == seq["PP1"]]
        if not xs:
            xs = [i + 1 for i, c in enumerate(cLegendLoop) if c == seq["PP2"]]
    else:
        xs = [i + 1 for i, c in enumerate(call_X) if c == seq["CabXfer"]]
    sOpmSeq = "Opm Seq : " + "_".join(f"{k}-{v}" for k, v in seq.items())
    if not xs:
        tit1 = seq["PP1"] + "+" + seq["PP2"]
        xs = [i + 1 for i, c in enumerate(cLegendLoop) if c == tit1]
        sOpmSeq = "Opm Seq : " + tit1
    if xs:
        ax.plot(xs, [outOpmFM["RMSEP_Opm"]] * len(xs), color=(1, 0.5, 0), marker="h", ms=15, ls="none")
        lab.append("Opm Final Model RMSEP=" + roundns(outOpmFM["RMSEP_Opm"], 3))
    ax.legend(lab, fontsize=8)
    if len(call_X) == 1 and len(cLegendLoop) > 1:
        ticks = [c.replace("{PRO}", "\n               MN-PRO").replace("+", "\n") for c in cLegendLoop]
        ax.set_xticks(range(1, len(cLegendLoop) + 1))
        ax.set_xticklabels(ticks, rotation=-45, fontsize=8, ha="left")
        stit_call = call_X[0]
    else:
        ax.set_xticks(range(1, len(call_X) + 1))
        ax.set_xticklabels(call_X, rotation=-45, fontsize=10, ha="left")
        stit_call = ""
    ax.set_ylabel("RMSEP")
    sPP1 = "PP1-" + (cPP1[0] if len(cPP1) == 1 else cPP1[-1])
    stit = os.path.basename(os.path.normpath(st.path_XLSX))
    sPP2_alt = ""
    if len(cPP2) == 1 and last_pp[1] != "SNV":
        sPP2_alt = "   PP2-" + last_pp[1]
    stit2 = st.Spectra_Avg_Method.replace("Spectra_Avg_", "Spectra-") + sPP2_alt
    ax.set_title(f"{sPP1}  {stit}    Xfer scheme -->{cCabXfer[0]}  {stit_call}\n{stit2}\n{sOpmSeq}", fontsize=9)
    figs.show(f)


def check_XTick_before_store(sfolder: str, cXT_new: List[str], policy: str = "divert") -> Dict[str, Any]:
    """CmpR_check_XTick_before_store: never mix two different x-axis lists in
    one comparison folder ('divert' -> <folder>_XTickSet2, _XTickSet3, ...)."""
    out = {"dest_folder": sfolder, "save_yes": True, "match_yes": True, "msg": ""}
    if not cXT_new:
        return out
    ok, bad = _folder_xtick_compatible(sfolder, cXT_new)
    out["match_yes"] = ok
    if ok:
        return out
    msg = (f"[XTick guard] the new run does NOT share the x-axis (PP1+PP2 scheme) list of "
           f"{len(bad)} file(s) already in {sfolder} --> {', '.join(bad)}")
    if policy == "save-anyway":
        out["msg"] = msg + " -- stored in the same folder anyway"
    elif policy == "skip":
        out["save_yes"] = False
        out["msg"] = msg + " -- NOT stored"
    else:
        for k in range(2, 100):
            cand = f"{sfolder}_XTickSet{k}"
            if not os.path.isdir(cand):
                os.makedirs(cand, exist_ok=True)
                out["dest_folder"] = cand
                break
            ok2, _ = _folder_xtick_compatible(cand, cXT_new)
            if ok2:
                out["dest_folder"] = cand
                break
        out["msg"] = msg + " -- stored in a separate folder instead --> " + out["dest_folder"]
    note(out["msg"])
    return out


def _folder_xtick_compatible(sfolder: str, cXT_new: List[str]):
    if not os.path.isdir(sfolder):
        return True, []
    new = sorted(unique_appear_order(cXT_new))
    bad = []
    for f in sorted(os.listdir(sfolder)):
        if "Results_AQP" not in f or not f.endswith(".mat"):
            continue
        try:
            old = [str(o.get("cLegendLoop_jLine", "")) for o in read_results_mat(os.path.join(sfolder, f))]
        except Exception:
            continue
        old = [o for o in old if o]
        if not old:
            continue
        if sorted(unique_appear_order(old)) != new:
            bad.append(f)
    return (len(bad) == 0), bad


# =============================================================================
#   READING THE PANEL  (what Run_AQP does with the uicontrols)
# =============================================================================
def strip_default(s: str) -> str:
    return strip_after(s, "(").strip()


def popup_default(menu: List[str], sel: str, title: str) -> str:
    """Set_GUI_PopUpMenu_parameter_Default: the title row means '(default)'."""
    if sel == title:
        k = [i for i, s in enumerate(menu) if "(default)" in s.lower().replace(" ", "")]
        if not k:
            raise AQPError(f"no '(default)' entry in the {title} menu")
        return strip_default(menu[k[0]])
    return strip_default(sel)


def settings_from_panel(sel: Dict[str, str], path_XLSX: str, workdir: str,
                        pick_PPn_xlsx=None, pick_user_PLS=None, menus=None) -> RunSettings:
    """sel: {tag: selected string} for the popups of the panel."""
    menus = menus or GUI_MENUS
    aqp_class = sel.get("AQP_class", "pro")
    s1, s2 = sel["cPP1"], sel["cPP2"]
    PPn_xlsx = None
    if aqp_class != "lite":
        if s1 == "PP1_PP2_xlsx" or s2 == "PP1_PP2_xlsx":
            PPn_xlsx = pick_PPn_xlsx() if pick_PPn_xlsx else None
            if not PPn_xlsx:
                raise AQPError("no PP1_PP2*.xlsx picked")
            cPP1, cPP2 = read_PP1_PP2_xlsx(PPn_xlsx)
        else:
            cPP1 = [s1]
    else:
        cPP1 = [s1]
        if s1 == "PP1":
            menu = menus["cPP1"]
            k = [i for i, s in enumerate(menu) if "default" in s.lower()]
            if len(k) == 1:
                s1 = menu[k[0]]
            elif not k:
                s1 = menu[menu.index("PP1") + 1]
            else:
                raise AQPError("can not handle case that multiple default found")
        if "default" in s1.lower():
            s1 = (find_keyword_between_markers(s1, "", "(") or s1).strip()
        cPP1 = [s1]
    if s2 == "PP1_PP2_xlsx" or s1 == "PP1_PP2_xlsx":
        if PPn_xlsx is None:
            PPn_xlsx = pick_PPn_xlsx() if pick_PPn_xlsx else None
            if not PPn_xlsx:
                raise AQPError("no PP1_PP2*.xlsx picked")
            _, cPP2 = read_PP1_PP2_xlsx(PPn_xlsx)
    else:
        if s2 == "PP2":
            menu = menus["cPP2"]
            k = [i for i, s in enumerate(menu) if "default" in s.lower()]
            if len(k) == 1:
                s2 = menu[k[0]]
            elif not k:
                s2 = menu[menu.index("PP2") + 1] if menu.index("PP2") + 1 < len(menu) else "SNV"
            else:
                raise AQPError("can not handle case that multiple default found")
        if "default" in s2.lower():
            s2 = find_keyword_between_markers(s2, "", "(") or s2
        cPP2 = [s2.strip()]
    # UDMas_scheme
    menu = menus["UDMas_scheme"]
    su = sel.get("UDMas_scheme", menu[0])
    if su == "UDMas_scheme":
        cUDM = [popup_default(menu, su, "UDMas_scheme")]
    elif su == "ALL_UDM_schemes":
        dash = [i for i, s in enumerate(menu) if s.startswith("---")]
        if len(dash) != 1:
            raise AQPError('can not hanlde this case for "ALL_UDM_schemes"')
        cUDM = menu[dash[0] + 1:-1]
    elif su == "ALL_Above_UDM_schemes":
        cUDM = menu[1:menu.index(su)]
    else:
        cUDM = [strip_default(su)]
    # CabXfer_scheme
    sx = sel["CabXfer_scheme"]
    if sx == "CabXfer_scheme":
        raise AQPError("pls pick a calibration transfer scheme")
    if sx == "ALL_CabXfer_schemes":
        m = menus["CabXfer_scheme"]
        cX = m[1:m.index(sx)]
    else:
        cX = [sx]
    # Lines--Xticks
    sl = sel.get("Lines_Xticks", menus["Lines_Xticks"][0])
    if sl == "Lines--Xticks":
        sl = popup_default(menus["Lines_Xticks"], sl, "Lines--Xticks")
    else:
        sl = strip_default(sl)
    tcv = popup_default(menus["TcvModelParaOpmScheme"],
                        sel.get("TcvModelParaOpmScheme", "TcvModelParaOpmScheme"), "TcvModelParaOpmScheme")
    user_pick = float("nan")
    if tcv == "User-Pick" and aqp_class == "pro":
        user_pick = pick_user_PLS(INTERNAL_SETTINGS["PlsfactorScan"]) if pick_user_PLS else float("nan")
        if user_pick is None or (isinstance(user_pick, float) and math.isnan(user_pick)):
            raise AQPError("you can only pick one PLSfactor !!!")
    df = sel.get("cDataFlow", "Xfer-PP")
    return RunSettings(path_XLSX=path_XLSX, AnaName=None, cPP1=cPP1, cPP2=cPP2, cCabXfer_scheme=cX,
                       cDataFlow=[df], cUDMas_scheme=cUDM,
                       Spectra_Avg_Method=sel.get("Spectra_Avg_Method", "Spectra_Avg_All"),
                       Lines_Xticks_picked=sl, CurTcvModelParaOpmScheme=tcv,
                       PLSfactor_Opm_User_Pick=user_pick, AQP_class=aqp_class,
                       pathfname_PPn_xlsx=PPn_xlsx, workdir=workdir)


# =============================================================================
#   COMPARE RESULTS  (Cmp_Results_AQP and its CmpR_* helpers)
# =============================================================================
def _distinct_tails(c0: List[str]) -> List[str]:
    c = list(c0)
    if len(c) < 2 or all(x == c[0] for x in c):
        return c
    nmin = min(len(x) for x in c)
    nh = 0
    for k in range(nmin):
        if all(x[k] == c[0][k] for x in c):
            nh = k + 1
        else:
            break
    nt = 0
    for k in range(1, nmin - nh + 1):
        if all(x[-k] == c[0][-k] for x in c):
            nt = k
        else:
            break
    out = []
    for i, x in enumerate(c):
        s = x[nh:len(x) - nt] if nt else x[nh:]
        s = re.sub(r"^[_\-+ ]+", "", s)
        s = re.sub(r"[_\-+ ]+$", "", s)
        out.append(s if s else c0[i])
    return out


def _join_unique(c: List[str]) -> str:
    return " ".join(x for x in unique_appear_order(c) if x)


def cmp_build_series(files: List[str], split_by_CabXfer=True):
    S = []
    for ifile, pfn in enumerate(files):
        try:
            O = read_results_mat(pfn)
        except Exception as e:
            warn(f"could not load --> {pfn} ({e})")
            continue
        cXT = [str(o.get("cLegendLoop_jLine", "")) for o in O]
        v = []
        for o in O:
            x = o.get("Results4sRMSE_Val", np.nan)
            try:
                v.append(float(x))
            except (TypeError, ValueError):
                v.append(float("nan"))
        cXf = [str(o.get("CabXfer_scheme", "")) for o in O]
        fname = os.path.basename(pfn)
        sSA = _join_unique([str(o.get("Spectra_Avg_Method", "")) for o in O]).replace("Spectra_Avg_", "")
        if not sSA:
            k = fname.rfind("Spectra_Avg_")
            if k >= 0:
                sSA = fname[k + len("Spectra_Avg_"):]
                p = sSA.find("_{")
                if p >= 0:
                    sSA = sSA[:p]
                sSA = sSA.replace(".mat", "")
        sAna = _join_unique([str(o.get("AnaName", "")) for o in O])
        uX = unique_appear_order(cXf)
        groups = uX if (split_by_CabXfer and len(uX) > 1) else [""]
        for g in groups:
            idx = list(range(len(cXT))) if g == "" else [i for i, x in enumerate(cXf) if x == g]
            cx, vv, nd = [], [], 0
            for i in idx:
                if cXT[i] in cx:
                    nd += 1
                    continue
                cx.append(cXT[i]); vv.append(v[i])
            if nd:
                warn(f"{nd} duplicated x-axis entries in {fname} (kept the first of each)")
            m = re.findall(r"\{[^{}]*-PP[12]\}", fname)
            core = fname.replace(".mat", "")
            k = core.find("Results_AQP_")
            if k >= 0:
                core = core[k + len("Results_AQP_"):]
            S.append({"pathfname": pfn, "fname": fname, "ifile": ifile, "cXT": cx, "RMSE": vv,
                      "sXfer": _join_unique([cXf[i] for i in idx]), "sSA": sSA, "sAna": sAna,
                      "sPP": m[-1] if m else "", "sCore": core, "ndup": nd})
    return S


def cmp_build_legend(S, detailed: bool = True) -> List[str]:
    """CmpR_build_legend.  detailed=False -> exactly the MATLAB rule (only what
    differs between the series, so a single-CabXfer comparison shows just '11').
    detailed=True (default, 20 Oct 2026) -> every entry always carries
    'CabXfer -->X' and the distinct part of its file name, e.g.
    'CabXfer -->MDC  |  11[PO2]_Xfer-PP_MDC'."""
    def varies(key):
        c = [s[key] for s in S]
        return not all(x == c[0] for x in c) and not all(x == "" for x in c)
    vAn, vSA, vXf, vPP = varies("sAna"), varies("sSA"), varies("sXfer"), varies("sPP")
    cCo = _distinct_tails([s["sCore"] for s in S])
    if detailed and len(S) > 1:
        # cut only the common HEAD (keep the tail, e.g. '_Xfer-PP_MDC') and the
        # Spectra_Avg token, which the title already carries
        c0 = [re.sub(r"_?Spectra_Avg_[A-Za-z]+", "", s["sCore"]) for s in S]
        nh = 0
        for k in range(min(len(x) for x in c0)):
            if all(x[k] == c0[0][k] for x in c0):
                nh = k + 1
            else:
                break
        cCo = [re.sub(r"[_\-+ ]+$", "", re.sub(r"^[_\-+ ]+", "", x[nh:])) or cCo[i]
               for i, x in enumerate(c0)]
    leg = []
    for i, s in enumerate(S):
        p = []
        if vAn and s["sAna"]:
            p.append("Analyte -->" + s["sAna"])
        if vSA and s["sSA"]:
            p.append("Spectra -->" + s["sSA"])
        if (vXf or detailed) and s["sXfer"]:
            p.append("CabXfer -->" + s["sXfer"])
        if vPP and s["sPP"]:
            p.append("PP -->" + s["sPP"])
        if detailed or not p:
            p.append(cCo[i])
        leg.append("  |  ".join(p))
    if len(set(leg)) < len(S):
        leg = [l + "  |  " + (s["sCore"] if detailed else cCo[i]) for i, (l, s) in enumerate(zip(leg, S))]
    return leg


def cmp_common_title(S) -> str:
    p = []
    uAn = unique_appear_order([s["sAna"] for s in S if s["sAna"]])
    if len(uAn) == 1:
        p.append("[" + uAn[0] + "]")
    elif len(uAn) > 1:
        p.append("Analyte --> " + " vs ".join(uAn) + "  (RMSEP scales differ !)")
    uX = unique_appear_order([s["sXfer"] for s in S if s["sXfer"]])
    if len(uX) == 1:
        p.append(uX[0])
    elif len(uX) > 1:
        p.append("CabXfer --> " + " vs ".join(uX) + "  (see legend)")
    uS = unique_appear_order([s["sSA"] for s in S if s["sSA"]])
    if len(uS) == 1:
        p.append("Spectra_Avg_" + uS[0])
    return "    ".join(p)


def enable_datatips(fig, lines, pickradius: float = 7):
    """MATLAB-style data tips on a matplotlib figure.

    Left-click a marker -> a box 'X <x>' / 'Y <y>' stays pinned to it (several
    may be open at once, as in MATLAB).  Click the same marker again to remove
    its tip; Delete / Escape removes them all.  Ignored while the toolbar's
    pan or zoom mode is on.  The figures here are bare Figure objects whose Tk
    canvas is made later, so the wiring is queued in fig._aqp_on_canvas and run
    by _show_figure once the canvas exists (PNG/CLI output: nothing happens)."""
    lines = [ln for ln in lines if ln is not None]
    if fig is None or not lines:
        return
    for ln in lines:
        ln.set_picker(True)
        ln.set_pickradius(pickradius)

    def _fmt(v):
        v = float(v)
        if v == int(v) and abs(v) < 1e15:
            return str(int(v))
        return f"{v:.4g}"                        # MATLAB datatip shows 4 significant digits

    def connect(canvas):
        tips = {}                                # (line id, point index) -> (annotation, dot)

        def _busy():
            tb = getattr(canvas, "toolbar", None)
            return bool(tb is not None and getattr(tb, "mode", ""))

        def on_pick(ev):
            ln = ev.artist
            if ln not in lines or _busy() or ev.mouseevent.button != 1 or not len(ev.ind):
                return
            ax = ln.axes
            xd, yd = np.asarray(ln.get_xdata(), float), np.asarray(ln.get_ydata(), float)
            # several points inside the pick radius -> the one nearest the click (pixels)
            pts = ax.transData.transform(np.column_stack([xd[ev.ind], yd[ev.ind]]))
            me = ev.mouseevent
            i = int(ev.ind[int(np.argmin(np.hypot(pts[:, 0] - me.x, pts[:, 1] - me.y)))])
            key = (id(ln), i)
            if key in tips:                      # second click on the same point -> remove
                for a in tips.pop(key):
                    a.remove()
                canvas.draw_idle()
                return
            x, y = xd[i], yd[i]
            # keep the box inside the axes: flip left near the right edge, down near the top
            fx, fy = ax.transAxes.inverted().transform(ax.transData.transform((x, y)))
            dx, ha = (-14, "right") if fx > 0.72 else (14, "left")
            dy, va = (-14, "top") if fy > 0.80 else (14, "bottom")
            dot, = ax.plot([x], [y], "o", ms=5, mfc="k", mec="w", mew=1, zorder=10)
            ann = ax.annotate(f"X  {_fmt(x)}\nY  {_fmt(y)}", xy=(x, y), xytext=(dx, dy),
                              textcoords="offset points", ha=ha, va=va, fontsize=11,
                              fontweight="bold", color="#0B6FD6", zorder=11, annotation_clip=False,
                              bbox=dict(boxstyle="round,pad=0.45", fc="#F8F8FB", ec="#8C8C8C", lw=1),
                              arrowprops=dict(arrowstyle="-", color="#8C8C8C", lw=0.8))
            tips[key] = (ann, dot)
            canvas.draw_idle()

        def on_key(ev):
            if ev.key in ("delete", "escape", "backspace") and tips:
                for pair in tips.values():
                    for a in pair:
                        a.remove()
                tips.clear()
                canvas.draw_idle()

        canvas.mpl_connect("pick_event", on_pick)
        canvas.mpl_connect("key_press_event", on_key)

    hooks = getattr(fig, "_aqp_on_canvas", None) or []
    hooks.append(connect)
    fig._aqp_on_canvas = hooks


def cmp_results(folder: str, figs: FigMgr, XTick_order="sortnat", strict_XTick=False,
                connect_lines=False, wrap=True, markersize=12):
    """Cmp_Results_AQP: RMSEP of every Results_AQP*.mat in one folder, one series
    per (file x CabXfer scheme)."""
    files = sorted([os.path.join(folder, f) for f in os.listdir(folder)
                    if "Results_AQP" in f and f.endswith(".mat")], key=lambda x: os.path.basename(x).upper())
    if len(files) < 2:
        raise AQPError(f'Need at least two "Results_AQP~.mat" files in this folder --> {folder}')
    S = cmp_build_series(files)
    if len(S) < 2:
        raise AQPError(f"Could not read two usable series out of --> {folder}")
    # [LEGEND ORDER, 20 Oct 2026] series (hence legend, colours and the [i]
    # numbering of the gap report) in sortnat order of their legend text
    leg = cmp_build_legend(S)
    isrt = sorted(range(len(S)), key=lambda i: _natkey(leg[i]))
    S, leg = [S[i] for i in isrt], [leg[i] for i in isrt]
    same = all(s["cXT"] == S[0]["cXT"] for s in S)
    if same:
        Q = list(S[0]["cXT"])
    else:
        Q = unique_appear_order([x for s in S for x in s["cXT"]])
    # [XTICK ORDER, 20 Oct 2026] applied whether or not the files agree.
    # default 'sortnat' -> SGFL5, 7, 9, 11, 13 (plain sort gives 11, 13, 5, 7, 9)
    if XTick_order == "sortnat":
        Q = sortnat(Q)
    elif XTick_order == "sorted":
        Q = sorted(Q)
    if not same:
        lines = ["x-axis (cLegendLoop_jLine) entries are NOT identical across the compared files.",
                 f"Union used for the x axis has {len(Q)} entries."]
        for i, s in enumerate(S):
            miss = [q for q in Q if q not in s["cXT"]]
            lines.append(f"  [{i + 1}] {s['fname']}  -->  " + ("complete" if not miss else
                         f"MISSING {len(miss)} : " + ", ".join(miss)))
        msg = "\n".join(lines)
        if strict_XTick:
            raise AQPError(msg)
        note(msg)
    f = figs.new("Compare_Results", size=(11, 6))
    if f is None:
        return {"S": S, "Q": Q, "legend": leg}
    ax = f.add_subplot(111)
    ax.grid(True)
    cols = ["b", "r", "k", "m", (0, 0.6, 0), "c", (0.85, 0.33, 0.10), (0.49, 0.18, 0.56)]
    mks = ["*", "o", "s", "d", "^", "v", "p", "h"]
    hl = []
    for k, s in enumerate(S):
        col = cols[k % len(cols)]
        mk = mks[(k // len(cols)) % len(mks)]
        xs = [i + 1 for i, q in enumerate(Q) if q in s["cXT"]]
        ys = [s["RMSE"][s["cXT"].index(q)] for q in Q if q in s["cXT"]]
        hl += ax.plot(xs, ys, ls="-" if connect_lines else "none", marker=mk, color=col, ms=markersize, lw=2)
    ax.legend(leg, loc="best", fontsize=8)
    enable_datatips(f, hl)          # MATLAB-style data tips (GUI window only)
    labs = []
    for q in Q:
        s = q.replace("{PRO}", "\n               MN-PRO").replace("+", "\n") if wrap else q
        labs.append(s)
    ax.set_xticks(range(1, len(Q) + 1))
    ax.set_xticklabels(labs, rotation=-45, fontsize=8, ha="left")
    ax.set_xlim(0.5, len(Q) + 0.5)
    ax.set_ylabel("RMSEP")
    m = re.findall(r"\{[^{}]*\}", os.path.basename(files[0]))
    stit1 = m[0] if m else ""
    title = [t for t in [stit1, os.path.basename(folder), cmp_common_title(S)] if t]
    if not same:
        title.append(f"NOTE: x-axis entries differ between files -> union of {len(Q)} schemes shown")
    ax.set_title("\n".join(title), fontsize=9)
    figs.show(f)
    return {"S": S, "Q": Q, "legend": leg, "fig": f}


# =============================================================================
#   STEP BY STEP PLOTS  (StepbyStep_plots_AQP, Calc_Qres_T2_AQPlite, ...)
# =============================================================================
def _load_at(pfn):
    import scipy.io as sio
    return sio.loadmat(pfn, squeeze_me=True, struct_as_record=False)


def _hist_matlab(x, nbins=20):
    """[N,X] = hist(x,20): bin centres and counts like MATLAB's hist."""
    x = np.ravel(x)
    x = x[np.isfinite(x)]
    if x.size == 0:
        return np.zeros(nbins), np.arange(nbins, dtype=float)
    lo, hi = float(np.min(x)), float(np.max(x))
    if lo == hi:
        lo, hi = lo - 0.5, hi + 0.5
    w = (hi - lo) / nbins
    centers = lo + w * (np.arange(nbins) + 0.5)
    edges = np.concatenate([[-np.inf], (centers[:-1] + centers[1:]) / 2, [np.inf]])
    N, _ = np.histogram(x, edges)
    return N.astype(float), centers


def calc_Qres_T2(LAT, opm: int):
    """Calc_Qres_T2_AQPlite: Q residuals and T2 (sum of squared scores) at opm."""
    XT = np.atleast_2d(LAT["Atrainpk"])
    try:
        sc = np.atleast_1d(LAT["PLS"].Tset.saConc)
        YT = np.array([float(s.Conc) for s in sc])
    except Exception:
        YT = None
    Xl, _, Xs, _, _, _ = plsregress(XT, YT, opm)
    meanXT = XT.mean(axis=0)
    XP = np.atleast_2d(LAT["Apred"])
    XPs = (np.linalg.inv(Xl.T @ Xl) @ Xl.T @ (XP - meanXT).T).T
    XPr = XPs @ Xl.T + meanXT
    XTr = Xs @ Xl.T + meanXT
    return {"ss_XT_Residual": np.sum((XT - XTr) ** 2, axis=1), "ss_XP_Residual": np.sum((XP - XPr) ** 2, axis=1),
            "ss_XT_scores": np.sum(Xs ** 2, axis=1), "ss_XP_scores": np.sum(XPs ** 2, axis=1), "FalsePos_yes": 0}


def plot_LogStdScale(figs: FigMgr, q, which: str, sfname: str, sDS: str):
    Ns = 3
    if which == "Qres":
        t, p = q["ss_XT_Residual"], q["ss_XP_Residual"]
    else:
        t, p = q["ss_XT_scores"], q["ss_XP_scores"]
    lt, lp = np.log(t), np.log(p)
    med, sd = np.median(lt), np.std(lt, ddof=1)
    thr = med + sd * Ns
    Ps, Ts, thr_s = (lp - med) / sd, (lt - med) / sd, (thr - med) / sd
    NT, XT_ = _hist_matlab(Ts, 20); NT = NT / Ts.size * 100
    NP, XP_ = _hist_matlab(Ps, 20); NP = NP / Ps.size * 100
    f = figs.new(f"LogStdScale_{which}", size=(10.6, 5.3))
    if f is not None:
        ax = f.add_subplot(111)
        ax.stem(XT_, NT, linefmt="b-", markerfmt="b.", basefmt=" ")
        st_ = ax.stem(XP_, NP, linefmt="-", markerfmt="*", basefmt=" ")
        st_.stemlines.set_color(ML_G); st_.markerline.set_color(ML_G)
        over = XP_ > thr_s
        if over.any():
            ax.plot(XP_[over], NP[over], "o", color=ML_M, mfc=ML_M, ls="none")
        ax.axvline(thr_s, color="r")
        ax.legend(["CS or Tset", "Val or Pset"] + (["Over Threshold"] if over.any() else []) + ["Threshold"])
        ax.set_xlabel(f"{which} of Pset scaled by Tset Sigma"); ax.set_ylabel("percent of samples")
        n_over = int(np.sum(lp > thr))
        ax.set_title(f"{sfname}\n{sDS}\n{which}    Val Over Thredhold = {roundns(n_over / p.size * 100, 2)}%  "
                     f"({n_over}/{p.size})", fontsize=8)
        figs.show(f)
    return {"T_scaled": Ts, "P_scaled": Ps}


def stepbystep_plots(kind: str, out_SbS: str, path_DS: str, figs: FigMgr, cab_scheme: str = "",
                     results_folder: Optional[str] = None, results_fname: Optional[str] = None):
    """StepbyStep_plots_AQP.  Like the MATLAB code, a missing input only prints a
    message (no error)."""
    disp("work on plotting spectra in StepbyStep folder")
    if not out_SbS or not os.path.isdir(out_SbS):
        note("no StepByStep folder yet - press Run AQP first")
        return
    files = sorted(f for f in os.listdir(out_SbS) if f.startswith("Atrainpketc_") and f.endswith(".mat"))
    sDS = os.path.basename(os.path.normpath(path_DS)) if path_DS else ""
    if kind == "CS_XRSmst_Bef_MGs":
        fcs = [f for f in files if "CS-Before-MatchGrids" in f]
        if len(fcs) != 1:
            note("no CS-Before-MatchGrids file in the StepByStep folder (the dataset has neither XRS nor Val)")
            return
        L = _load_at(os.path.join(out_SbS, fcs[0]))
        wvl = np.ravel(L["wvl_standardize"])
        famg = [f for f in files if "_After-MatchGrids_" in f]
        wmn, typ = MN_WVL, "MN grids"
        if famg:
            w2 = np.ravel(_load_at(os.path.join(out_SbS, famg[0]))["wvl_standardize"])
            if not is_MN_wvl(w2):
                wmn, typ = w2, "non-MN grids"
        f = figs.new("CS_XRSmst_Bef_MGs", size=(8.4, 6.3))
        if f is None:
            return
        ax = f.add_subplot(111)
        h1 = ax.plot(wvl, np.atleast_2d(L["Atrainpk"]).T, "b-", lw=0.5)
        hs = [h1[0]]; lab = ["CS"]
        fxs = [f2 for f2 in files if "XSmst-Before-MatchGrids" in f2]
        if len(fxs) == 1:
            Lx = _load_at(os.path.join(out_SbS, fxs[0]))
            h2 = ax.plot(wvl, np.atleast_2d(Lx["Atrainpk"]).T, "-", color=ML_C, lw=0.5)
            hs.append(h2[0]); lab.append("XRSmst")
        y0, y1 = ax.get_ylim()
        # plot_vline_stem: pure-yellow stems at MATLAB's default 0.5 pt, drawn behind
        # the spectra so every spectrum stays visible through the grid band
        # (a pale-yellow band underneath keeps the band even at any window size / DPI,
        #  where 125 sub-pixel lines alone would beat into stripes)
        ax.axvspan(float(np.min(wmn)), float(np.max(wmn)), ymin=0, ymax=1, color=(1, 1, 0.62), lw=0, zorder=0.9)
        h3 = ax.vlines(wmn, y0 if y0 < 0 else 0, y1, colors=[ML_Y], linewidths=0.5, zorder=1)
        ax.set_ylim(y0, y1)
        hs.append(h3); lab.append(typ)
        ax.legend(hs, lab)
        ax.set_title(f"{sDS}\n{kind}"); ax.set_ylabel("Spectra"); ax.set_xlabel("wvl")
        figs.show(f)
    elif kind == "CS_Val_CabXfer_PP":
        fv = [f for f in files if "CS&Val_for-AAQP" in f]
        if len(fv) != 1:
            note("NO appropriate spectra of CS vs Val were found, hence NO plot will be generated")
            return
        L = _load_at(os.path.join(out_SbS, fv[0]))
        f = figs.new("CS_Val_CabXfer_PP", size=(8.4, 6.3))
        if f is not None:
            ax = f.add_subplot(111)
            h1 = ax.plot(np.atleast_2d(L["Atrainpk"]).T, "b-", lw=0.5)
            h2 = ax.plot(np.atleast_2d(L["Apred"]).T, "r-", lw=0.5)
            ax.legend([h1[0], h2[0]], ["CS", "Val"])
            ax.set_title(f"{sDS}  {cab_scheme}\n{kind}"); ax.set_ylabel("Spectra"); ax.set_xlabel("wvl")
            figs.show(f)
        results_file = os.path.join(results_folder or "", results_fname or "")
        if not (results_folder and results_fname and os.path.isfile(results_file)):
            cand = ([os.path.join(results_folder, f) for f in os.listdir(results_folder)
                     if f.startswith("Results_") and f.endswith(".mat")]
                    if results_folder and os.path.isdir(results_folder) else [])
            if len(cand) != 1:
                note("can  Not uniquely find Results_~.mat file,  hence can Not show Qres and T2 "
                     "for CS_Val_CabXfer_PP")
                return
            results_file = cand[0]
        O = read_results_mat(results_file)
        opm = int(float(O[0]["Results4sOpmPLSfactor"]))
        q = calc_Qres_T2(L, opm)
        sf = fv[0]
        a = plot_LogStdScale(figs, q, "Qres", sf, sDS)
        b = plot_LogStdScale(figs, q, "T2", sf, sDS)
        f = figs.new("T2_vs_Qres", size=(10.6, 6.9))
        if f is not None:
            ax = f.add_subplot(111)
            ax.plot(b["T_scaled"], a["T_scaled"], "bo", mfc="none", ls="none", label="CS or Tset")
            ax.plot(b["P_scaled"], a["P_scaled"], "*", color=ML_G, ls="none", label="Val or Pset")
            ax.axvline(3, color="r"); ax.axhline(3, color="r")
            ax.legend()
            ax.set_xlabel("Hotelling T-square (scaled by StdLog)"); ax.set_ylabel("Q Residuals (scaled by StdLog)")
            ax.set_title(f"{sf}\n{sDS}   OpmPLSfactor = {opm}", fontsize=8)
            figs.show(f)
    else:
        raise AQPError("SbS_Spectra_type Not supported")


# =============================================================================
#   SCAN THRU  (ScanThru_* : press Run AQP once per combination of menu options)
#   GUI-independent core, used by the panel and by --cli --scan
# =============================================================================
SCANTHRU_SETTINGS = {
    "menu_String": ["Scan Thru ...", "Set up & run scan", "Re-run last scan",
                    "Show last summary", "Collect results"],
    # popup menus that are never offered as scan parameters (action menus)
    "exclude_tags": ["ScanThru", "Load_DS", "StepbyStep_plots"],
    # the order the parameters are listed in (anything else follows)
    "order_tags": ["Spectra_Avg_Method", "CabXfer_scheme", "UDMas_scheme", "cPP1", "cPP2",
                   "TcvModelParaOpmScheme", "Lines_Xticks", "cDataFlow", "AQP_class"],
    # menu entries dropped from every option list (they open dialogs of their own)
    "blocked_option_kw": ["ChooseFrom", "User-Pick"],
    # 0 -> the one-prompt-per-sweep answers are forgotten when a sweep starts/ends
    "cache_across_sweeps": 0,
    "max_combinations": 64,          # ask for confirmation beyond this many runs
    "close_figs_between_runs": 1,    # initial state of the picker's checkbox
    "stop_on_error": 0,              # 1 = abandon the sweep on the first failure
    "consolidate_policy": "ask",     # 'ask' | 'collect' | 'largest' | 'leave'
}


def is_dash_line(t: str) -> bool:
    t = t.strip()
    return len(t) >= 3 and set(t) == {"-"}


def scan_is_title_row(t: str, tag: str) -> bool:
    tag2 = tag[1:] if tag.startswith("c") else tag
    t2 = t.replace("--", "_")
    return any(a.lower() == b.lower() for a in (t, t2) for b in (tag, tag2))


def scan_clean_options(tag: str, menu: List[str]):
    """ScanThru_clean_options: the real, pickable entries and their menu rows."""
    copt, vidx = [], []
    for i, t in enumerate(menu):
        t = t.strip()
        if (not t or is_dash_line(t) or scan_is_title_row(t, tag)
                or any(k in t for k in SCANTHRU_SETTINGS["blocked_option_kw"]) or t in copt):
            continue
        copt.append(t)
        vidx.append(i)
    return copt, vidx


def scan_registry(menus: Dict[str, List[str]], values: Dict[str, int]) -> List[Dict[str, Any]]:
    """ScanThru_registry: the scannable popup menus.  values = {tag: current row}."""
    R = []
    for tag, menu in menus.items():
        if tag in SCANTHRU_SETTINGS["exclude_tags"]:
            continue
        copt, vidx = scan_clean_options(tag, menu)
        if len(copt) < 2:
            continue
        v = values.get(tag, 0)
        if v in vidx:
            cur, resolved = vidx.index(v), 0
        else:
            cur, resolved = 0, 1
        R.append({"tag": tag, "options": copt, "vidx": vidx, "cur": cur, "resolved": resolved})
    order = SCANTHRU_SETTINGS["order_tags"]
    rank = [order.index(r["tag"]) if r["tag"] in order else 1000 + i for i, r in enumerate(R)]
    return [r for _, r in sorted(zip(rank, R), key=lambda x: x[0])]


def scan_combinations(P: List[Dict[str, Any]]) -> List[Tuple[int, ...]]:
    """the cartesian product; the LAST parameter varies fastest (odometer)."""
    cnt = [max(1, len(p["sel"])) for p in P]
    if not cnt:
        return []
    return list(itertools.product(*[range(c) for c in cnt]))


def scan_combo_label(P, row) -> str:
    c = [f"{p['tag']} = {p['options'][p['sel'][row[i]]]}" for i, p in enumerate(P) if len(p["sel"]) > 1]
    return "  |  ".join(c) if c else "(single run -- nothing varied)"


def scan_header(P, N) -> str:
    L = ["", "=" * 71, f"  SCAN THRU  --  {N} run(s)   {_dt.datetime.now():%d-%b-%Y %H:%M:%S}", "-" * 71]
    nres = 0
    for p in P:
        k = len(p["sel"])
        sres = " *" if p["resolved"] else ""
        nres += bool(p["resolved"])
        L.append(f"  {'SCAN ' if k > 1 else 'fixed'}  {p['tag']:<24s} {k} option(s){sres}")
        L += [f"            {p['options'][j]}" for j in p["sel"]]
    if nres:
        L += ["-" * 71, "  * the panel is parked on that menu's title row -- Scan Thru uses the",
              "    option printed above it instead, which is what the picker showed."]
    L.append("=" * 71)
    return "\n".join(L)


def scan_count_results(folder: Optional[str]) -> int:
    if not folder or not os.path.isdir(folder):
        return 0
    return sum(1 for f in os.listdir(folder)
               if "Results_AQP" in f and f.endswith(".mat") and os.path.isfile(os.path.join(folder, f)))


def scan_distinct_folders(LOG) -> List[str]:
    cf = []
    for f in (LOG or {}).get("folder", []):
        if f and f not in cf:
            cf.append(f)
    return cf


def scan_sweep(P, run_combo, status=None, after_run=None, stop_on_error=None) -> Dict[str, Any]:
    """ScanThru_sweep.  run_combo(row) applies the combination and runs AQP; it
    returns (folder, fname) of the comparison folder in use after the run, and
    raises on failure.  A failing run is logged and the sweep carries on."""
    if stop_on_error is None:
        stop_on_error = SCANTHRU_SETTINGS["stop_on_error"]
    Cidx = scan_combinations(P)
    N = len(Cidx)
    LOG = {"spec": P, "Cidx": Cidx, "when": _dt.datetime.now().strftime("%d-%b-%Y %H:%M:%S"),
           "nrun": N, "label": [""] * N, "status": ["not run"] * N, "folder": [""] * N,
           "file": [""] * N, "msg": [""] * N}
    print(scan_header(P, N), flush=True)
    nDone = N
    for r, row in enumerate(Cidx):
        slab = scan_combo_label(P, row)
        LOG["label"][r] = slab
        if status:
            status(f"run {r + 1}/{N}", f"Scan Thru: run {r + 1} of {N} -- {slab}")
        print(f"\n[Scan Thru {r + 1}/{N}]  {slab}", flush=True)
        LOG["status"][r] = "ok"
        folder = fname = ""
        try:
            folder, fname = run_combo(row)
        except Exception as e:                  # noqa: BLE001 - one bad combination must not stop the rest
            LOG["status"][r] = "FAILED"
            LOG["msg"][r] = str(e)
            traceback.print_exc()
            folder, fname = getattr(e, "_aqp_folder", ("", ""))
        LOG["folder"][r] = folder or ""
        LOG["file"][r] = fname or ""
        if after_run:
            after_run()
        if LOG["status"][r] == "FAILED" and stop_on_error:
            print(f"[Scan Thru] stop_on_error is set -- the sweep stopped at run {r + 1} of {N}")
            nDone = r + 1
            break
    LOG["nrun"] = nDone
    return LOG


def scan_unique_target(dest: str, name: str) -> str:
    base, ext = os.path.splitext(name)
    t = os.path.join(dest, name)
    k = 1
    while os.path.isfile(t) and k <= 99:
        k += 1
        t = os.path.join(dest, f"{base}_{k}{ext}")
    return t


def scan_collect(cf: List[str]) -> str:
    """ScanThru_collect: COPY every Results_AQP*.mat into <parent>/ScanThru_<stamp>."""
    if not cf:
        return ""
    parent = os.path.dirname(cf[0])
    if not parent:
        return ""
    dest = os.path.join(parent, "ScanThru_" + _dt.datetime.now().strftime("%Y%m%d_%H%M%S"))
    try:
        os.makedirs(dest)
    except OSError:
        return ""
    for f in cf:
        if not os.path.isdir(f):
            continue
        for n in sorted(os.listdir(f)):
            if "Results_AQP" in n and n.endswith(".mat") and os.path.isfile(os.path.join(f, n)):
                try:
                    shutil.copy2(os.path.join(f, n), scan_unique_target(dest, n))
                except OSError:
                    print(f"[Scan Thru] could not copy {n}")
    return dest


def scan_summary(LOG, cmp_folder: Optional[str]) -> str:
    if not LOG:
        return "[Scan Thru] no scan has been run in this session yet"
    n = LOG["nrun"]
    L = ["", "=" * 71, f"  SCAN THRU SUMMARY   (started {LOG['when']})", "-" * 71]
    nok = 0
    for r in range(n):
        nok += LOG["status"][r] == "ok"
        L.append(f"  [{r + 1:2d}/{n:2d}] {LOG['status'][r]:<8s} {LOG['label'][r]}")
        if LOG["file"][r]:
            L.append(f"           -> {LOG['file'][r]}")
        if LOG["msg"][r]:
            L.append(f"           !! {LOG['msg'][r]}")
    L += ["-" * 71, f"  {nok} of {n} run(s) completed"]
    for f in scan_distinct_folders(LOG):
        L += [f"  {scan_count_results(f)} Results_AQP_~.mat in", f"     {f}"]
    if cmp_folder:
        L += ["  Compare Results is pointing at", f"     {cmp_folder}",
              '  Press "Compare Results" to plot the sweep.']
    L.append("=" * 71)
    return "\n".join(L)


def list_results_text(folder: Optional[str]) -> Tuple[str, int]:
    """CmpR_list_Results_files: what is in the comparison folder right now."""
    L = ["=" * 71]
    n = 0
    if not folder:
        L += ["Results_AQP_~.mat  -->  no results folder yet (press Run AQP first)", "=" * 71]
        return "\n".join(L), 0
    L += [f"Results_AQP_~.mat  -->  {folder}", "-" * 71]
    if not os.path.isdir(folder):
        L += ["   ( that folder does not exist any more )", "=" * 71]
        return "\n".join(L), 0
    for f in sorted(os.listdir(folder)):
        p = os.path.join(folder, f)
        if not ("Results_AQP" in f and f.endswith(".mat") and os.path.isfile(p)):
            continue
        n += 1
        try:
            nx = len(unique_appear_order([str(o.get("cLegendLoop_jLine", "")) for o in read_results_mat(p)]))
        except Exception:
            nx = 0
        stt = os.stat(p)
        L += [f"  [{n}] {f}",
              f"        {round(stt.st_size / 1024)} kB   "
              f"{_dt.datetime.fromtimestamp(stt.st_mtime):%d-%b-%Y %H:%M:%S}   x-axis entries: {nx}"]
    if n == 0:
        L.append("  ( none )")
    L += ["-" * 71, f"  # results = {n}", "=" * 71]
    return "\n".join(L), n


def clear_results(folder: Optional[str]) -> int:
    """CmpR_clear_Results_files (the confirmation is asked by the caller)."""
    if not folder or not os.path.isdir(folder):
        return 0
    print("=" * 71 + f"\n[clear results] deleting from --> {folder}")
    for f in sorted(os.listdir(folder)):
        p = os.path.join(folder, f)
        if "Results_AQP" in f and f.endswith(".mat") and os.path.isfile(p):
            try:
                os.remove(p)
                print(f"   deleted  {f}")
            except OSError as e:
                print(f"   could not delete {f} ({e})")
    n = scan_count_results(folder)
    print(f"[clear results] # results = {n}\n" + "=" * 71)
    return n


def menus_for_session() -> Dict[str, List[str]]:
    """the panel menus; like AQP_gui_OutputFcn, the last Spectra_Avg_Method entry
    (Spectra_Avg_T-Mean_P-All) is only offered to the administrator account."""
    m = copy.deepcopy(GUI_MENUS)
    if os.environ.get("USERNAME", os.environ.get("USER", "")) != "hsi60346":
        if len(m["Spectra_Avg_Method"]) == 3:
            m["Spectra_Avg_Method"] = m["Spectra_Avg_Method"][:-1]
    return m


# =============================================================================
#   COMMAND LINE  (python AQP_gui.py --cli --ds <folder> ...)
# =============================================================================
_CLI_AVG = {"All": "Spectra_Avg_All", "Mean": "Spectra_Avg_Mean", "T-Mean_P-All": "Spectra_Avg_T-Mean_P-All"}
_CLI_OPM = {"KneePt+1": "KneePt+1_RMSECV (default)", "KneePt": "KneePt-RMSECV", "Min": "Min-RMSECV"}


def _split_list(s: Optional[str]) -> List[str]:
    return [x.strip() for x in (s or "").split(",") if x.strip()]


def build_arg_parser() -> argparse.ArgumentParser:
    ap = argparse.ArgumentParser(
        prog="AQP_gui.py",
        description="AutoQuant Data-Analytic Pipeline (Python port of AQP_gui_merged.m). "
                    "Without --cli the GUI opens.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="examples:\n"
               "  python AQP_gui.py\n"
               "  python AQP_gui.py --cli --ds data/DS1 --ana Brix\n"
               "  python AQP_gui.py --cli --ds data/DS1 --ana Brix --pp1 SNV --pp2 1stDerSGFL7[PO2] --cabxfer PDS99[w5]\n"
               "  python AQP_gui.py --cli --ds data/DS1 --ana Brix --pp-xlsx pp1_pp2_4schemes.xlsx --cabxfer MDC,PDS99[w5]\n"
               "  python AQP_gui.py --cli --ds data/DS1 --ana Brix --scan cPP2=SNV,none --scan CabXfer_scheme=MDC,PDS99[w1] --compare\n")
    ap.add_argument("--cli", action="store_true", help="run without the GUI")
    ap.add_argument("-d", type=int, default=None, metavar="0|1", help="print progress (default: 1 in --cli, 0 in GUI)")
    ap.add_argument("-w", type=int, default=None, metavar="0|1", help="print warnings (default 0)")
    ap.add_argument("--workdir", default=None, help="where TMP_AQP / Result4* folders go (default: current folder)")
    g = ap.add_argument_group("run settings (--cli)")
    g.add_argument("--ds", help="dataset folder holding CS_*.xlsx [XRS_*.xlsx] [VAL_*.xlsx]")
    g.add_argument("--ana", help="analyte column of the CS sheet (asked / first one when omitted)")
    g.add_argument("--pp1", default="1stDerSGFL7[PO2]", help="PP1 (comma list allowed); default 1stDerSGFL7[PO2]")
    g.add_argument("--pp2", default="SNV", help="PP2 (comma list allowed); default SNV")
    g.add_argument("--pp-xlsx", help="PP1_PP2*.xlsx with PP1 / PP2 columns (same as picking PP1_PP2_xlsx)")
    g.add_argument("--cabxfer", default="MDC",
                   help="woCabXfer | MDC | PDS99[w1] | PDS99[w5] | PDS99[w9] (comma list = one x tick each); default MDC")
    g.add_argument("--avg", default="All", choices=list(_CLI_AVG), help="Spectra_Avg_Method (default All)")
    g.add_argument("--opm", default="KneePt+1", choices=list(_CLI_OPM), help="optimal PLS factor rule")
    g.add_argument("--class", dest="aqp_class", default="pro", choices=["pro", "lite"], help="AQP_class")
    g.add_argument("--figs", default="png", choices=["png", "none"],
                   help="png: save every figure under <workdir>/Figs_AQP_cli (default); none: no figures")
    g.add_argument("--scan", action="append", default=[], metavar="TAG=opt1,opt2",
                   help="Scan Thru: sweep a panel menu (tags: " + ", ".join(SCANTHRU_SETTINGS["order_tags"]) +
                        "); repeat for several menus")
    g.add_argument("--compare", nargs="?", const="__last__", default=None, metavar="FOLDER",
                   help="Compare Results (RMSEP plot) of a results folder; no value = the folder of this run")
    g.add_argument("--stepbystep", choices=GUI_MENUS["StepbyStep_plots"][1:], action="append", default=[],
                   help="StepbyStep plots after the run")
    return ap


def _cli_sel(args) -> Dict[str, str]:
    return {"AQP_class": args.aqp_class,
            "cPP1": "PP1_PP2_xlsx" if args.pp_xlsx else _split_list(args.pp1)[0],
            "cPP2": "PP1_PP2_xlsx" if args.pp_xlsx else _split_list(args.pp2)[0],
            "cDataFlow": "Xfer-PP", "UDMas_scheme": "woUDM",
            "CabXfer_scheme": _split_list(args.cabxfer)[0],
            "TcvModelParaOpmScheme": _CLI_OPM[args.opm],
            "Lines_Xticks": "Lines_PP1--Xticks_CabXfer",
            "Spectra_Avg_Method": _CLI_AVG[args.avg]}


def _cli_pick_analyte(args):
    def pick(analytes):
        if args.ana:
            if args.ana in analytes:
                return args.ana
            raise AQPError(f"analyte '{args.ana}' is not in the CS sheet; it has: {', '.join(analytes)}")
        if sys.stdin.isatty():
            print("pick the target analyte:")
            for i, a in enumerate(analytes, 1):
                print(f"  [{i}] {a}")
            k = input("number> ").strip()
            return analytes[int(k) - 1]
        note(f"no --ana given: using the first analyte '{analytes[0]}' (of {', '.join(analytes)})")
        return analytes[0]
    return pick


def cli_run_once(args, sel: Dict[str, str], figdir: Optional[str], lists=True) -> BatchResult:
    workdir = os.path.abspath(args.workdir or os.getcwd())
    st = settings_from_panel(sel, os.path.abspath(args.ds), workdir,
                             pick_PPn_xlsx=lambda: os.path.abspath(args.pp_xlsx) if args.pp_xlsx else None)
    if lists and not args.pp_xlsx:
        if len(_split_list(args.pp1)) > 1 and sel["cPP1"] == _split_list(args.pp1)[0]:
            st.cPP1 = _split_list(args.pp1)
        if len(_split_list(args.pp2)) > 1 and sel["cPP2"] == _split_list(args.pp2)[0]:
            st.cPP2 = _split_list(args.pp2)
    if lists and len(_split_list(args.cabxfer)) > 1 and sel["CabXfer_scheme"] == _split_list(args.cabxfer)[0]:
        st.cCabXfer_scheme = _split_list(args.cabxfer)
    if figdir:
        tmp_folder_rm_mk(os.path.basename(figdir), os.path.dirname(figdir))
    figs = FigMgr("file", figdir) if figdir else FigMgr("none")
    res = batch_run(st, figs, progress=lambda s: disp(s), pick_analyte=_cli_pick_analyte(args))
    res.figs_saved = list(figs.saved)
    return res


def _print_run(res: BatchResult):
    print("-" * 71)
    for o in res.OUT_cln:
        print(f"  {o['cLegendLoop_jLine']:<34s} {o['CabXfer_scheme']:<12s} "
              f"OpmPLSfactor = {o['Results4sOpmPLSfactor']:<3s} RMSEP = {o['Results4sRMSE_Val']}")
    if res.outOpmFM and res.outOpmFM.get("opm_file"):
        print(f"  optimal final model: {res.outOpmFM['opm_file']}")
    print(f"  results folder ({res.nfile_CR} Results_AQP file(s)): {res.Path4OUT_cln_AQP_FinalSubfolder}")
    print(f"  elapsed {seconds2human(res.elapsed)}")
    print("-" * 71, flush=True)


def cli_main(args) -> int:
    VERB.disp = 1 if args.d is None else args.d
    VERB.warn = 0 if args.w is None else args.w
    if not args.ds:
        print("--cli needs --ds <dataset folder>", file=sys.stderr)
        return 2
    if not os.path.isdir(args.ds):
        print(f"dataset folder not found: {args.ds}", file=sys.stderr)
        return 2
    workdir = os.path.abspath(args.workdir or os.getcwd())
    root = find_last_nonTMP_path(workdir)
    figroot = os.path.join(root, "Figs_AQP_cli") if args.figs == "png" else None
    sel = _cli_sel(args)
    last = None
    cmp_folder = None
    try:
        if args.scan:
            menus = copy.deepcopy(GUI_MENUS)
            rows = {t: (menus[t].index(v) if v in menus[t] else 0) for t, v in sel.items() if t in menus}
            R = {r["tag"]: r for r in scan_registry(menus, rows)}
            P = []
            for r in R.values():
                r = dict(r)
                r["sel"] = [r["cur"]]
                P.append(r)
            for spec in args.scan:
                if "=" not in spec:
                    raise AQPError(f"--scan wants TAG=opt1,opt2 (got {spec})")
                tag, opts = spec.split("=", 1)
                tag = tag.strip()
                if tag not in R:
                    raise AQPError(f"--scan: '{tag}' is not a scannable menu ({', '.join(R)})")
                p = next(q for q in P if q["tag"] == tag)
                want = _split_list(opts)
                if want == ["ALL"]:
                    p["sel"] = list(range(len(p["options"])))
                else:
                    bad = [w for w in want if w not in p["options"]]
                    if bad:
                        raise AQPError(f"--scan {tag}: unknown option(s) {bad}; options: {p['options']}")
                    p["sel"] = [p["options"].index(w) for w in want]
            N = len(scan_combinations(P))
            state = {"folder": "", "fname": ""}

            def run_combo(row):
                s2 = dict(sel)
                for i, p in enumerate(P):
                    s2[p["tag"]] = p["options"][p["sel"][row[i]]]
                k = run_combo.k = getattr(run_combo, "k", 0) + 1
                fd = os.path.join(figroot, f"ScanRun_{k:02d}") if figroot else None
                if fd:
                    os.makedirs(figroot, exist_ok=True)
                try:
                    res = cli_run_once(args, s2, fd, lists=False)
                except Exception as e:
                    e._aqp_folder = (state["folder"], state["fname"])
                    raise
                state["folder"], state["fname"] = res.Path4OUT_cln_AQP_FinalSubfolder, res.fname_Results_AQP_OUT_cln
                _print_run(res)
                return state["folder"], state["fname"]

            print(f"[Scan Thru] {N} run(s)")
            LOG = scan_sweep(P, run_combo)
            cmp_folder = state["folder"]
            cf = scan_distinct_folders(LOG)
            if len(cf) > 1:
                dest = scan_collect(cf)
                if dest:
                    cmp_folder = dest
                    print(f"[Scan Thru] collected {scan_count_results(dest)} file(s) into\n   {dest}")
            print(scan_summary(LOG, cmp_folder))
        else:
            last = cli_run_once(args, sel, figroot)
            _print_run(last)
            cmp_folder = last.Path4OUT_cln_AQP_FinalSubfolder
            figs = FigMgr("file", figroot) if figroot else FigMgr("none")
            figs.counter = len(getattr(last, "figs_saved", []))
            for kind in args.stepbystep:
                stepbystep_plots(kind, last.out_SbS, args.ds, figs, sel["CabXfer_scheme"],
                                 last.Path4OUT_cln_AQP_FinalSubfolder, last.fname_Results_AQP_OUT_cln)
        if args.compare is not None:
            folder = cmp_folder if args.compare == "__last__" else os.path.abspath(args.compare)
            figs = FigMgr("file", os.path.join(root, "Figs_AQP_cli")) if args.figs == "png" else FigMgr("none")
            figs.counter = 900
            out = cmp_results(folder, figs)
            print(f"Compare Results: {len(out['S'])} series over {len(out['Q'])} x-axis entries"
                  + (f" -> {figs.saved[-1]}" if figs.saved else ""))
        if figroot and os.path.isdir(figroot):
            print(f"figures: {figroot}")
        return 0
    except AQPError as e:
        print(f"AQP: {e}", file=sys.stderr)
        return 1


# =============================================================================
#   THE PANEL  (AQP_gui_merged.fig, rebuilt in tkinter)
# =============================================================================
class _Tooltip:
    def __init__(self, widget, text):
        self.w, self.text, self.tip = widget, text, None
        widget.bind("<Enter>", self._show, add="+")
        widget.bind("<Leave>", self._hide, add="+")

    def _show(self, _e=None):
        if self.tip or not self.text:
            return
        import tkinter as tk
        x = self.w.winfo_rootx() + 10
        y = self.w.winfo_rooty() + self.w.winfo_height() + 4
        self.tip = tk.Toplevel(self.w)
        self.tip.wm_overrideredirect(True)
        self.tip.wm_geometry(f"+{x}+{y}")
        tk.Label(self.tip, text=self.text, background="#ffffe0", relief="solid", borderwidth=1,
                 font=("TkDefaultFont", 9), justify="left").pack()

    def _hide(self, _e=None):
        if self.tip:
            self.tip.destroy()
            self.tip = None


class _PopupMenu:
    """A MATLAB (Windows) popupmenu: white box with a thin grey border, the
    string at the left, a chevron at the right; clicking opens the list right
    below it.  Offers the bits of the ttk.Combobox API the panel uses:
    current(), configure(state=...), cget('state'), bind(), place()."""
    BORDER, BORDER_HOT = "#ABADB3", "#0078D7"

    def __init__(self, master, values, font, height_px, bg="white", command=None):
        import tkinter as tk
        import tkinter.font as tkfont
        self.values, self.command, self._i, self._state, self.bg = list(values), command, 0, "readonly", bg
        self._tkfont, self._padded_to = tkfont.Font(font=font), 0
        self.frame = tk.Frame(master, bg=self.BORDER, bd=0, highlightthickness=0, height=height_px)
        self.frame.pack_propagate(False)
        inner = tk.Frame(self.frame, bg=bg, bd=0)
        inner.pack(fill="both", expand=True, padx=1, pady=1)
        aw = max(12, int(round(height_px * 0.62)))
        self.cv = tk.Canvas(inner, width=aw, bg=bg, highlightthickness=0, bd=0)
        self.cv.pack(side="right", fill="y")
        self.lab = tk.Label(inner, text=self.values[0], font=font, bg=bg, fg="black", anchor="w",
                            padx=max(2, height_px // 12), pady=0, bd=0)
        self.lab.pack(side="left", fill="both", expand=True)
        self.cv.bind("<Configure>", self._draw_chevron)
        self.menu = tk.Menu(master, tearoff=0, font=font, bg="white", fg="black", bd=1, relief="solid",
                            activebackground="#0078D7", activeforeground="white", activeborderwidth=0)
        for i, v in enumerate(self.values):
            self.menu.add_command(label=v, command=lambda i=i: self._pick(i))
        for w in (self.frame, inner, self.lab, self.cv):
            w.bind("<Button-1>", self._open)
        self.frame.bind("<Enter>", lambda e: self._hot(True), add="+")
        self.frame.bind("<Leave>", lambda e: self._hot(False), add="+")

    def _draw_chevron(self, _e=None):
        c = self.cv
        c.delete("all")
        w, h = c.winfo_width(), c.winfo_height()
        d = max(3, h // 7)
        cx, cy = w / 2 - 1, h / 2
        col = "#A0A0A0" if self._state == "disabled" else "#505050"
        c.create_line(cx - d, cy - d / 2, cx, cy + d / 2, cx + d, cy - d / 2, fill=col, width=1)

    def _hot(self, on):
        if self._state != "disabled":
            self.frame.configure(bg=self.BORDER_HOT if on else self.BORDER)

    def _open(self, _e=None):
        if self._state == "disabled":
            return
        x, y = self.frame.winfo_rootx(), self.frame.winfo_rooty() + self.frame.winfo_height()
        self._fit_width()
        self.menu.after(1, lambda: self.menu.activate(self._i))
        try:
            self.menu.tk_popup(x, y)
        finally:
            self.menu.grab_release()

    def _fit_width(self):
        """make the list as wide as the control, as MATLAB's is"""
        target = self.frame.winfo_width() - 28
        if target <= 0 or target == self._padded_to:
            return
        sp = max(1, self._tkfont.measure(" "))
        for i, v in enumerate(self.values):
            n = max(0, (target - self._tkfont.measure(v)) // sp)
            self.menu.entryconfigure(i, label=v + " " * n)
        self._padded_to = target

    def _pick(self, i):
        self.current(i)
        if self.command:
            self.command()

    def current(self, i=None):
        if i is None:
            return self._i
        self._i = int(i)
        self.lab.configure(text=self.values[self._i])
        return self._i

    def get(self):
        return self.values[self._i]

    def configure(self, **kw):
        st = kw.pop("state", None)
        if st is not None:
            self._state = "disabled" if st == "disabled" else "readonly"
            self.lab.configure(fg="#6D6D6D" if self._state == "disabled" else "black")
            self.frame.configure(bg="#CCCCCC" if self._state == "disabled" else self.BORDER)
            self._draw_chevron()
        if kw:
            self.lab.configure(**kw)

    config = configure

    def cget(self, key):
        return self._state if key == "state" else self.lab.cget(key)

    def __getattr__(self, name):          # place, bind, winfo_* ... go to the frame
        return getattr(self.frame, name)


class AQPApp:
    """The AQP panel, rebuilt from AQP_gui_merged.fig.

    Geometry is the .fig's, in MATLAB 'characters' units (figure 112 x 32.3).
    One MATLAB character unit on Windows is ~3.98 pt wide and ~10.43 pt high
    (measured from the MATLAB panel: 7.95 x 20.87 px at 150 % display scaling),
    and uicontrol FontSize is in points, so everything is laid out in points
    and scales with the display the same way MATLAB does.
    """
    CW_PT, CH_PT = 3.975, 10.43     # one MATLAB character unit, in points
    FW, FH = 112.0, 32.3            # figure size in characters (.fig Position)
    BG = "#F0F0F0"                  # figure Color [0.94 0.94 0.94]

    # tag: (Position [x y w h] in characters, FontSize, HorizontalAlignment, Visible)
    # straight from AQP_gui_merged.fig; the Compare/Show/Clear Results, Scan Thru
    # positions are the ones CmpR_layout_* / ScanThru_SETTINGS set at start-up.
    LAYOUT = {
        "Load_DS": ([9.25, 28.667, 50.0, 2.762], 12, "center", True),
        "Loaded_DS_folder": ([9.25, 26.476, 74.625, 1.762], 7, "left", True),
        "AQP_class": ([60.75, 29.810, 10.625, 1.381], 12, "center", False),
        "Spectra_Avg_Method": ([74.571, 29.0, 35.143, 1.278], 8, "center", True),
        "suserInitials": ([88.0, 31.278, 23.857, 0.944], 8, "center", True),
        "Val_Exist": ([88.0, 25.857, 18.875, 1.714], 8, "left", False),
        "cPP1": ([-0.125, 21.667, 43.0, 3.429], 12, "center", True),
        "Cur_PP1": ([-0.125, 21.333, 41.875, 1.667], 12, "center", True),
        "cPP2": ([46.75, 23.429, 41.75, 1.667], 12, "center", True),
        "Cur_PP2": ([46.75, 21.952, 38.0, 1.190], 10, "left", True),
        "cDataFlow": ([92.875, 23.048, 19.0, 2.048], 12, "center", False),
        "CabXfer_scheme": ([4.25, 14.0, 103.25, 3.952], 12, "center", True),
        "UDMas_scheme": ([71.0, 12.762, 28.375, 1.619], 8, "center", False),
        "TcvModelParaOpmScheme": ([70.875, 11.238, 29.125, 1.333], 8, "center", False),
        "ScanThru": ([3.0, 10.1, 46.0, 3.4], 18, "center", True),
        "ScanThruState": ([50.5, 10.3, 19.5, 3.0], 16, "left", True),
        "Lines_Xticks": ([37.375, 8.095, 37.625, 1.619], 8, "center", False),
        "Cmp_Results": ([78.0, 7.55, 31.38, 2.6], 12, "center", True),
        "ShowResults": ([63.0, 6.3, 14.5, 1.7], 9, "center", True),
        "ClearResults": ([63.0, 4.4, 14.5, 1.7], 9, "center", True),
        "sNumResults": ([78.0, 4.9, 31.38, 2.4], 16, "center", True),
        "CurStatus": ([39.0, 4.190, 20.5, 2.333], 15, "center", True),
        "Run_AQP": ([-0.125, 3.667, 37.625, 3.619], 12, "center", True),
        "StepbyStep_plots": ([37.375, 1.762, 37.75, 1.905], 12, "center", True),
        "close_figs": ([77.75, 0.095, 25.25, 3.048], 12, "center", True),
    }
    POPUPS = ["Load_DS", "AQP_class", "Spectra_Avg_Method", "cPP1", "cPP2", "cDataFlow", "CabXfer_scheme",
              "UDMas_scheme", "TcvModelParaOpmScheme", "ScanThru", "Lines_Xticks", "StepbyStep_plots"]

    def __init__(self, root, workdir: str):
        import tkinter as tk
        import tkinter.font as tkfont
        from tkinter import ttk
        global MSGBOX_HOOK
        self.tk, self.ttk, self.root = tk, ttk, root
        self.workdir = os.path.abspath(workdir)
        self.menus = menus_for_session()
        self.menus["ScanThru"] = list(SCANTHRU_SETTINGS["menu_String"])
        self.path_DS = self.path_XLSX = ""
        self.OUT_Cmp_Results: Optional[Dict[str, str]] = None
        self.out_SbS: Optional[str] = None
        self.last_result: Optional[BatchResult] = None
        self.scan_LOG = None
        self.cache = {"armed": False, "store": {}}
        self.fig_windows: List[Any] = []
        self.figmgr = FigMgr("gui", on_new=self._show_figure)
        self.busy = False
        # character unit -> pixels on this display (MATLAB scales with DPI the same way)
        ppt = self.ppt = root.winfo_fpixels("1p")
        self.CW, self.CH = self.CW_PT * ppt, self.CH_PT * ppt
        fams = set(tkfont.families(root))
        # MATLAB draws uicontrol text in a Helvetica/Arial-style face
        self.family = next((f for f in ("Arial", "Helvetica", "Liberation Sans", "Arimo") if f in fams),
                           tkfont.nametofont("TkDefaultFont").actual()["family"])
        root.title(os.path.splitext(os.path.basename(__file__))[0])
        root.geometry(f"{round(self.FW * self.CW)}x{round(self.FH * self.CH)}")
        root.resizable(False, False)            # .fig Resize 'off'
        root.configure(bg=self.BG)
        self.w: Dict[str, Any] = {}
        self.v: Dict[str, Any] = {}
        for tag in self.POPUPS:
            self._make_popup(tag)
        self._label("Loaded_DS_folder", "Loaded DS folder")
        self._label("suserInitials", "PowerUser")
        self._label("Cur_PP1", "Current Pretreatment")
        self._label("Cur_PP2", "Current PP2")
        self._label("CurStatus", "CurStatus")
        self._label("ScanThruState", "idle", bold=True)
        self._label("sNumResults", "# results = 0", bold=True)
        # Val_Exist: a hidden radiobutton in the .fig that the run sets (kept as state only)
        self.v["Val_Exist"] = tk.IntVar(value=0)
        self.w["Val_Exist"] = tk.Checkbutton(root, text="Val_Exist", variable=self.v["Val_Exist"],
                                             bg=self.BG, font=self._font(8))
        self._button("Run_AQP", "Run AQP", self.on_run_aqp)
        self._button("close_figs", "close figs", self.close_all_figs)
        self._button("Cmp_Results", "Compare Results", self.on_cmp_results)
        self._button("ShowResults", "Show Results", self.on_show_results)
        self._button("ClearResults", "Clear Results", self.on_clear_results, fg="#990000")
        _Tooltip(self.w["ShowResults"], "List the Results_AQP_~.mat files of the results folder in the command window")
        _Tooltip(self.w["ClearResults"], "DELETE every Results_AQP_~.mat in the results folder")
        _Tooltip(self.w["ScanThru"], "Pick hyperparameters and their options, then run every combination")
        self.update_cmp_button(0)
        MSGBOX_HOOK = self._msgbox
        root.protocol("WM_DELETE_WINDOW", self._quit)

    # ---------------------------------------------------------------- layout
    def _font(self, fs, bold=False):
        """MATLAB FontSize is in points; Tk positive sizes are points too."""
        return (self.family, int(fs), "bold" if bold else "normal")

    def _visible(self, tag) -> bool:
        return self.LAYOUT[tag][3]

    def _place(self, widget, tag, natural_height=False):
        """place at the .fig Position.  A MATLAB popupmenu keeps the height its
        font needs and sits at the TOP of its Position box; a text uicontrol
        draws its string from the top of the box."""
        pos = self.LAYOUT[tag][0]
        if not self._visible(tag):
            return
        x = max(0.0, pos[0]) * self.CW
        w = pos[2] * self.CW
        top = (self.FH - pos[1] - pos[3]) * self.CH
        if natural_height:
            widget.place(x=round(x), y=round(top), width=round(w), height=int(widget.frame.cget("height")))
        else:
            widget.place(x=round(x), y=round(top), width=round(w), height=round(pos[3] * self.CH))

    def _make_popup(self, tag):
        tk = self.tk
        fs = self.LAYOUT[tag][1]
        f = self._font(fs, bold=(tag == "ScanThru"))
        h = round((1.35 * fs + 2.5) * self.ppt)          # MATLAB's popupmenu height for this FontSize
        cb = _PopupMenu(self.root, self.menus[tag], f, h, bg="#FFFFCC" if tag == "ScanThru" else "white",
                        command=lambda t=tag: self._on_popup(t))
        self._place(cb, tag, natural_height=True)
        self.w[tag], self.v[tag] = cb, tk.StringVar(value=self.menus[tag][0])

    def _label(self, tag, text, bold=False):
        pos, fs, ha, _ = self.LAYOUT[tag]
        anchor, justify = {"left": ("nw", "left"), "right": ("ne", "right")}.get(ha, ("n", "center"))
        lab = self.tk.Label(self.root, text=text, font=self._font(fs, bold), anchor=anchor, justify=justify,
                            bg=self.BG, fg="black", padx=0, pady=0, bd=0,
                            wraplength=round(pos[2] * self.CW))
        self._place(lab, tag)
        self.w[tag] = lab

    def _button(self, tag, text, cmd, fg=None):
        """Windows-style MATLAB pushbutton: near-white face, thin light-grey border."""
        tk = self.tk
        fr = tk.Frame(self.root, bg="#CCCCCC", bd=0, highlightthickness=0)
        b = tk.Button(fr, text=text, command=cmd, font=self._font(self.LAYOUT[tag][1]),
                      bg="#FDFDFD", activebackground="#E5F1FB", fg=fg or "black",
                      disabledforeground="#A0A0A0", relief="flat", bd=0, highlightthickness=0,
                      overrelief="flat", padx=1, pady=0)
        b.pack(fill="both", expand=True, padx=1, pady=1)
        b.bind("<Enter>", lambda e: fr.configure(bg="#0078D7") if str(b.cget("state")) != "disabled" else None,
               add="+")
        b.bind("<Leave>", lambda e: fr.configure(bg="#CCCCCC"), add="+")
        self._place(fr, tag)
        self.w[tag] = b

    def _set(self, tag, text):
        self.w[tag].configure(text=text)
        self.root.update_idletasks()

    def _row(self, tag) -> int:
        return self.w[tag].current()

    def _sel(self) -> Dict[str, str]:
        return {t: self.menus[t][max(0, self._row(t))] for t in self.POPUPS}

    # ---------------------------------------------------------------- dialogs
    def _msgbox(self, msg):
        tk = self.tk
        top = tk.Toplevel(self.root)
        top.withdraw()                          # shown by _center_on_screen
        top.title("")
        tk.Label(top, text=msg, padx=24, pady=12).pack()
        tk.Button(top, text="OK", width=8, command=top.destroy).pack(pady=(0, 10))
        self._center_on_screen(top)
        self.root.update()

    def _error(self, title, lines):
        from tkinter import messagebox
        messagebox.showerror(title, "\n".join(lines), parent=self.root)

    def _text_window(self, title, text):
        tk = self.tk
        top = tk.Toplevel(self.root)
        top.title(title)
        t = tk.Text(top, width=100, height=min(40, text.count("\n") + 3), font=("Courier", 9))
        t.insert("1.0", text)
        t.configure(state="disabled")
        t.pack(fill="both", expand=True)
        tk.Button(top, text="Close", command=top.destroy).pack(pady=4)

    def _cache_get(self, key):
        if self.cache["armed"] and key in self.cache["store"]:
            return True, self.cache["store"][key]
        return False, None

    def _cache_set(self, key, val):
        if self.cache["armed"]:
            self.cache["store"][key] = val

    def _work_area(self):
        """(left, top, right, bottom) of the screen the AQP panel is on.  On Windows
        this is the monitor's work area (taskbar excluded) from the Win32 API, since
        Tk only knows the primary screen; elsewhere the Tk screen size."""
        r = self.root
        if sys.platform == "win32":
            try:
                import ctypes
                from ctypes import wintypes

                class MONITORINFO(ctypes.Structure):
                    _fields_ = [("cbSize", wintypes.DWORD), ("rcMonitor", wintypes.RECT),
                                ("rcWork", wintypes.RECT), ("dwFlags", wintypes.DWORD)]
                u32 = ctypes.windll.user32
                u32.MonitorFromPoint.restype = wintypes.HMONITOR
                u32.MonitorFromPoint.argtypes = [wintypes.POINT, wintypes.DWORD]
                pt = wintypes.POINT(r.winfo_rootx() + r.winfo_width() // 2, r.winfo_rooty() + r.winfo_height() // 2)
                hmon = u32.MonitorFromPoint(pt, 2)              # MONITOR_DEFAULTTONEAREST
                mi = MONITORINFO()
                mi.cbSize = ctypes.sizeof(MONITORINFO)
                if u32.GetMonitorInfoW(hmon, ctypes.byref(mi)):
                    w = mi.rcWork
                    return w.left, w.top, w.right, w.bottom
            except Exception:
                pass
        return 0, 0, r.winfo_screenwidth(), r.winfo_screenheight()

    def _center_on_screen(self, top, w=None, h=None):
        """Show a dialog in the middle of the screen, in front, so it cannot be missed.
        The dialog is positioned while hidden and then shown: on Windows a position
        given to a window that is already mapped (or not yet mapped and then shown
        by the window manager) can be ignored, which parks it in the top-left
        corner.  The position is applied once more just after it appears."""
        top.withdraw()
        top.update_idletasks()
        w = int(w or top.winfo_reqwidth())
        h = int(h or top.winfo_reqheight())
        L, T, R, B = self._work_area()
        x = L + max(0, (R - L - w) // 2)
        y = T + max(0, (B - T - h) // 2)
        geo = f"{w}x{h}+{x}+{y}"
        top.geometry(geo)
        top.deiconify()
        top.lift()
        try:
            top.attributes("-topmost", True)
            top.after(400, lambda: top.winfo_exists() and top.attributes("-topmost", False))
        except Exception:
            pass

        def again():
            try:
                if top.winfo_exists():
                    top.geometry(geo)
            except Exception:
                pass
        top.after(20, again)
        top.after(150, again)
        top.focus_force()

    def _pick_analyte(self, analytes):
        """uicellect: pick the target analyte (a large dialog in the middle of the screen)"""
        key = "analyte|" + "|".join(analytes)
        hit, val = self._cache_get(key)
        if hit:
            return val
        tk = self.tk
        ppt = self.ppt
        top = tk.Toplevel(self.root, bg=self.BG)
        top.withdraw()                          # shown by _center_on_screen
        top.title("UICELLECT  --  pick the target analyte")
        top.transient(self.root)
        top.resizable(True, True)
        fl, fi, fb = self._font(16, True), self._font(16), self._font(14, True)
        tk.Label(top, text="Select from the following:", font=fl, bg=self.BG, anchor="w"
                 ).pack(fill="x", padx=round(10 * ppt), pady=(round(10 * ppt), round(6 * ppt)))
        n = len(analytes)
        lb = tk.Listbox(top, height=min(12, max(4, n)), exportselection=False, font=fi, activestyle="none",
                        selectbackground="#0078D7", selectforeground="white", bd=1, relief="solid",
                        highlightthickness=0)
        for a in analytes:
            lb.insert("end", "  " + a)
        lb.selection_set(0)
        lb.activate(0)
        lb.pack(fill="both", expand=True, padx=round(10 * ppt))
        out = {"v": None}

        def ok(_e=None):
            s_ = lb.curselection()
            out["v"] = analytes[s_[0]] if s_ else None
            top.destroy()
        fr = tk.Frame(top, bg=self.BG)
        fr.pack(fill="x", padx=round(10 * ppt), pady=round(10 * ppt))
        tk.Button(fr, text="Cancel", font=fb, width=9, command=top.destroy).pack(side="right")
        tk.Button(fr, text="FINISH", font=fb, width=9, bg="#FDFDFD", command=ok).pack(side="right",
                                                                                     padx=round(8 * ppt))
        lb.bind("<Double-Button-1>", ok)
        top.bind("<Return>", ok)
        top.bind("<Escape>", lambda e: top.destroy())
        top.update_idletasks()
        w = max(round(230 * ppt), top.winfo_reqwidth())      # ~ 2.5x the old 160 x 130 px box
        h = max(round(150 * ppt), top.winfo_reqheight())
        self._center_on_screen(top, w, h)
        lb.focus_set()
        top.grab_set()
        self.root.wait_window(top)
        if out["v"] is None:
            raise AQPError("no analyte picked")
        self._cache_set(key, out["v"])
        return out["v"]

    def _pick_ppn_xlsx(self):
        hit, val = self._cache_get("PP1_PP2_xlsx")
        if hit:
            return val
        from tkinter import filedialog
        p = filedialog.askopenfilename(parent=self.root, title="Pick PP1_PP2*.xlsx", initialdir=os.getcwd(),
                                       filetypes=[("PP1_PP2 xlsx", "*.xlsx"), ("all", "*")])
        p = p or None
        if p:
            self._cache_set("PP1_PP2_xlsx", p)
        return p

    def _pick_user_pls(self, scan):
        from tkinter import simpledialog
        k = simpledialog.askinteger("User-Pick", f"PLSfactor ({scan[0]}..{scan[-1]})", parent=self.root,
                                    minvalue=scan[0], maxvalue=scan[-1])
        return float("nan") if k is None else float(k)

    # ---------------------------------------------------------------- figures
    def _show_figure(self, fig, name):
        tk = self.tk
        from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg, NavigationToolbar2Tk
        top = tk.Toplevel(self.root)
        top.title(f"Figure {fig._aqp_num}: {name}")
        k = len(self.fig_windows) % 10
        top.geometry(f"+{60 + 30 * k}+{40 + 30 * k}")
        canvas = FigureCanvasTkAgg(fig, master=top)
        try:
            fig.tight_layout()
        except Exception:
            pass
        canvas.draw()
        NavigationToolbar2Tk(canvas, top).update()
        canvas.get_tk_widget().pack(fill="both", expand=True)
        for hook in getattr(fig, "_aqp_on_canvas", None) or []:   # e.g. enable_datatips
            try:
                hook(canvas)
            except Exception:                                      # noqa: BLE001
                traceback.print_exc()
        # key presses (Delete/Escape for data tips) need keyboard focus on the canvas
        canvas.get_tk_widget().bind("<Button-1>", lambda e: e.widget.focus_set(), add="+")
        fig._aqp_window = top
        self.fig_windows.append(top)
        top.protocol("WM_DELETE_WINDOW", lambda: self._close_window(top))
        self.root.update()

    def _close_window(self, top):
        try:
            top.destroy()
        except Exception:
            pass
        if top in self.fig_windows:
            self.fig_windows.remove(top)

    def close_all_figs(self):
        """close figs  (cab: every figure but the panel)"""
        for w in list(self.fig_windows):
            self._close_window(w)
        self.fig_windows = []
        self.figmgr.figs, self.figmgr.kept = [], []

    # ---------------------------------------------------------------- state
    def _progress(self, s):
        if s.startswith("iter-"):
            self._set("CurStatus", s)
        self.root.update()

    def _lock(self, on: bool):
        for t in ["Run_AQP", "Cmp_Results", "ScanThru", "ShowResults", "ClearResults", "StepbyStep_plots",
                  "Load_DS"]:
            w = self.w[t]
            if on:
                w._prev_state = str(w.cget("state"))
                w.configure(state="disabled")
            else:
                w.configure(state=getattr(w, "_prev_state", "normal"))
        self.busy = on
        self.root.update_idletasks()

    def _cmp_folder(self):
        return (self.OUT_Cmp_Results or {}).get("folder") or ""

    def update_cmp_button(self, n=None):
        """CmpR_update_Cmp_Results_button"""
        if n is None:
            n = scan_count_results(self._cmp_folder())
        self._set("sNumResults", f"# results = {n}")
        b = self.w["Cmp_Results"]
        if n >= 2:
            b.configure(state="normal", bg="#66cc66", fg="black", activebackground="#7ad67a")
        else:
            b.configure(state="disabled", bg="#D4D4D4", disabledforeground="#737373")
        for t in ("ShowResults", "ClearResults"):
            self.w[t].configure(state="normal" if n >= 1 else "disabled")
        return n

    # ---------------------------------------------------------------- callbacks
    def _on_popup(self, tag):
        if self.busy:
            return
        if tag == "Load_DS":
            self.on_load_ds()
        elif tag == "StepbyStep_plots":
            self.on_stepbystep()
        elif tag == "ScanThru":
            self.on_scanthru()

    def on_load_ds(self):
        from tkinter import filedialog
        disp_with_border("Loading Dataset by --> " + self.w["suserInitials"].cget("text"))
        p = filedialog.askdirectory(parent=self.root, initialdir=os.getcwd(), title="Pls Pick XLSX Dataset folder")
        if not p:
            disp("Load_DS: cancelled by user - dataset folder left unchanged.")
            return
        self.path_XLSX = self.path_DS = os.path.abspath(p)
        self._set("Loaded_DS_folder", self.path_DS)

    def _run_aqp_core(self) -> BatchResult:
        """Run_AQP: read the panel, run the batch, update the panel."""
        if not self.path_XLSX:
            raise AQPError("No dataset folder has been loaded yet - use Load_DS first")
        sel = self._sel()
        if sel["CabXfer_scheme"] == "CabXfer_scheme":
            raise AQPError("pls pick a calibration transfer scheme")
        self.close_all_figs()
        st = settings_from_panel(sel, self.path_XLSX, self.workdir, pick_PPn_xlsx=self._pick_ppn_xlsx,
                                 pick_user_PLS=self._pick_user_pls, menus=self.menus)
        if st.pathfname_PPn_xlsx:      # both menus point at the xlsx, as in Run_AQP
            for t in ("cPP1", "cPP2"):
                self.w[t].current(self.menus[t].index("PP1_PP2_xlsx"))
        self._set("Cur_PP1", f"Scan Thru  -->  {len(st.cPP1)}  PP1 schemes" if len(st.cPP1) > 1 else st.cPP1[0])
        self._set("Cur_PP2", f"Scan Thru  -->  {len(st.cPP2)}  PP2 schemes" if len(st.cPP2) > 1
                  else "PP2 --> " + st.cPP2[0])
        res = batch_run(st, self.figmgr, progress=self._progress, pick_analyte=self._pick_analyte)
        self.OUT_Cmp_Results = {"folder": res.Path4OUT_cln_AQP_FinalSubfolder,
                                "fname": res.fname_Results_AQP_OUT_cln}
        self.out_SbS = res.out_SbS
        self.last_result = res
        if res.dataset is not None:
            self.v["Val_Exist"].set(1 if res.dataset.Val is not None else 0)
        self.update_cmp_button()
        for o in res.OUT_cln:
            print(f"  {o['cLegendLoop_jLine']}  {o['CabXfer_scheme']}  OpmPLSfactor={o['Results4sOpmPLSfactor']}"
                  f"  RMSEP={o['Results4sRMSE_Val']}", flush=True)
        return res

    def on_run_aqp(self):
        if self.busy:
            return
        self._lock(True)
        try:
            self._run_aqp_core()
        except AQPError as e:
            if str(e) == "pls pick a calibration transfer scheme":
                from tkinter import messagebox
                messagebox.showwarning("Run AQP", "please pick a calibration transfer scheme", parent=self.root)
            else:
                self._set("CurStatus", "AQP run FAILED - see command window")
                print("AQP:", e, file=sys.stderr)
                self._error("Run AQP failed", ["The AQP run did not complete.", "", str(e)])
        except Exception as e:     # noqa: BLE001
            self._set("CurStatus", "AQP run FAILED - see command window")
            traceback.print_exc()
            self._error("Run AQP failed", ["The AQP run did not complete.", "", f"{type(e).__name__}: {e}"])
        finally:
            self._lock(False)
            self.update_cmp_button()

    def on_cmp_results(self):
        if not self.OUT_Cmp_Results:
            self._error("Nothing to compare", ["There are no results to compare yet.", "",
                                               "Press Run AQP first."])
            return
        n = self.update_cmp_button()
        if n < 2:
            self._error("Nothing to compare", ["There are fewer than two Results_AQP_~.mat files to compare.",
                                               "", "Compare Results needs at least two of them in",
                                               self._cmp_folder()])
            return
        try:
            cmp_results(self._cmp_folder(), self.figmgr)
        except Exception as e:     # noqa: BLE001
            traceback.print_exc()
            self._error("Compare Results failed", ["The results comparison did not complete.", "", str(e)])

    def on_show_results(self):
        txt, n = list_results_text(self._cmp_folder() or None)
        print(txt, flush=True)
        self._text_window("Results_AQP_~.mat", txt)
        self.update_cmp_button(n)

    def on_clear_results(self):
        from tkinter import messagebox
        folder = self._cmp_folder()
        if not folder or not os.path.isdir(folder):
            messagebox.showwarning("Clear Results", "There is no results folder yet - press Run AQP first.",
                                   parent=self.root)
            self.update_cmp_button(0)
            return
        n = scan_count_results(folder)
        if n == 0:
            print(f"[clear results] nothing to delete in --> {folder}")
            self.update_cmp_button(0)
            return
        if not messagebox.askyesno("Clear Results", f"Delete all {n} Results_AQP_~.mat file(s) in\n{folder}\n\n"
                                   "This cannot be undone.", default="no", parent=self.root):
            print("[clear results] cancelled")
            self.update_cmp_button()
            return
        self.update_cmp_button(clear_results(folder))

    def on_stepbystep(self):
        kind = self.menus["StepbyStep_plots"][self._row("StepbyStep_plots")]
        self.w["StepbyStep_plots"].current(0)
        if kind == "StepbyStep":
            print("=" * 31 + "\npls pick one of the following : \n" + "-" * 30)
            print("\n".join(self.menus["StepbyStep_plots"][1:]) + "\n" + "=" * 31)
            return
        try:
            stepbystep_plots(kind, self.out_SbS, self.path_DS, self.figmgr,
                             self.menus["CabXfer_scheme"][max(0, self._row("CabXfer_scheme"))],
                             (self.OUT_Cmp_Results or {}).get("folder"), (self.OUT_Cmp_Results or {}).get("fname"))
        except Exception as e:     # noqa: BLE001
            traceback.print_exc()
            self._error("StepbyStep_plots failed", ["The step-by-step plots did not complete.", "", str(e)])

    # ---------------------------------------------------------------- Scan Thru
    def _scan_status(self, short, full=None):
        self._set("ScanThruState", short)

    def on_scanthru(self):
        v = self._row("ScanThru")
        self.w["ScanThru"].current(0)
        self.root.update_idletasks()
        try:
            if v == 1:
                self._scan_run(None)
            elif v == 2:
                if not self.scan_LOG:
                    self._error("Scan Thru", ["There is no earlier scan to re-run.", "",
                                              'Pick "Set up & run scan" first.'])
                else:
                    self._scan_run(self.scan_LOG["spec_full"])
            elif v == 3:
                txt = scan_summary(self.scan_LOG, self._cmp_folder())
                print(txt, flush=True)
                self._text_window("Scan Thru summary", txt)
            elif v == 4:
                self._scan_collect_dialog()
        except Exception as e:     # noqa: BLE001
            traceback.print_exc()
            self._error("Scan Thru failed", ["Scan Thru did not complete.", "", str(e)])

    def _scan_run(self, spec_in):
        from tkinter import messagebox
        S = SCANTHRU_SETTINGS
        if not self.path_XLSX:
            self._error("Scan Thru", ["No dataset folder has been loaded yet.", "",
                                      "Scan Thru does nothing more than press Run AQP once per",
                                      "combination, so it needs the same input - use Load_DS first."])
            return
        if spec_in is not None:
            spec = spec_in
        else:
            R = scan_registry({t: self.menus[t] for t in self.POPUPS}, {t: self._row(t) for t in self.POPUPS})
            if not R:
                self._error("Scan Thru", ["No scannable popup menus were found on this panel."])
                return
            spec = self._scan_setup_dialog(R)
            if spec is None:
                return
        P = spec["par"]
        N = len(scan_combinations(P))
        if N < 1:
            return
        if N > S["max_combinations"] and not messagebox.askyesno(
                "Scan Thru", f"That is {N} runs.\n\nEach one is a full Run AQP, so this can take a long time.\n"
                             "Carry on?", default="no", parent=self.root):
            return
        V0 = {p["tag"]: self._row(p["tag"]) for p in P}
        self._lock(True)
        self.cache["armed"] = True
        if not S["cache_across_sweeps"]:
            self.cache["store"] = {}
        LOG = None
        err = None

        def run_combo(row):
            for i, p in enumerate(P):
                self.w[p["tag"]].current(p["vidx"][p["sel"][row[i]]])
            self.root.update()
            try:
                self._run_aqp_core()
            except Exception as e:
                e._aqp_folder = (self._cmp_folder(), (self.OUT_Cmp_Results or {}).get("fname", ""))
                raise
            return self._cmp_folder(), (self.OUT_Cmp_Results or {}).get("fname", "")

        try:
            LOG = scan_sweep(P, run_combo, status=self._scan_status,
                             after_run=(self.close_all_figs if spec["close_figs"] else None))
        except Exception as e:     # noqa: BLE001
            err = e
        for t, r in V0.items():
            self.w[t].current(r)
        self._lock(False)
        self.cache["armed"] = False
        if not S["cache_across_sweeps"]:
            self.cache["store"] = {}
        if err is not None:
            self._scan_status("error")
            traceback.print_exception(type(err), err, err.__traceback__)
            self._error("Scan Thru failed", ["The scan stopped on an error.", "", str(err), "",
                                             "The menus have been put back where they were."])
            return
        LOG["spec_full"] = spec
        self.scan_LOG = LOG
        self._scan_finish(LOG)
        print(scan_summary(LOG, self._cmp_folder()), flush=True)
        nOK = sum(1 for s in LOG["status"][:LOG["nrun"]] if s == "ok")
        self._scan_status(f"{nOK}/{LOG['nrun']} done")
        self.update_cmp_button()

    def _point_at(self, dest):
        if not self.OUT_Cmp_Results:
            self.OUT_Cmp_Results = {"folder": dest, "fname": ""}
        else:
            self.OUT_Cmp_Results["folder"] = dest

    def _scan_finish(self, LOG):
        cf = scan_distinct_folders(LOG)
        if len(cf) <= 1:
            return
        print(f"\n[Scan Thru] the sweep wrote into {len(cf)} different comparison folders:")
        for i, f in enumerate(cf, 1):
            print(f"   [{i}] {f}\n        {scan_count_results(f)} Results_AQP_~.mat")
        policy = SCANTHRU_SETTINGS["consolidate_policy"]
        if policy == "ask":
            policy = self._ask3("Scan Thru -- more than one results folder",
                                f"This scan wrote into {len(cf)} different comparison folders.\n\n"
                                "Compare Results reads ONE folder, so it can only see part of the\n"
                                "sweep as things stand.  Collecting copies into a single folder\n"
                                "lets it plot them all - the x-axis then becomes the union of the\n"
                                "runs, and Compare Results reports whatever is missing per series.",
                                [("Collect into one folder", "collect"), ("Point at the fullest", "largest"),
                                 ("Leave as is", "leave")])
        if policy == "collect":
            dest = scan_collect(cf)
            if not dest:
                print("[Scan Thru] could not create the collection folder -- left as is")
                return
            self._point_at(dest)
            print(f"[Scan Thru] collected {scan_count_results(dest)} file(s) into\n   {dest}")
        elif policy == "largest":
            dest = max(cf, key=scan_count_results)
            self._point_at(dest)
            print(f"[Scan Thru] Compare Results now points at\n   {dest}")
        else:
            print("[Scan Thru] folders left as they are")

    def _ask3(self, title, msg, choices):
        tk = self.tk
        top = tk.Toplevel(self.root)
        top.withdraw()                          # shown by _center_on_screen
        top.title(title)
        top.transient(self.root)
        tk.Label(top, text=msg, justify="left", padx=14, pady=10).pack()
        out = {"v": choices[-1][1]}
        fr = tk.Frame(top)
        fr.pack(pady=8)
        for lab, val in choices:
            tk.Button(fr, text=lab, command=lambda v=val: (out.update(v=v), top.destroy())).pack(side="left", padx=4)
        self._center_on_screen(top)
        top.grab_set()
        self.root.wait_window(top)
        return out["v"]

    def _scan_collect_dialog(self):
        if not self.scan_LOG:
            self._error("Scan Thru", ["There is no scan to collect from yet."])
            return
        cf = scan_distinct_folders(self.scan_LOG)
        if not cf:
            self._error("Scan Thru", ["The last scan produced no results folder."])
            return
        dest = scan_collect(cf)
        if not dest:
            self._error("Scan Thru", ["Could not create the collection folder."])
            return
        self._point_at(dest)
        print(f"[Scan Thru] collected {scan_count_results(dest)} file(s) into\n   {dest}")
        self.update_cmp_button()

    def _scan_setup_dialog(self, R):
        """ScanThru_setup_dialog: parameters left, their options right (multi-select)."""
        tk = self.tk
        S = SCANTHRU_SETTINGS
        top = tk.Toplevel(self.root)
        top.withdraw()                          # shown by _center_on_screen
        top.title("Scan Thru  --  pick the hyperparameters and their options")
        top.transient(self.root)
        top.geometry("1000x620")
        st = {"sel": [[r["cur"]] for r in R], "cur": 0, "ok": False}
        F = lambda fs, b=False: ("TkDefaultFont", fs, "bold" if b else "normal")   # noqa: E731
        tk.Label(top, text="Tick the options to scan.  One option = that parameter stays fixed.",
                 font=F(14, True), anchor="w").place(relx=0.02, rely=0.02, relwidth=0.96, relheight=0.06)
        tk.Label(top, text="Hyperparameter", font=F(11), anchor="w").place(relx=0.02, rely=0.09, relwidth=0.43,
                                                                           relheight=0.05)
        tk.Label(top, text="Options  (ctrl-click / shift-click)", font=F(11), anchor="w").place(
            relx=0.47, rely=0.09, relwidth=0.51, relheight=0.05)
        lbP = tk.Listbox(top, exportselection=False, font=("Courier", 11))
        lbP.place(relx=0.02, rely=0.14, relwidth=0.43, relheight=0.56)
        lbO = tk.Listbox(top, exportselection=False, selectmode="extended", font=("Courier", 11))
        lbO.place(relx=0.47, rely=0.14, relwidth=0.51, relheight=0.56)
        cnt = tk.Label(top, text="# runs = 1", font=F(16, True), anchor="w")
        cnt.place(relx=0.02, rely=0.80, relwidth=0.50, relheight=0.09)
        nt = tk.Label(top, text="", font=F(9), anchor="w", justify="left", wraplength=720)
        nt.place(relx=0.02, rely=0.90, relwidth=0.73, relheight=0.08)
        cf_var = tk.IntVar(value=S["close_figs_between_runs"])
        tk.Checkbutton(top, text="close plot figures between runs", variable=cf_var, font=F(11),
                       anchor="w").place(relx=0.02, rely=0.72, relwidth=0.43, relheight=0.065)

        def refresh():
            N, nvary, nres = 1, 0, 0
            lbP.delete(0, "end")
            for i, r in enumerate(R):
                if not st["sel"][i]:
                    st["sel"][i] = [r["cur"]]
                k = len(st["sel"][i])
                N *= k
                nvary += k > 1
                nres += bool(r["resolved"])
                tail = f"{k} options" if k > 1 else r["options"][st["sel"][i][0]]
                lbP.insert("end", f"{'*' if r['resolved'] else ' '}{r['tag']:<21s} [{k:2d}]  {tail}")
            lbP.selection_clear(0, "end")
            lbP.selection_set(st["cur"])
            lbO.delete(0, "end")
            for o in R[st["cur"]]["options"]:
                lbO.insert("end", o)
            for j in st["sel"][st["cur"]]:
                lbO.selection_set(j)
            cnt.configure(text=f"# runs = {N}", fg="#b30000" if N >= 200 else "black")
            s = "Nothing varied yet -- this would be one ordinary Run AQP." if nvary == 0 else \
                f"{nvary} parameter(s) varying."
            if nres:
                s += "   * panel is on that menu's title row; the option shown is used."
            nt.configure(text=s)

        def on_param(_e=None):
            s = lbP.curselection()
            if s:
                st["cur"] = s[0]
                refresh()

        def on_opt(_e=None):
            s = list(lbO.curselection())
            st["sel"][st["cur"]] = s if s else [R[st["cur"]]["cur"]]
            refresh()

        def sel_all():
            st["sel"][st["cur"]] = list(range(len(R[st["cur"]]["options"])))
            refresh()

        def reset():
            st["sel"][st["cur"]] = [R[st["cur"]]["cur"]]
            refresh()

        def run():
            st["ok"] = True
            top.destroy()
        lbP.bind("<<ListboxSelect>>", on_param)
        lbO.bind("<<ListboxSelect>>", on_opt)
        tk.Button(top, text="Select all", font=F(11), command=sel_all).place(relx=0.47, rely=0.72, relwidth=0.17,
                                                                             relheight=0.065)
        tk.Button(top, text="Reset to current", font=F(11), command=reset).place(relx=0.655, rely=0.72,
                                                                                 relwidth=0.225, relheight=0.065)
        tk.Button(top, text="Run the scan", font=F(14, True), bg="#66cc66", command=run).place(
            relx=0.755, rely=0.80, relwidth=0.225, relheight=0.10)
        tk.Button(top, text="Cancel", font=F(11), command=top.destroy).place(relx=0.755, rely=0.91,
                                                                             relwidth=0.225, relheight=0.07)
        refresh()
        self._center_on_screen(top, 1000, 620)
        top.grab_set()
        self.root.wait_window(top)
        if not st["ok"]:
            return None
        P = []
        for i, r in enumerate(R):
            p = dict(r)
            p["sel"] = list(st["sel"][i])
            P.append(p)
        return {"par": P, "close_figs": bool(cf_var.get())}

    def _quit(self):
        self.close_all_figs()
        self.root.destroy()


def gui_main(args) -> int:
    VERB.disp = 0 if args.d is None else args.d
    VERB.warn = 0 if args.w is None else args.w
    try:
        import tkinter as tk
    except ImportError:
        print("tkinter is not available in this Python - install it (e.g. python3-tk) or use --cli",
              file=sys.stderr)
        return 2
    if sys.platform == "win32":     # be DPI aware like MATLAB, so text is crisp at 125 / 150 %
        try:
            import ctypes
            ctypes.windll.shcore.SetProcessDpiAwareness(1)
        except Exception:
            try:
                ctypes.windll.user32.SetProcessDPIAware()
            except Exception:
                pass
    try:
        root = tk.Tk()
    except tk.TclError as e:
        print(f"cannot open the GUI ({e}); use --cli for a run without a display", file=sys.stderr)
        return 2
    import matplotlib
    matplotlib.use("TkAgg")
    app = AQPApp(root, args.workdir or os.getcwd())
    if args.ds:
        app.path_XLSX = app.path_DS = os.path.abspath(args.ds)
        app._set("Loaded_DS_folder", app.path_DS)
    disp_with_border("Current User is --> " + app.w["suserInitials"].cget("text"))
    root.mainloop()
    return 0


def main(argv=None) -> int:
    args = build_arg_parser().parse_args(argv)
    warnings.filterwarnings("ignore", category=RuntimeWarning)
    if args.cli:
        return cli_main(args)
    return gui_main(args)


if __name__ == "__main__":
    sys.exit(main())
